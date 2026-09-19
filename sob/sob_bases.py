# -*- coding: utf-8 -*-
"""
sob_bases.py - Mapeamento e Resolução de Bases Técnicas da Enel Rio e FABAO.
Responsabilidade única: Conversão de nomes de bases, CECOs e municípios
para siglas canônicas (ex: Petrópolis -> PE, Cabo Frio -> CF) e geração do
nome padronizado do arquivo DWG ({BASE}_{SOB}_DESENHO.dwg).
"""

import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# Caminho padrão da pasta FABAO
CAMINHO_FABAO_PADRAO = (
    Path(os.environ.get("USERPROFILE", "C:/Users/jonat"))
    / "OneDrive"
    / "Área de Trabalho"
    / "FABAO"
)

# Tabela Canônica de Siglas de Bases (Enel Distribuição Rio)
TABELA_BASES_PADRAO: Dict[str, str] = {
    "PETRÓPOLIS": "PE",
    "PETROPOLIS": "PE",
    "CABO FRIO": "CF",
    "ANGRA DOS REIS": "AN",
    "ANGRA": "AN",
    "RESENDE": "RE",
    "NITERÓI": "NI",
    "NITEROI": "NI",
    "MAGÉ": "MA",
    "MAGE": "MA",
    "SÃO GONÇALO": "SG",
    "SAO GONCALO": "SG",
    "TERESÓPOLIS": "TE",
    "TERESOPOLIS": "TE",
    "CAMPOS DOS GOYTACAZES": "CA",
    "CAMPOS": "CA",
    "MACAÉ": "MC",
    "MACAE": "MC",
    "ITAPERUNA": "IT",
    "SANTO ANTÔNIO DE PÁDUA": "PA",
    "PADUA": "PA",
}

# Lista Ordenada para Menus e Caixas de Seleção Gráfica
LISTA_BASES_PADRAO: List[Tuple[str, str]] = [
    ("PE", "Petrópolis"),
    ("CF", "Cabo Frio"),
    ("AN", "Angra dos Reis"),
    ("RE", "Resende"),
    ("NI", "Niterói"),
    ("MA", "Magé"),
    ("SG", "São Gonçalo"),
    ("TE", "Teresópolis"),
    ("CA", "Campos"),
    ("MC", "Macaé"),
    ("IT", "Itaperuna"),
    ("PA", "Santo Antônio de Pádua"),
]

# Mapeamento CECO -> Base
TABELA_CECO_BASE: Dict[str, str] = {
    "1550": "AN",  # Angra
    "1560": "RE",  # Resende
    "1610": "NI",  # Niterói
    "1650": "MA",  # Magé
    "1660": "SG",  # São Gonçalo
    "1670": "PE",  # Petrópolis
    "1680": "TE",  # Teresópolis
}

# Lista amigável para exibição em interfaces
OPCOES_BASES_UI: List[Tuple[str, str]] = [
    ("PE", "Petrópolis"),
    ("CF", "Cabo Frio"),
    ("AN", "Angra dos Reis"),
    ("RE", "Resende"),
    ("NI", "Niterói"),
    ("MA", "Magé"),
    ("SG", "São Gonçalo"),
    ("TE", "Teresópolis"),
    ("CA", "Campos"),
    ("MC", "Macaé"),
    ("IT", "Itaperuna"),
]


def normalizar_sigla_base(base_input: Optional[str]) -> str:
    """
    Normaliza a sigla da base para 2 caracteres maiúsculos.
    Exemplos: 'pe' -> 'PE', 'Petrópolis' -> 'PE', '1670' -> 'PE'.
    Se não reconhecida ou vazia, retorna 'PE' como padrão.
    """
    if not base_input:
        return "PE"

    txt = str(base_input).strip().upper()

    # Se já for uma sigla conhecida de 2 letras
    for _, sigla in TABELA_BASES_PADRAO.items():
        if txt == sigla:
            return sigla

    # Se for CECO numérico
    if txt in TABELA_CECO_BASE:
        return TABELA_CECO_BASE[txt]

    # Busca no dicionário de nomes
    for chave, sigla in TABELA_BASES_PADRAO.items():
        if chave in txt or txt in chave:
            return sigla

    # Se tiver exatamente 2 letras, usa as duas letras
    if len(txt) == 2 and txt.isalpha():
        return txt

    return "PE"


def gerar_nome_dwg_sob(base: Optional[str], sob_id: str) -> str:
    """
    Gera o nome canônico do arquivo DWG conforme a regra:
    {BASE}_{SOB}_DESENHO.dwg
    Exemplo: PE_A057350816_DESENHO.dwg
    """
    sigla = normalizar_sigla_base(base)
    sob_limpa = str(sob_id).strip().upper()
    if sob_limpa.startswith("SOB_"):
        sob_limpa = sob_limpa[4:]

    return f"{sigla}_{sob_limpa}_DESENHO.dwg"


def carregar_bases_fabao(caminho_fabao: Optional[Path] = None) -> Dict[str, str]:
    """
    Tenta carregar mapeamento adicional da planilha gomnet_18-08-2022.xlsm no diretório FABAO.
    Retorna cópia atualizada do dicionário de bases.
    """
    pasta = caminho_fabao or CAMINHO_FABAO_PADRAO
    mapa = dict(TABELA_BASES_PADRAO)

    planilha = pasta / "Controle" / "gomnet_18-08-2022.xlsm"
    if not planilha.is_file():
        return mapa

    try:
        import openpyxl

        wb = openpyxl.load_workbook(str(planilha), data_only=True, read_only=True)
        if "REF" in wb.sheetnames:
            ws = wb["REF"]
            # Lê linhas de CECO e Base (Colunas 4 e 5)
            for row in ws.iter_rows(min_row=2, max_row=20, min_col=4, max_col=5, values_only=True):
                ceco, base_nome = row[0], row[1]
                if base_nome:
                    nome_str = str(base_nome).strip().upper()
                    sigla = normalizar_sigla_base(nome_str)
                    mapa[nome_str] = sigla
                    if ceco:
                        TABELA_CECO_BASE[str(ceco).strip()] = sigla
    except Exception:
        pass

    return mapa
