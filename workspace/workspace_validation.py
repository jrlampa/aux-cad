"""Validação baseada em evidências dos entregáveis existentes do Workspace."""

from dataclasses import dataclass, replace
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

import openpyxl
from sob.sob_core import calcular_hash_arquivo


class ValidationStatus(str, Enum):
    OK = "OK"
    ALERTA = "ALERTA"
    ERRO = "ERRO"
    NAO_DISPONIVEL = "NÃO DISPONÍVEL"


@dataclass(frozen=True)
class ValidationSource:
    tipo: str
    arquivo: Optional[str]
    localizacao: Optional[str]
    origem: str
    sha256: Optional[str] = None
    mtime_ns: Optional[int] = None

    def to_dict(self) -> Dict[str, Optional[str]]:
        return {
            "tipo": self.tipo,
            "arquivo": self.arquivo,
            "localizacao": self.localizacao,
            "origem": self.origem,
            "sha256": self.sha256,
            "mtime_ns": self.mtime_ns,
        }


@dataclass(frozen=True)
class ValidationContext:
    execution_id: str
    correlation_id: str
    sob_id: str
    dwg_path: Optional[str]
    dwg_identity: Optional[str]
    database_fingerprint: Optional[str]
    project_root: Path
    cqt_path: Optional[Path]
    revision: Optional[str] = None
    origin: str = "workspace.cli"
    host_verified: bool = False


@dataclass(frozen=True)
class ValidationResult:
    rule_id: str
    description: str
    status: ValidationStatus
    value: Any
    criterion: Any
    source: ValidationSource
    timestamp: str
    message: str
    context: ValidationContext

    def to_dict(self) -> Dict[str, Any]:
        return {
            "rule_id": self.rule_id,
            "description": self.description,
            "status": self.status.value,
            "value": self.value,
            "criterion": self.criterion,
            "source": self.source.to_dict(),
            "timestamp": self.timestamp,
            "message": self.message,
            "execution_id": self.context.execution_id,
            "correlation_id": self.context.correlation_id,
            "sob_id": self.context.sob_id,
            "dwg_path": self.context.dwg_path,
            "dwg_identity": self.context.dwg_identity,
            "database_fingerprint": self.context.database_fingerprint,
            "revision": self.context.revision,
            "origin": self.context.origin,
        }


class ValidationRule:
    rule_id = ""
    description = ""

    def validate(self, context: ValidationContext) -> List[ValidationResult]:
        raise NotImplementedError

    @staticmethod
    def result(
        context: ValidationContext,
        rule_id: str,
        description: str,
        status: ValidationStatus,
        value: Any,
        criterion: Any,
        source: ValidationSource,
        message: str,
    ) -> ValidationResult:
        return ValidationResult(
            rule_id=rule_id,
            description=description,
            status=status,
            value=value,
            criterion=criterion,
            source=source,
            timestamp=datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            message=message,
            context=context,
        )


def _verificar_contexto_validacao(context: ValidationContext) -> None:
    from workspace.workspace_context_cli import _resolver_project_scope

    root = context.project_root.resolve()
    dwg = Path(context.dwg_path).resolve() if context.dwg_path else None
    if (not dwg or not dwg.is_file() or dwg.parent != root
            or not context.sob_id
            or not context.host_verified and _resolver_project_scope(
                root, str(dwg), context.sob_id
            )["status"] != "Bound"):
        raise ValueError("VALIDATION_CONTEXT_INVALID")
    if (context.dwg_identity and context.database_fingerprint
            and context.dwg_identity.casefold()
            != context.database_fingerprint.casefold()):
        raise ValueError("VALIDATION_CONTEXT_INVALID")
    if (context.cqt_path is not None
            and not context.cqt_path.resolve().is_relative_to(root)):
        raise ValueError("VALIDATION_SOURCE_CONTEXT_MISMATCH")


class CqtMaximumVoltageDropRule(ValidationRule):
    rule_id = "CQT.TENSAO_MAXIMA"
    description = "Queda de tensão acumulada máxima por cenário"

    def validate(self, context: ValidationContext) -> List[ValidationResult]:
        _verificar_contexto_validacao(context)
        path = context.cqt_path
        if path is None or not path.is_file():
            return [self._not_available(context, path, "Planilha CQT não encontrada.")]

        from cqt.cqt_live_calc import LIMITE_NORMATIVO_QT
        from cqt.cqt_quadros_generator import extract_trechos_from_sheets

        stat_before = path.stat()
        source_hash = calcular_hash_arquivo(path)

        def source_for(scenario: str, segment_id: str) -> ValidationSource:
            return replace(
                self._source_for(path, workbook, scenario, segment_id),
                sha256=source_hash,
                mtime_ns=stat_before.st_mtime_ns,
            )

        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            extracted = extract_trechos_from_sheets(workbook)
            if not extracted:
                return [self._not_available(
                    context,
                    path,
                    "Não há valores de queda de tensão armazenados nas abas CQT.",
                )]

            results: List[ValidationResult] = []
            scenario_fields = (
                ("ATUAL", "qt_atual"),
                ("ATUAL+NC", "qt_atual_nc"),
                ("PROJETADO", "qt_projetado"),
            )
            for scenario, field in scenario_fields:
                values = []
                for segment in extracted:
                    raw = segment.get(field)
                    if raw in (None, "", "-"):
                        continue
                    try:
                        value = float(str(raw).replace("%", "").replace(",", "."))
                    except (TypeError, ValueError):
                        source = source_for(scenario, segment["trecho"])
                        results.append(self.result(
                            context,
                            self.rule_id,
                            self.description,
                            ValidationStatus.ERRO,
                            raw,
                            {"operator": "<=", "value": LIMITE_NORMATIVO_QT, "unit": "%"},
                            source,
                            f"Valor não numérico no cenário {scenario}.",
                        ))
                        continue
                    values.append((value, segment["trecho"]))

                if not values:
                    continue

                maximum, segment_id = max(values, key=lambda item: item[0])
                source = source_for(scenario, segment_id)
                status = (
                    ValidationStatus.ALERTA
                    if maximum > LIMITE_NORMATIVO_QT
                    else ValidationStatus.OK
                )
                results.append(self.result(
                    context,
                    self.rule_id,
                    f"{self.description} ({scenario})",
                    status,
                    {"percentual": maximum, "trecho": segment_id},
                    {"operator": "<=", "value": LIMITE_NORMATIVO_QT, "unit": "%"},
                    source,
                    f"Cenário {scenario}: trecho {segment_id}, {maximum:.2f}%.",
                ))

            return results or [self._not_available(
                context,
                path,
                "As abas CQT existem, mas não há valores numéricos armazenados.",
            )]
        finally:
            workbook.close()
            stat_after = path.stat()
            if (stat_before.st_size, stat_before.st_mtime_ns) != (
                stat_after.st_size, stat_after.st_mtime_ns
            ) or calcular_hash_arquivo(path) != source_hash:
                raise RuntimeError("SOURCE_CHANGED_DURING_VALIDATION")

    def _not_available(
        self,
        context: ValidationContext,
        path: Optional[Path],
        message: str,
    ) -> ValidationResult:
        return self.result(
            context,
            self.rule_id,
            self.description,
            ValidationStatus.NAO_DISPONIVEL,
            None,
            {"operator": "<=", "value": 6.0, "unit": "%"},
            ValidationSource("excel", str(path) if path else None, None, "CQT existente"),
            message,
        )

    @staticmethod
    def _source_for(
        path: Path,
        workbook: openpyxl.Workbook,
        scenario: str,
        segment_id: str,
    ) -> ValidationSource:
        sheet_name = next((
            candidate for candidate in (
                f"CQT {scenario}", f"CQT_{scenario}", f"CQT-{scenario}", scenario
            ) if candidate in workbook.sheetnames
        ), None)
        if sheet_name is None:
            return ValidationSource("excel", str(path), None, "valor armazenado da planilha CQT")

        sheet = workbook[sheet_name]
        for row in range(11, 35):
            letter = str(sheet.cell(row, 1).value or "").strip()
            length = sheet.cell(row, 3).value
            if letter != str(segment_id) or length in (None, "", 0, "0"):
                continue
            column = 12 if sheet.cell(row, 12).value not in (None, "", " ") else 11
            cell = sheet.cell(row, column).coordinate
            return ValidationSource(
                "excel",
                str(path),
                f"{sheet_name}!{cell}",
                "valor CQT armazenado (data_only)",
            )

        return ValidationSource("excel", str(path), sheet_name, "valor CQT armazenado (data_only)")


class ValidationEngine:
    def __init__(self, rules: Optional[Iterable[ValidationRule]] = None):
        self._rules = {rule.rule_id: rule for rule in (rules or [CqtMaximumVoltageDropRule()])}

    def validate(
        self,
        context: ValidationContext,
        requested_rules: Optional[Iterable[str]] = None,
    ) -> List[ValidationResult]:
        _verificar_contexto_validacao(context)
        rule_ids = list(requested_rules) if requested_rules is not None else list(self._rules)
        results = []
        for rule_id in rule_ids:
            rule = self._rules.get(rule_id)
            if rule is None:
                results.append(ValidationRule.result(
                    context,
                    rule_id,
                    "Regra de validação não registrada",
                    ValidationStatus.NAO_DISPONIVEL,
                    None,
                    None,
                    ValidationSource("none", None, None, "registry de validadores"),
                    "Nenhum validador comprovado para esta regra.",
                ))
                continue
            results.extend(rule.validate(context))
        return results