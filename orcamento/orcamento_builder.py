"""
Módulo Extrator e Compilador de Dados Reais de Orçamento (IM3 Brasil / Concessionária).
Lê planilhas oficiais (sisORÇA.xlsx, KITS E MATERIAIS REV_31, Mão de Obra 31-05-2024.xlsm),
extrai materiais precificados, regras de kits e tabelas de mão de obra contratuais,
e compila em banco SQLite indexado (orcamento_dados.sqlite3) e cache JSON (orcamento_dados.json).
"""

import json
import os
import re
import sqlite3
import sys
from typing import Any, Dict, List, Optional, Tuple

_curr_dir = os.path.dirname(os.path.abspath(__file__))
_root_dir = os.path.abspath(os.path.join(_curr_dir, ".."))
if _root_dir not in sys.path:
    sys.path.insert(0, _root_dir)

DEFAULT_SISORCA_PATH = r"C:\Users\jonat\OneDrive - IM3 Brasil\sisORÇA.xlsx"
DEFAULT_AUX_DIR = r"C:\Users\jonat\OneDrive - IM3 Brasil\Arquivos para auxílio\Orçamento"
DB_PATH = os.path.join(_curr_dir, "orcamento_dados.sqlite3")
JSON_PATH = os.path.join(_curr_dir, "orcamento_dados.json")


def _init_sqlite(db_path: str) -> sqlite3.Connection:
    """Cria tabelas no SQLite se não existirem."""
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    cur.execute("PRAGMA journal_mode = WAL;")
    cur.execute("PRAGMA synchronous = NORMAL;")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS materiais (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo_sie TEXT,
        codigo_e4e TEXT,
        codigo_curto TEXT,
        descricao TEXT NOT NULL,
        unidade TEXT DEFAULT 'UN',
        preco_unitario REAL DEFAULT 0.0,
        categoria TEXT DEFAULT 'Geral',
        tipo TEXT DEFAULT '',
        flag_instala INTEGER DEFAULT 1,
        flag_retira INTEGER DEFAULT 0,
        flag_reuse INTEGER DEFAULT 0
    );
    """)

    cur.execute("CREATE INDEX IF NOT EXISTS idx_mat_sie ON materiais(codigo_sie);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mat_e4e ON materiais(codigo_e4e);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mat_curto ON materiais(codigo_curto);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mat_cat ON materiais(categoria);")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS kits (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        codigo_kit TEXT NOT NULL,
        nome_kit TEXT,
        codigo_material TEXT,
        nome_material TEXT,
        qtd_material REAL DEFAULT 1.0,
        ambiente TEXT DEFAULT 'NA'
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_kits_cod ON kits(codigo_kit);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_kits_amb ON kits(ambiente);")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS mao_de_obra (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        contrato TEXT NOT NULL,
        empreiteira TEXT NOT NULL,
        baremo TEXT,
        voz TEXT NOT NULL,
        texto_breve TEXT,
        territorio TEXT NOT NULL,
        unidade TEXT DEFAULT 'UN',
        custo_unitario REAL DEFAULT 0.0,
        grupo TEXT DEFAULT 'LINHAS AÉREAS',
        descricao_completa TEXT
    );
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mo_emp_terr ON mao_de_obra(empreiteira, territorio);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mo_voz ON mao_de_obra(voz);")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_mo_baremo ON mao_de_obra(baremo);")

    cur.execute("""
    CREATE TABLE IF NOT EXISTS contratos (
        contrato TEXT PRIMARY KEY,
        empreiteira TEXT NOT NULL,
        regiao TEXT NOT NULL,
        unidade_base TEXT
    );
    """)

    conn.commit()
    return conn


def carregar_materiais_sisorca(caminho_sisorca: str) -> List[Dict[str, Any]]:
    """Extrai materiais precificados da aba DADOS e gerais da LISTA GERAL do sisORÇA.xlsx."""
    if not os.path.exists(caminho_sisorca):
        print(f"[AVISO] sisORÇA não encontrado em: {caminho_sisorca}")
        return []

    import openpyxl
    wb = openpyxl.load_workbook(caminho_sisorca, data_only=True, read_only=True)
    materiais: List[Dict[str, Any]] = []

    # 1. Aba DADOS (Postes: colunas 0..6 / Condutores: colunas 8..14)
    if "DADOS" in wb.sheetnames:
        ws_dados = wb["DADOS"]
        for row in ws_dados.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            # Tabela 1 (Postes)
            tipo1 = str(row[0] or "").strip()
            cod_sie1 = str(row[1] or "").strip() if row[1] is not None else ""
            cod_e4e1 = str(row[2] or "").strip() if row[2] is not None else ""
            cod_curto1 = str(row[3] or "").strip() if row[3] is not None else ""
            desc1 = str(row[4] or "").strip()
            un1 = str(row[5] or "UN").strip()
            preco1 = 0.0
            try:
                if row[6] is not None:
                    preco1 = float(str(row[6]).replace(",", "."))
            except (ValueError, TypeError):
                preco1 = 0.0

            if desc1 and (cod_sie1 or cod_e4e1 or cod_curto1):
                materiais.append({
                    "codigo_sie": cod_sie1,
                    "codigo_e4e": cod_e4e1,
                    "codigo_curto": cod_curto1,
                    "descricao": desc1,
                    "unidade": un1,
                    "preco_unitario": preco1,
                    "categoria": "Postes",
                    "tipo": tipo1 or "POSTE",
                    "flag_instala": 1,
                    "flag_retira": 0,
                    "flag_reuse": 0
                })

            # Tabela 2 (Condutores - colunas 8..14)
            if len(row) > 14 and row[12]:
                tipo2 = str(row[8] or "").strip()
                cod_sie2 = str(row[9] or "").strip() if row[9] is not None else ""
                cod_e4e2 = str(row[10] or "").strip() if row[10] is not None else ""
                cod_curto2 = str(row[11] or "").strip() if row[11] is not None else ""
                desc2 = str(row[12] or "").strip()
                un2 = str(row[13] or "M").strip()
                preco2 = 0.0
                try:
                    if row[14] is not None:
                        preco2 = float(str(row[14]).replace(",", "."))
                except (ValueError, TypeError):
                    preco2 = 0.0

                if desc2:
                    materiais.append({
                        "codigo_sie": cod_sie2,
                        "codigo_e4e": cod_e4e2,
                        "codigo_curto": cod_curto2,
                        "descricao": desc2,
                        "unidade": un2,
                        "preco_unitario": preco2,
                        "categoria": "Condutores",
                        "tipo": tipo2 or "CONDUTOR",
                        "flag_instala": 1,
                        "flag_retira": 0,
                        "flag_reuse": 0
                    })

    # 2. Aba LISTA GERAL (13.808 materiais cadastrados para lookup e códigos oficiais)
    if "LISTA GERAL" in wb.sheetnames:
        ws_lg = wb["LISTA GERAL"]
        codigos_ja_inseridos = {m["codigo_sie"] for m in materiais if m["codigo_sie"]}
        codigos_ja_inseridos.update({m["codigo_e4e"] for m in materiais if m["codigo_e4e"]})
        for row in ws_lg.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 5:
                continue
            sie = str(row[0] or "").strip() if row[0] is not None else ""
            e4e = str(row[1] or "").strip() if row[1] is not None else ""
            desc = str(row[2] or "").strip()
            un = str(row[3] or "UN").strip()
            macro = str(row[4] or "MATERIAIS DE REDE").strip()

            if (sie in codigos_ja_inseridos or e4e in codigos_ja_inseridos) or not desc:
                continue

            materiais.append({
                "codigo_sie": sie,
                "codigo_e4e": e4e,
                "codigo_curto": "",
                "descricao": desc,
                "unidade": un,
                "preco_unitario": 0.0,
                "categoria": "Materiais de Rede" if "REDE" in macro.upper() else "Geral",
                "tipo": macro,
                "flag_instala": 1 if len(row) > 5 and row[5] else 0,
                "flag_retira": 1 if len(row) > 6 and row[6] else 0,
                "flag_reuse": 1 if len(row) > 7 and row[7] else 0
            })
            if sie:
                codigos_ja_inseridos.add(sie)
            if e4e:
                codigos_ja_inseridos.add(e4e)

    wb.close()
    return materiais


def carregar_kits_oficiais(pasta_aux: str) -> List[Dict[str, Any]]:
    """Extrai kits de estruturas, trafos, aterramentos e estais das planilhas auxiliares."""
    if not os.path.exists(pasta_aux):
        print(f"[AVISO] Pasta de arquivos auxiliares não encontrada: {pasta_aux}")
        return []

    import openpyxl
    kits: List[Dict[str, Any]] = []

    # 1. KITS E MATERIAIS REV_31.xlsx ou REV_32
    rev31_path = os.path.join(pasta_aux, "KITS E MATERIAIS REV_31.xlsx")
    if not os.path.exists(rev31_path):
        rev31_path = os.path.join(pasta_aux, "E4E_KITS E MATERIAIS REV_32 06 05 20.xlsx")

    if os.path.exists(rev31_path):
        wb = openpyxl.load_workbook(rev31_path, data_only=True, read_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 5:
                    continue
                # Estrutura típica: CODIGO KIT, NOME KIT, CODIGO MATERIAL, NOME MATERIAL, QTD MATERIAL KIT
                c_kit = str(row[0] or "").strip()
                n_kit = str(row[1] or "").strip()
                c_mat = str(row[2] or "").strip() if row[2] is not None else ""
                n_mat = str(row[3] or "").strip() if row[3] is not None else ""
                try:
                    qtd = float(str(row[4]).replace(",", ".")) if row[4] is not None else 1.0
                except (ValueError, TypeError):
                    qtd = 1.0

                if c_kit and c_mat and c_kit != "CODIGO KIT":
                    amb = "AA" if ("AA-" in c_kit or "AGRESSIVO" in n_kit.upper()) else "NA"
                    kits.append({
                        "codigo_kit": c_kit,
                        "nome_kit": n_kit,
                        "codigo_material": c_mat,
                        "nome_material": n_mat,
                        "qtd_material": qtd,
                        "ambiente": amb
                    })

                # Verifica colunas laterais da mesma linha (ex: kits agressivos ao lado em colunas 6 a 10)
                if len(row) >= 10:
                    c_kit2 = str(row[6] or "").strip()
                    n_kit2 = str(row[7] or "").strip()
                    c_mat2 = str(row[8] or "").strip() if row[8] is not None else ""
                    n_mat2 = str(row[9] or "").strip() if row[9] is not None else ""
                    try:
                        qtd2 = float(str(row[10]).replace(",", ".")) if len(row) > 10 and row[10] is not None else 1.0
                    except (ValueError, TypeError):
                        qtd2 = 1.0

                    if c_kit2 and c_mat2 and c_kit2 != "CODIGO KIT":
                        kits.append({
                            "codigo_kit": c_kit2,
                            "nome_kit": n_kit2,
                            "codigo_material": c_mat2,
                            "nome_material": n_mat2,
                            "qtd_material": qtd2,
                            "ambiente": "AA"
                        })
        wb.close()

    # 2. Rede convencional cruzeta polimerica Meio e Beco
    for nome_arq in ["Rede convencional poste CC cruzeta polimerica_Meio.xlsx", "Rede convencional poste CC cruzeta polimerica_Beco.xlsx"]:
        arq_path = os.path.join(pasta_aux, nome_arq)
        if os.path.exists(arq_path):
            wb = openpyxl.load_workbook(arq_path, data_only=True, read_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                for row in ws.iter_rows(min_row=3, values_only=True):
                    if not row or len(row) < 4:
                        continue
                    c_kit = str(row[0] or "").strip()
                    n_kit = str(row[1] or "").strip()
                    c_mat = str(row[2] or "").strip() if row[2] is not None else ""
                    try:
                        qtd = float(str(row[3]).replace(",", ".")) if row[3] is not None else 1.0
                    except (ValueError, TypeError):
                        qtd = 1.0

                    if c_kit and c_mat and c_kit != "CODIGO KIT":
                        kits.append({
                            "codigo_kit": c_kit,
                            "nome_kit": n_kit,
                            "codigo_material": c_mat,
                            "nome_material": n_kit,
                            "qtd_material": qtd,
                            "ambiente": "NA"
                        })
            wb.close()

    return kits


def carregar_mao_de_obra(pasta_aux: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Extrai contratos, baremos e custos de mão de obra por território (A, B, C) de Mão de Obra 31-05-2024.xlsm."""
    mo_path = os.path.join(pasta_aux, "Mão de Obra 31-05-2024.xlsm")
    if not os.path.exists(mo_path):
        print(f"[AVISO] Mão de obra não encontrada em: {mo_path}")
        return [], []

    import openpyxl
    wb = openpyxl.load_workbook(mo_path, data_only=True, read_only=True)

    contratos = [
        {"contrato": "JA10110552", "empreiteira": "INDICA", "regiao": "Cantagalo / Pádua", "unidade_base": "0010552"},
        {"contrato": "JA10114811", "empreiteira": "ENGELMIG", "regiao": "Itaperuna / Campos", "unidade_base": "0002042"},
        {"contrato": "52L0002297", "empreiteira": "DÍNAMO", "regiao": "Lagos", "unidade_base": "0002297"},
        {"contrato": "JA10083389", "empreiteira": "PROGEN", "regiao": "Macaé", "unidade_base": "0002043"}
    ]

    mapa_abas_detalhadas = {
        "INDICA 0552 (PÁDUA E CANTAGALO)": ("INDICA", "JA10110552"),
        "ENGELMIG ENERGIA 4811(CP E IT)": ("ENGELMIG", "JA10114811"),
        "DINAMO LAGOS 0376": ("DÍNAMO", "52L0002297"),
        "PROGEN (MACAÉ)": ("PROGEN", "JA10083389")
    }

    itens_mo: List[Dict[str, Any]] = []

    for sname, (emp_nome, num_contrato) in mapa_abas_detalhadas.items():
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 9:
                continue
            # Estrutura típica: Contrato, Item, Baremo, Voz, Grupo de instalação, Território, Texto breve, Unidade, Custo unitário, Descrição Completa
            baremo = str(row[2] or "").strip() if row[2] is not None else ""
            voz = str(row[3] or "").strip() if row[3] is not None else ""
            grupo = str(row[4] or "LINHAS AÉREAS").strip()
            terr = str(row[5] or "A").strip().upper()
            texto_brv = str(row[6] or "").strip()
            un = str(row[7] or "UN").strip()
            try:
                custo = float(str(row[8]).replace(",", ".")) if row[8] is not None else 0.0
            except (ValueError, TypeError):
                custo = 0.0

            desc_comp = str(row[9] or "").strip() if len(row) > 9 and row[9] else texto_brv

            if voz and custo > 0:
                itens_mo.append({
                    "contrato": num_contrato,
                    "empreiteira": emp_nome,
                    "baremo": baremo,
                    "voz": voz,
                    "texto_breve": texto_brv,
                    "territorio": terr,
                    "unidade": un,
                    "custo_unitario": custo,
                    "grupo": grupo,
                    "descricao_completa": desc_comp
                })

    wb.close()
    return itens_mo, contratos


def compilar_base_dados(
    caminho_sisorca: Optional[str] = None,
    pasta_aux: Optional[str] = None,
    db_saida: Optional[str] = None,
    json_saida: Optional[str] = None
) -> Dict[str, int]:
    """
    Função principal de compilação: extrai dados das planilhas e salva no SQLite e JSON.
    """
    p_sisorca = caminho_sisorca or DEFAULT_SISORCA_PATH
    p_aux = pasta_aux or DEFAULT_AUX_DIR
    out_db = db_saida or DB_PATH
    out_json = json_saida or JSON_PATH

    print(f"[COMPILAÇÃO] Conectando ao SQLite: {out_db}")
    conn = _init_sqlite(out_db)
    cur = conn.cursor()

    # Limpa dados anteriores
    cur.execute("DELETE FROM materiais;")
    cur.execute("DELETE FROM kits;")
    cur.execute("DELETE FROM mao_de_obra;")
    cur.execute("DELETE FROM contratos;")

    # 1. Materiais
    print(f"[COMPILAÇÃO] Lendo materiais do sisORÇA...")
    materiais = carregar_materiais_sisorca(p_sisorca)
    for m in materiais:
        cur.execute("""
        INSERT INTO materiais (codigo_sie, codigo_e4e, codigo_curto, descricao, unidade, preco_unitario, categoria, tipo, flag_instala, flag_retira, flag_reuse)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            m["codigo_sie"], m["codigo_e4e"], m["codigo_curto"], m["descricao"],
            m["unidade"], m["preco_unitario"], m["categoria"], m["tipo"],
            m["flag_instala"], m["flag_retira"], m["flag_reuse"]
        ))
    print(f" -> {len(materiais)} materiais inseridos.")

    # 2. Kits
    print(f"[COMPILAÇÃO] Lendo kits de estruturas, trafos e aterramentos...")
    kits = carregar_kits_oficiais(p_aux)
    for k in kits:
        cur.execute("""
        INSERT INTO kits (codigo_kit, nome_kit, codigo_material, nome_material, qtd_material, ambiente)
        VALUES (?, ?, ?, ?, ?, ?);
        """, (
            k["codigo_kit"], k["nome_kit"], k["codigo_material"], k["nome_material"],
            k["qtd_material"], k["ambiente"]
        ))
    print(f" -> {len(kits)} itens de kits inseridos.")

    # 3. Mão de Obra e Contratos
    print(f"[COMPILAÇÃO] Lendo mão de obra e contratos das empreiteiras...")
    itens_mo, contratos = carregar_mao_de_obra(p_aux)
    for c in contratos:
        cur.execute("""
        INSERT OR REPLACE INTO contratos (contrato, empreiteira, regiao, unidade_base)
        VALUES (?, ?, ?, ?);
        """, (c["contrato"], c["empreiteira"], c["regiao"], c["unidade_base"]))

    for mo in itens_mo:
        cur.execute("""
        INSERT INTO mao_de_obra (contrato, empreiteira, baremo, voz, texto_breve, territorio, unidade, custo_unitario, grupo, descricao_completa)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            mo["contrato"], mo["empreiteira"], mo["baremo"], mo["voz"],
            mo["texto_breve"], mo["territorio"], mo["unidade"], mo["custo_unitario"],
            mo["grupo"], mo["descricao_completa"]
        ))
    print(f" -> {len(itens_mo)} itens de mão de obra inseridos para {len(contratos)} empreiteiras.")

    conn.commit()
    conn.close()

    # Exporta cache JSON otimizado
    print(f"[COMPILAÇÃO] Gerando cache JSON: {out_json}")
    dados_cache = {
        "materiais": materiais,
        "kits": kits[:1000],  # amostra otimizada ou índices
        "contratos": contratos,
        "total_materiais": len(materiais),
        "total_kits": len(kits),
        "total_mo": len(itens_mo)
    }
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(dados_cache, f, ensure_ascii=False, indent=2)

    print(f"[SUCESSO] Base de dados de orçamento compilada com sucesso!")
    return {
        "materiais": len(materiais),
        "kits": len(kits),
        "mao_de_obra": len(itens_mo),
        "contratos": len(contratos)
    }


if __name__ == "__main__":
    res = compilar_base_dados()
    print("Resultado da compilação:", res)
