"""
Módulo de Exportação Excel para Queda de Tensão (CQT) no Modo Clandestino (Light S.A.).
Gera planilha Excel (.xlsx) profissional de 11 colunas com formatação idêntica
ao padrão operacional da Light:
  1. TRECHO (rótulo do vão: TR, P1 / P1, P2...)
  2. PONTO (sequência numérica)
  3. TRECHO jusante (nó pai / montante: 0 para Trafo, SEQ do montante)
  4. M (comprimento do vão em metros)
  5. CLIENTES (clientes locais)
  6. CLT ACUMULAD (clientes acumulados)
  7. TOTAL DO TRECHO (demanda local em kVA)
  8. ACUMULADA (demanda acumulada em kVA)
  9. LADO (circuito: TRAFO, LADO 1, LADO 2...)
 10. CONDUTOR (cabo adotado)
 11. CQT NO PONTO (tensão calculada com indicador ✔/✖)
"""

import argparse
import json
import os
import sys
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def parse_float_safe(val: Any) -> float:
    """Converte valor para float de forma defensiva (suporta vírgula e strings com unidades)."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace("kVA", "").replace("V", "").replace("%", "").replace("*", "").strip()
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_int_safe(val: Any) -> int:
    """Converte valor para int de forma defensiva."""
    if val is None:
        return 0
    if isinstance(val, int):
        return val
    try:
        return int(float(str(val).strip().replace(",", ".")))
    except ValueError:
        return 0


def formatar_cqt_ponto(v_127: float, qt_perc: float, status_cqt: str = "") -> str:
    """Formata a célula CQT NO PONTO com indicador visual de conformidade (✔/✖)."""
    if "ERRO" in str(status_cqt).upper() or v_127 <= 0.0:
        return "✖ - (ERRO)"
    v_str = f"{v_127:.2f}".replace(".", ",")
    if qt_perc <= 7.0001:
        return f"✔ {v_str}"
    return f"✖ {v_str}"


def exportar_cqt_excel(
    pontos: List[Dict[str, Any]],
    caminho_saida: str,
    abrir_arquivo: bool = False
) -> str:
    """
    Gera pasta de trabalho Excel (.xlsx) com a tabela de 11 colunas de CQT Clandestino.
    
    Args:
        pontos: Lista de dicionários de pontos enriquecidos pelo motor CQT.
        caminho_saida: Caminho do arquivo .xlsx a ser gerado.
        abrir_arquivo: Se True, tenta abrir a planilha no aplicativo padrão.
        
    Returns:
        Caminho absoluto do arquivo gerado.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CQT CLANDESTINO"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de Estilos
    header_fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    
    green_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    blue_fill = PatternFill(start_color="CFE2F3", end_color="CFE2F3", fill_type="solid")
    
    font_normal = Font(name="Arial", size=10, bold=False, color="000000")
    font_trafo = Font(name="Arial", size=10, bold=True, color="C00000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    
    thin_border_side = Side(style="thin", color="B2B2B2")
    thin_border = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=thin_border_side
    )

    # 1. Cabeçalhos das 11 Colunas
    headers = [
        "TRECHO",
        "PONTO",
        "TRECHO\njusante",
        "M",
        "CLIENTES",
        "CLT\nACUMULAD",
        "TOTAL DO\nTRECHO",
        "ACUMULADA",
        "LADO",
        "CONDUTOR",
        "CQT NO\nPONTO",
    ]

    ws.row_dimensions[1].height = 28
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # 2. Preenchimento das Linhas de Dados
    for row_offset, p in enumerate(pontos):
        row_idx = row_offset + 2
        ws.row_dimensions[row_idx].height = 20

        # Identificação se é o nó raiz Trafo
        lado_str = str(p.get("LADO", "")).strip().upper()
        seq_num = parse_int_safe(p.get("SEQ", row_offset + 1))
        eh_trafo = (lado_str == "TRAFO" or seq_num == 1)

        trecho_label = str(p.get("TRECHO", "")).strip()
        if not trecho_label:
            poste_nome = str(p.get("POSTE", "")).strip()
            trecho_label = f"TR, {poste_nome}" if eh_trafo else poste_nome

        montante_seq_raw = p.get("TRECHO_JUSANTE", 0 if eh_trafo else 1)
        montante_seq = parse_int_safe(montante_seq_raw) if str(montante_seq_raw).isdigit() else montante_seq_raw

        comp_val = parse_float_safe(p.get("COMPRIMENTO_VÃO", 4.0 if eh_trafo else 30.0))
        cli_local = parse_int_safe(p.get("CLIENTES", 0))
        cli_acum = parse_int_safe(p.get("CLT_ACUMULADO", p.get("CLIENTES_ACUMULADOS", cli_local)))

        demanda_local = parse_float_safe(p.get("DEMANDA", 0.0))
        demanda_acum = parse_float_safe(p.get("ACUMULADA", 0.0))

        condutor_str = str(p.get("CONDUTOR_PREVISTO", "240 Cu" if eh_trafo else "70 Al - MX")).strip()
        
        # CQT / Tensão
        v_127 = parse_float_safe(p.get("TENSAO_PREVISTA", 127.0))
        qt_perc = parse_float_safe(p.get("CQT_PREVISTO", 0.0))
        status_cqt = str(p.get("STATUS_CQT", ""))
        cqt_formatado = p.get("CQT_NO_PONTO")
        if not cqt_formatado:
            cqt_formatado = formatar_cqt_ponto(v_127, qt_perc, status_cqt)

        # Definição dos valores das 11 colunas
        row_values = [
            trecho_label,                           # 1. TRECHO
            seq_num,                                # 2. PONTO
            montante_seq,                           # 3. TRECHO jusante
            comp_val,                               # 4. M
            cli_local,                              # 5. CLIENTES
            cli_acum,                               # 6. CLT ACUMULAD
            demanda_local,                          # 7. TOTAL DO TRECHO
            demanda_acum,                           # 8. ACUMULADA
            lado_str if lado_str else ("TRAFO" if eh_trafo else "LADO 1"), # 9. LADO
            condutor_str,                           # 10. CONDUTOR
            str(cqt_formatado),                     # 11. CQT NO PONTO
        ]

        # Estilização de cada célula da linha
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            
            # Fonte: Trafo usa destaque vermelho/bold; não-conformidade CQT (✖) usa destaque vermelho; demais normal
            if eh_trafo and col_idx in (5, 6, 7, 8, 9, 10):
                cell.font = font_trafo
            elif col_idx == 11 and "✖" in str(val):
                cell.font = Font(name="Arial", size=10, bold=True, color="C00000")
            else:
                cell.font = font_normal

            # Cores de fundo (Fills) por coluna
            if col_idx in (2, 3, 9, 10):  # PONTO, TRECHO jusante, LADO, CONDUTOR
                cell.fill = green_fill
            elif col_idx in (4, 11):      # M, CQT NO PONTO
                cell.fill = blue_fill

            # Alinhamentos e Formatos numéricos
            if col_idx in (7, 8):  # Demandas kVA
                cell.alignment = align_right
                cell.number_format = "#,##0.00"
            elif col_idx in (4,):   # M
                cell.alignment = align_center
                cell.number_format = "0.0" if isinstance(val, float) and val % 1 != 0 else "0"
            else:
                cell.alignment = align_center

    # 3. Ajuste Automático de Largura das Colunas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 4. Gravação no Disco
    os.makedirs(os.path.dirname(os.path.abspath(caminho_saida)), exist_ok=True)
    wb.save(caminho_saida)

    if abrir_arquivo:
        try:
            if sys.platform.startswith("win"):
                os.startfile(caminho_saida)
        except Exception:
            pass

    return os.path.abspath(caminho_saida)


def main():
    parser = argparse.ArgumentParser(description="Exportador Excel de CQT Clandestino (Light S.A.)")
    parser.add_argument("--json", required=True, help="Caminho do arquivo JSON com a lista de pontos enriquecidos")
    parser.add_argument("--out", required=True, help="Caminho do arquivo Excel (.xlsx) de saída")
    parser.add_argument("--open", action="store_true", help="Abrir a planilha gerada no Excel")
    args = parser.parse_args()

    if not os.path.exists(args.json):
        sys.exit(f"Erro: Arquivo JSON não encontrado: {args.json}")

    with open(args.json, "r", encoding="utf-8") as f:
        pontos = json.load(f)

    saida = exportar_cqt_excel(pontos, args.out, abrir_arquivo=args.open)
    print(f"Sucesso: Planilha CQT gerada em: {saida}")


if __name__ == "__main__":
    main()
