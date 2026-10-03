from datetime import datetime
import json
import os
import sys

_curr_dir = os.path.dirname(os.path.abspath(__file__))
_root_dir = os.path.dirname(_curr_dir)
for _d in [_curr_dir, _root_dir]:
    if _d not in sys.path:
        sys.path.insert(0, _d)

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from typing import Dict, List, Optional
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def format_environmental_text(boxes_txt_path: str) -> str:
    """Lê as camadas do KMZ e gera parágrafo técnico formatado."""
    if not os.path.exists(boxes_txt_path):
        return "Nenhuma área de restrição ambiental interceptada pelo projeto."

    lines = []
    with open(boxes_txt_path, 'r', encoding='utf-8') as f:
        for line in f:
            line_str = line.strip()
            if not line_str:
                continue
            parts = line_str.split('|')
            cat = parts[0].strip()
            spec = parts[1].strip() if len(parts) > 1 else ""
            if spec:
                lines.append(f"• {cat} ({spec})")
            else:
                lines.append(f"• {cat}")

    if not lines:
        return "Nenhuma área de restrição ambiental interceptada pelo projeto."

    texto = (
        "O caminhamento e as estruturas projetadas intersectam as seguintes áreas e restrições "
        "ambientais identificadas no levantamento georreferenciado contra as bases cartográficas oficiais:\n"
        + "\n".join(lines) + "\n\n"
        "Todas as intervenções seguirão rigorosamente as normas ambientais vigentes, as diretrizes do órgão gestor "
        "competente e os procedimentos operacionais de baixo impacto da concessionária."
    )
    return texto


PREFERENCIAS_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "preferencias_ambiental.json")


def carregar_preferencias() -> Dict[str, str]:
    """Carrega o Smart Cache local de preferências do projetista."""
    padrao = {
        "projetista": "FABIO VARGAS",
        "empresa": "HBT",
        "polo": "Angra dos Reis",
        "municipio": "Angra dos Reis",
        "tipo_area": "Área Rural"
    }
    if os.path.exists(PREFERENCIAS_PATH):
        try:
            with open(PREFERENCIAS_PATH, 'r', encoding='utf-8') as f:
                saved = json.load(f)
                padrao.update(saved)
        except Exception:
            pass
    return padrao


def salvar_preferencias(prefs: Dict[str, str]):
    """Salva as preferências no Smart Cache local."""
    try:
        with open(PREFERENCIAS_PATH, 'w', encoding='utf-8') as f:
            json.dump(prefs, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def is_safe_writable_dir(path: Optional[str]) -> bool:
    """Verifica se o diretório existe, não é uma pasta protegida do SO e possui permissão de escrita."""
    if not path or not isinstance(path, str):
        return False
    try:
        clean = path.strip().strip('"').strip("'")
        if not clean:
            return False
        abs_p = os.path.abspath(clean)
        if not os.path.exists(abs_p) or not os.path.isdir(abs_p):
            return False

        p_lower = abs_p.lower()
        win_dir = os.environ.get("WINDIR", r"C:\Windows").lower()
        sys_drive = os.environ.get("SystemDrive", "C:").lower()
        prog_files = os.environ.get("ProgramFiles", r"C:\Program Files").lower()
        prog_files_x86 = os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)").lower()

        if p_lower.startswith(win_dir) or p_lower == sys_drive + "\\" or p_lower == sys_drive:
            return False
        if prog_files and p_lower.startswith(prog_files):
            return False
        if prog_files_x86 and p_lower.startswith(prog_files_x86):
            return False

        test_file = os.path.join(abs_p, f".test_perm_{os.getpid()}.tmp")
        try:
            with open(test_file, 'w', encoding='utf-8') as f:
                f.write("ok")
            os.remove(test_file)
            return True
        except (PermissionError, OSError):
            return False
    except Exception:
        return False


def determinar_pasta_saida(candidate: Optional[str] = None, template_path: Optional[str] = None) -> str:
    """
    Determina a pasta de saída ideal para salvar o Checklist e Memorial.
    Prioridades:
    1. Subpasta de licenciamento ambiental (ex: LIC.AMB, Licen.Ambiental) dentro do candidato ou o próprio candidato.
    2. Diretório informado pelo CAD (DWGPREFIX).
    3. Diretório do template (se seguro e gravável, exceto se for System32/Windows).
    4. Área de Trabalho (Desktop) do usuário.
    5. Pasta Documentos do usuário.
    6. C:\\Temp (garantidamente gravável para CAD/engenharia).
    7. os.getcwd() (somente se seguro e gravável).
    """
    if candidate and is_safe_writable_dir(candidate):
        cand_abs = os.path.abspath(candidate.strip().strip('"').strip("'"))
        cand_name = os.path.basename(cand_abs).lower()
        if any(term in cand_name for term in ["lic.amb", "licen", "ambiental"]):
            return cand_abs

        subpastas = [
            "LIC.AMB", "Licen.Ambiental", "LICEN.AMBIENTAL", "LIC.AMBIENTAL",
            "LICEN.AMB", "Licenciamento Ambiental", "1_Documentos"
        ]
        for sub in subpastas:
            sub_p = os.path.join(cand_abs, sub)
            if os.path.isdir(sub_p) and is_safe_writable_dir(sub_p):
                return sub_p

        tem_manifesto = os.path.exists(os.path.join(cand_abs, ".sob_state.json"))
        tem_dwg = False
        try:
            if os.path.isdir(cand_abs):
                tem_dwg = any(f.lower().endswith(".dwg") for f in os.listdir(cand_abs) if os.path.isfile(os.path.join(cand_abs, f)))
        except OSError:
            pass

        if tem_manifesto or tem_dwg:
            lic_amb = os.path.join(cand_abs, "LIC.AMB")
            try:
                os.makedirs(lic_amb, exist_ok=True)
                if is_safe_writable_dir(lic_amb):
                    return lic_amb
            except Exception:
                pass

        return cand_abs

    if template_path:
        clean_tpl = template_path.strip().strip('"').strip("'")
        if clean_tpl:
            tpl_dir = os.path.dirname(os.path.abspath(clean_tpl))
            if is_safe_writable_dir(tpl_dir):
                return tpl_dir

    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    if is_safe_writable_dir(desktop):
        return desktop

    docs = os.path.join(os.path.expanduser("~"), "Documents")
    if is_safe_writable_dir(docs):
        return docs

    c_temp = r"C:\Temp"
    if os.path.exists(c_temp) and is_safe_writable_dir(c_temp):
        return c_temp

    cwd = os.getcwd()
    if is_safe_writable_dir(cwd):
        return os.path.abspath(cwd)

    return os.path.abspath(os.path.expanduser("~"))


def apply_formatted_text(p, new_text: str, default_font: str = "Arial", default_size_pt: float = 11.0):
    """
    Substitui o texto de um parágrafo preservando rigorosamente
    todas as configurações de estilo, alinhamento, espaçamento (antes/depois),
    entrelinha, recuo e fonte originais.
    """
    orig_style = p.style
    orig_align = p.alignment
    orig_line_spacing = p.paragraph_format.line_spacing
    orig_space_after = p.paragraph_format.space_after
    orig_space_before = p.paragraph_format.space_before
    orig_left_indent = p.paragraph_format.left_indent
    orig_first_line_indent = p.paragraph_format.first_line_indent

    # Tenta herdar fonte e tamanho do primeiro run original se existirem
    font_name = default_font
    font_size = Pt(default_size_pt)
    if p.runs and p.runs[0].font.name:
        font_name = p.runs[0].font.name
    if p.runs and p.runs[0].font.size:
        font_size = p.runs[0].font.size

    # Remove completamente runs antigos mantendo a estrutura do parágrafo
    for old_r in list(p.runs):
        p._element.remove(old_r._element)

    # Trata quebras de linha no texto gerado
    lines = new_text.split('\n')
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = font_name
        r._element.get_or_add_rPr().get_or_add_rFonts().set(docx.oxml.ns.qn('w:ascii'), font_name)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(docx.oxml.ns.qn('w:hAnsi'), font_name)
        r.font.size = font_size
        r.font.highlight_color = None
        r.font.color.rgb = RGBColor(0, 0, 0)
        if i < len(lines) - 1:
            r.add_break()

    # Restaura rigorosamente todas as propriedades do parágrafo
    p.style = orig_style
    p.alignment = orig_align
    p.paragraph_format.line_spacing = orig_line_spacing
    p.paragraph_format.space_after = orig_space_after
    p.paragraph_format.space_before = orig_space_before
    p.paragraph_format.left_indent = orig_left_indent
    p.paragraph_format.first_line_indent = orig_first_line_indent


def generate_docx_memorial(template_path: str, output_path: str, data: Dict[str, str]):
    """Gera o Memorial Descritivo com formatação fiel e rodapé atualizado."""
    doc = docx.Document(template_path)

    # 1. Monta o texto específico do caminhamento atendendo à divergência MT e BT
    has_mt = data.get("EXTENSAO_MT", "0") not in ["0", ""]
    has_bt = data.get("EXTENSAO_BT", "0") not in ["0", ""]

    if has_mt and has_bt:
        caminhamento_text = (
            f"O circuito de rede de distribuição de rede elétrica derivará da rede existente na {data['ENDERECO_DERIVACAO']}, "
            f"será instalado {data['EXTENSAO_MT']} metros do condutor de Media tensão {data['TENSAO_MT']}volts entre o poste "
            f"{data['POSTE_INI_MT']} de coordenadas UTM X ={data['POSTE_INI_MT_X']} e Y ={data['POSTE_INI_MT_Y']} e o poste "
            f"{data['POSTE_FIM_MT']} de coordenadas UTM X-{data['POSTE_FIM_MT_X']} e Y-{data['POSTE_FIM_MT_Y']} e "
            f"{data['EXTENSAO_BT']} metros de rede em baixa tensão 127/220V entre o poste "
            f"{data['POSTE_INI_BT']} de coordenadas X-{data['POSTE_INI_BT_X']} Y-{data['POSTE_INI_BT_Y']} e o poste "
            f"{data['POSTE_FIM_BT']} de coordenadas X-{data['POSTE_FIM_BT_X']} Y-{data['POSTE_FIM_BT_Y']} para atendimento a um cliente {data['TIPO_CLIENTE']}."
        )
    elif has_mt and not has_bt:
        caminhamento_text = (
            f"O circuito de rede de distribuição de rede elétrica derivará da rede existente na {data['ENDERECO_DERIVACAO']}, "
            f"será instalado {data['EXTENSAO_MT']} metros do condutor de Media tensão {data['TENSAO_MT']}volts entre o poste "
            f"{data['POSTE_INI_MT']} de coordenadas UTM X ={data['POSTE_INI_MT_X']} e Y ={data['POSTE_INI_MT_Y']} e o poste "
            f"{data['POSTE_FIM_MT']} de coordenadas UTM X-{data['POSTE_FIM_MT_X']} e Y-{data['POSTE_FIM_MT_Y']} para atendimento a um cliente {data['TIPO_CLIENTE']}."
        )
    else:
        caminhamento_text = (
            f"O circuito de rede de distribuição de rede elétrica derivará da rede existente na {data['ENDERECO_DERIVACAO']}, "
            f"será instalado {data['EXTENSAO_BT']} metros de rede em baixa tensão 127/220V entre o poste "
            f"{data['POSTE_INI_BT']} de coordenadas X-{data['POSTE_INI_BT_X']} Y-{data['POSTE_INI_BT_Y']} e o poste "
            f"{data['POSTE_FIM_BT']} de coordenadas X-{data['POSTE_FIM_BT_X']} Y-{data['POSTE_FIM_BT_Y']} para atendimento a um cliente {data['TIPO_CLIENTE']}."
        )

    data["TEXTO_CAMINHAMENTO"] = caminhamento_text

    # 2. Percorre os parágrafos do documento e substitui tags com formatação preservada
    for p in doc.paragraphs:
        # Se for o parágrafo de caminhamento
        if "{{TEXTO_CAMINHAMENTO}}" in p.text or "{{POSTE_INI_MT}}" in p.text or "O circuito de rede de distribuição" in p.text:
            apply_formatted_text(p, caminhamento_text, default_font="Arial", default_size_pt=11.0)
            continue

        # Verifica se alguma tag está presente no parágrafo
        needs_update = False
        p_text = p.text
        for key, val in data.items():
            tag = f"{{{{{key}}}}}"
            if tag in p_text:
                p_text = p_text.replace(tag, str(val))
                needs_update = True

        if needs_update:
            apply_formatted_text(p, p_text, default_font="Arial", default_size_pt=11.0)

    # 3. Atualiza Polo e Ano no Rodapé e remove realces em amarelo
    polo_val = data.get("POLO", "Polo Regional")
    ano_val = str(data.get("ANO", datetime.now().year))

    for s in doc.sections:
        footer_elem = s.footer._element
        # Substitui texto de Polo e Ano nos elementos <w:t> do rodapé
        for t in footer_elem.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
            if t.text:
                if 'xxxxx' in t.text:
                    t.text = t.text.replace('xxxxx', polo_val)
                elif '{{POLO}}' in t.text:
                    t.text = t.text.replace('{{POLO}}', polo_val)
                if '20xx' in t.text:
                    t.text = t.text.replace('20xx', ano_val)
                elif '{{ANO}}' in t.text:
                    t.text = t.text.replace('{{ANO}}', ano_val)

        # Remove destaque amarelo em todos os runs do rodapé
        for rPr in footer_elem.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr'):
            h = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight')
            if h is not None:
                rPr.remove(h)

    doc.save(output_path)
    return output_path


class MemorialGUI:
    def __init__(self, template_path: str, pts_path: str, boxes_path: str, fuso: int = 23, candidate_out_dir: Optional[str] = None):
        # Resolução resiliente de template (.docx)
        resolved_template = template_path
        if not resolved_template or not os.path.exists(resolved_template):
            candidatos_tpl = [
                r"G:\Meu Drive\1-HBT\04_MODELOS\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
                r"G:\Meu Drive\1-HBT\05_PROJETOS_HISTORICOS\SETEMBRO\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
                os.path.join(os.path.dirname(os.path.abspath(__file__)), "modelos", "MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx"),
                r"G:\Meu Drive\1-HBT\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
            ]
            for t_cand in candidatos_tpl:
                if os.path.exists(t_cand):
                    resolved_template = t_cand
                    break

        self.template_path = resolved_template
        self.pts_path = pts_path
        self.boxes_path = boxes_path
        self.fuso = fuso
        self.candidate_out_dir = candidate_out_dir
        self.default_out_dir = determinar_pasta_saida(candidate_out_dir, resolved_template)

        # Extrai postes e coordenadas
        self.pts_list = []
        self.postes_labels = []

        if os.path.exists(pts_path):
            try:
                with open(pts_path, 'r', encoding='utf-8') as f:
                    self.pts_list = json.load(f)
                    for pt in self.pts_list:
                        lbl = f"{pt.get('id', 'Poste')} (X: {float(pt['x']):.1f}, Y: {float(pt['y']):.1f})"
                        self.postes_labels.append(lbl)
            except Exception:
                pass

        # Análise Inteligente de Postes via blocos do CAD
        from postes_detalhador import analisar_postes_projeto
        self.analise_postes = analisar_postes_projeto(self.pts_list)
        self.total_postes = self.analise_postes["total_postes"]
        self.detalhe_postes_txt = self.analise_postes["texto_formatado"]

        self.areas_ambientais_text = format_environmental_text(boxes_path)

        # Analisa restrição ambiental via KMZ
        from checklist_ambiental_manager import analisar_incidencias_kmz
        self.analise_ambiental = analisar_incidencias_kmz(boxes_path)
        self.tem_restricao = (
            self.analise_ambiental["em_app"]
            or self.analise_ambiental["em_uc"]
            or self.analise_ambiental["em_exclusao"]
        )

        # Janela Tkinter Moderna
        self.root = tk.Tk()
        self.root.title("AUX CAD - Memorial Técnico Descritivo Ambiental (ENEL)")
        self.root.geometry("1120x780")
        self.root.minsize(960, 640)
        self.root.configure(bg="#F8FAFC")

        try:
            self.root.deiconify()
            self.root.lift()
            self.root.attributes("-topmost", True)
            self.root.after_idle(self.root.attributes, "-topmost", False)
            self.root.focus_force()
        except Exception:
            pass

        style = ttk.Style()
        style.theme_use('clam')
        style.configure("TEntry", fieldbackground="#FFFFFF", foreground="#1E293B")
        style.configure("Vertical.TScrollbar", background="#CBD5E1", bordercolor="#E2E8F0")

        self._build_interface()

    def _build_interface(self):
        # 1. HEADER MODERNO
        header = tk.Frame(self.root, bg="#1B365D", height=68)
        header.pack(fill="x", side="top")
        header.pack_propagate(False)

        header_inner = tk.Frame(header, bg="#1B365D", padx=20)
        header_inner.pack(fill="both", expand=True)

        lbl_tit = tk.Label(
            header_inner,
            text="🌿 MEMORIAL TÉCNICO DESCRITIVO AMBIENTAL",
            font=("Segoe UI", 13, "bold"),
            fg="#FFFFFF",
            bg="#1B365D"
        )
        lbl_tit.pack(anchor="w", pady=(10, 0))

        lbl_sub = tk.Label(
            header_inner,
            text="Diagnóstico Automatizado de Restrições Cartográficas (KMZ) & Geração Técnica de Documentos",
            font=("Segoe UI", 9),
            fg="#94A3B8",
            bg="#1B365D"
        )
        lbl_sub.pack(anchor="w")

        # 2. CONTAINER COM CANVAS RESPONSIVO E SCROLLBAR
        container = tk.Frame(self.root, bg="#F8FAFC")
        container.pack(fill="both", expand=True)

        canvas = tk.Canvas(container, bg="#F8FAFC", highlightthickness=0, borderwidth=0)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)

        self.scrollable_frame = tk.Frame(canvas, bg="#F8FAFC", padx=16, pady=12)
        canvas_window = canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")

        def on_canvas_configure(event):
            canvas.itemconfig(canvas_window, width=event.width)
            canvas.configure(scrollregion=canvas.bbox("all"))

        canvas.bind("<Configure>", on_canvas_configure)
        self.scrollable_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.configure(yscrollcommand=scrollbar.set)

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
        canvas.bind_all("<MouseWheel>", _on_mousewheel)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Grid de 2 Colunas com peso idêntico
        self.scrollable_frame.columnconfigure(0, weight=1, uniform="group1")
        self.scrollable_frame.columnconfigure(1, weight=1, uniform="group1")

        # Smart Cache de Preferências
        prefs = carregar_preferencias()

        # =========================================================
        # COLUNA 0 - ESQUERDA: DADOS GERAIS E TRAÇADO DE REDE
        # =========================================================

        # CARD 1: Identificação do Projeto e Rodapé
        grp_proj = self._criar_card("📋 1. Identificação da Obra & Projeto", 0, 0)
        self.ent_projeto = self._add_field(grp_proj, "Número do Projeto (Ordem):", "A0", 0)
        self.ent_local = self._add_field(grp_proj, "Local da Obra (Endereço/Fazenda):", "Ilha Grande, Angra dos Reis - RJ", 1)
        self.ent_municipio = self._add_field(grp_proj, "Município (Checklist):", prefs.get("municipio", "Angra dos Reis"), 2)
        self.ent_projetista = self._add_field(grp_proj, "Projetista Responsável:", prefs.get("projetista", "FABIO VARGAS"), 3)
        self.ent_empresa = self._add_field(grp_proj, "Empresa Executante:", prefs.get("empresa", "HBT"), 4)
        self.ent_derivacao = self._add_field(grp_proj, "Endereço da Rede Existente (Derivação):", "Avenida Principal Nº 100", 5)
        self.ent_polo = self._add_field(grp_proj, "Polo Operativo (Rodapé):", prefs.get("polo", "Angra dos Reis"), 6)
        self.ent_ano = self._add_field(grp_proj, "Ano (Rodapé):", str(datetime.now().year), 7)
        self.ent_tipo_area = self._add_field(grp_proj, "Classificação da Área:", prefs.get("tipo_area", "Área Rural"), 8)
        self.ent_out_dir = self._add_dir_field(grp_proj, "Pasta de Destino (Salvar):", self.default_out_dir, 9)

        # CARD 2: Traçado da Rede Elétrica (MT e BT)
        grp_rede = self._criar_card("⚡ 2. Caminhamento da Rede Elétrica (MT / BT)", 1, 0)

        lbl_sec_mt = tk.Label(grp_rede, text="▸ Média Tensão (MT)", font=("Segoe UI", 9, "bold"), fg="#1E40AF", bg="#FFFFFF")
        lbl_sec_mt.grid(row=0, column=0, columnspan=2, sticky="w", pady=(2, 4), padx=6)

        self.ent_ext_mt = self._add_field(grp_rede, "Extensão MT (metros):", "291", 1)
        self.ent_tensao_mt = self._add_field(grp_rede, "Tensão MT (Volts):", "13.800", 2)
        self.ent_cabo_mt = self._add_field(grp_rede, "Condutor MT:", "2#50AAAC-13.8Kv engraxado", 3)

        ini_p = self.pts_list[0] if self.pts_list else {}
        self.ent_p_ini_mt = self._add_field(grp_rede, "Poste Inicial MT:", ini_p.get('id', 'Poste 01'), 4)
        self.ent_p_ini_mt_x = self._add_field(grp_rede, "UTM X Inicial MT:", f"{float(ini_p.get('x', 0)):.3f}" if ini_p else "000000", 5)
        self.ent_p_ini_mt_y = self._add_field(grp_rede, "UTM Y Inicial MT:", f"{float(ini_p.get('y', 0)):.3f}" if ini_p else "0000000", 6)

        fim_p = self.pts_list[min(7, len(self.pts_list)-1)] if self.pts_list else {}
        self.ent_p_fim_mt = self._add_field(grp_rede, "Poste Final MT:", fim_p.get('id', 'Poste 08'), 7)
        self.ent_p_fim_mt_x = self._add_field(grp_rede, "UTM X Final MT:", f"{float(fim_p.get('x', 0)):.3f}" if fim_p else "000000", 8)
        self.ent_p_fim_mt_y = self._add_field(grp_rede, "UTM Y Final MT:", f"{float(fim_p.get('y', 0)):.3f}" if fim_p else "0000000", 9)

        lbl_sec_bt = tk.Label(grp_rede, text="▸ Baixa Tensão (BT) - Traçado Específico", font=("Segoe UI", 9, "bold"), fg="#0D9488", bg="#FFFFFF")
        lbl_sec_bt.grid(row=10, column=0, columnspan=2, sticky="w", pady=(8, 4), padx=6)

        self.ent_ext_bt = self._add_field(grp_rede, "Extensão BT (metros):", "72", 11)
        self.ent_cabo_bt = self._add_field(grp_rede, "Condutor BT:", "PR3#35(54,6)AL", 12)

        ini_bt_p = fim_p if fim_p else (self.pts_list[0] if self.pts_list else {})
        self.ent_p_ini_bt = self._add_field(grp_rede, "Poste Inicial BT (Derivação/Trafo):", ini_bt_p.get('id', 'Poste 08'), 13)
        self.ent_p_ini_bt_x = self._add_field(grp_rede, "UTM X Inicial BT:", f"{float(ini_bt_p.get('x', 0)):.3f}" if ini_bt_p else "000000", 14)
        self.ent_p_ini_bt_y = self._add_field(grp_rede, "UTM Y Inicial BT:", f"{float(ini_bt_p.get('y', 0)):.3f}" if ini_bt_p else "0000000", 15)

        fim_bt_p = self.pts_list[-1] if self.pts_list else {}
        self.ent_p_fim_bt = self._add_field(grp_rede, "Poste Final BT:", fim_bt_p.get('id', 'Poste 11'), 16)
        self.ent_p_fim_bt_x = self._add_field(grp_rede, "UTM X Final BT:", f"{float(fim_bt_p.get('x', 0)):.3f}" if fim_bt_p else "000000", 17)
        self.ent_p_fim_bt_y = self._add_field(grp_rede, "UTM Y Final BT:", f"{float(fim_bt_p.get('y', 0)):.3f}" if fim_bt_p else "0000000", 18)

        # =========================================================
        # COLUNA 1 - DIREITA: EQUIPAMENTOS E ASPECTOS AMBIENTAIS
        # =========================================================

        # CARD 3: Equipamentos e Postes
        grp_eq = self._criar_card("🏗️ 3. Transformadores e Postes", 0, 1)
        self.ent_trafo = self._add_field(grp_eq, "Transformador:", "1-10kVa -saída a 3 fios.", 0)
        self.ent_qtd_postes = self._add_field(grp_eq, "Quantidade Total de Postes:", str(self.total_postes if self.total_postes > 0 else 10), 1)
        self.ent_detalhe_postes = self._add_field(grp_eq, "Detalhamento dos Postes:", self.detalhe_postes_txt, 2)

        # CARD 4: Meio Ambiente, Vegetação e KMZ
        grp_amb = self._criar_card("🌳 4. Diagnóstico Ambiental & Vegetação", 1, 1)

        # Banner de Alerta / Status
        if self.tem_restricao:
            alert_card = tk.Frame(grp_amb, bg="#FEF2F2", relief="solid", bd=1, padx=12, pady=10)
            alert_card.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10), padx=6)

            tk.Label(
                alert_card,
                text="⚠ ÁREA DE RESTRIÇÃO AMBIENTAL DETECTADA NO KMZ!",
                font=("Segoe UI", 10, "bold"),
                fg="#991B1B",
                bg="#FEF2F2"
            ).pack(anchor="w")

            tk.Label(
                alert_card,
                text="Serão gerados obrigatoriamente:\n• Memorial Descritivo (.docx) com detalhamento da intervenção\n• Checklist Ambiental (.xlsx) com apontamento das restrições e APP/UC.",
                font=("Segoe UI", 9),
                fg="#7F1D1D",
                bg="#FEF2F2",
                justify="left"
            ).pack(anchor="w", pady=(4, 0))
        else:
            alert_card = tk.Frame(grp_amb, bg="#F0FDF4", relief="solid", bd=1, padx=12, pady=10)
            alert_card.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10), padx=6)

            tk.Label(
                alert_card,
                text="✔ PROJETO FORA DE ÁREAS AMBIENTAIS RESTRITAS!",
                font=("Segoe UI", 10, "bold"),
                fg="#166534",
                bg="#F0FDF4"
            ).pack(anchor="w")

            tk.Label(
                alert_card,
                text="Conforme procedimento operacional da concessionária:\n• O Memorial Descritivo é opcional quando não há restrições ambientais.\n• O Checklist (.xlsx) será emitido atestando ausência de restrições.",
                font=("Segoe UI", 9),
                fg="#14532D",
                bg="#F0FDF4",
                justify="left"
            ).pack(anchor="w", pady=(4, 0))

            self.var_gerar_memorial = tk.BooleanVar(value=True)
            chk_mem = tk.Checkbutton(
                alert_card,
                text="Gerar também o Memorial Descritivo (.docx)",
                variable=self.var_gerar_memorial,
                font=("Segoe UI", 9, "bold"),
                fg="#14532D",
                bg="#F0FDF4",
                activebackground="#F0FDF4",
                cursor="hand2"
            )
            chk_mem.pack(anchor="w", pady=(6, 0))

        lbl_kmz = tk.Label(grp_amb, text="Áreas e Restrições Detectadas pelo KMZ:", font=("Segoe UI", 9, "bold"), fg="#1E293B", bg="#FFFFFF")
        lbl_kmz.grid(row=1, column=0, columnspan=2, sticky="w", pady=(4, 2), padx=6)

        txt_frame = tk.Frame(grp_amb, bg="#FFFFFF", bd=1, relief="solid")
        txt_frame.grid(row=2, column=0, columnspan=2, sticky="ew", pady=4, padx=6)
        grp_amb.columnconfigure(1, weight=1)

        self.txt_areas = tk.Text(txt_frame, height=5, font=("Segoe UI", 9), bg="#FFFFFF", fg="#1E293B", relief="flat", padx=6, pady=6)
        txt_scroll = ttk.Scrollbar(txt_frame, orient="vertical", command=self.txt_areas.yview)
        self.txt_areas.configure(yscrollcommand=txt_scroll.set)
        self.txt_areas.pack(side="left", fill="both", expand=True)
        txt_scroll.pack(side="right", fill="y")
        self.txt_areas.insert("1.0", self.areas_ambientais_text)

        # Campos de Intervenção Arbórea
        self.ent_podas = self._add_field(grp_amb, "Quantidade de Podas Previstas:", "0", 3)
        self.ent_supressoes = self._add_field(grp_amb, "Quantidade de Supressões Arbóreas:", "0", 4)

        # =========================================================
        # 3. BARRA INFERIOR FIXA COM BOTÕES DE AÇÃO
        # =========================================================
        bottom_bar = tk.Frame(self.root, bg="#FFFFFF", height=60, bd=1, relief="solid", padx=20)
        bottom_bar.pack(fill="x", side="bottom")
        bottom_bar.pack_propagate(False)

        btn_cancelar = tk.Button(
            bottom_bar,
            text="Cancelar",
            font=("Segoe UI", 9),
            bg="#F1F5F9",
            fg="#475569",
            activebackground="#E2E8F0",
            relief="flat",
            bd=0,
            padx=14,
            pady=7,
            cursor="hand2",
            command=self.root.destroy
        )
        btn_cancelar.pack(side="left", pady=12)

        btn_text = "✔ GERAR MEMORIAL DESCRITIVO E CHECKLIST" if self.tem_restricao else "✔ GERAR CHECKLIST AMBIENTAL (.XLSX)"
        btn_bg = "#1D4ED8" if self.tem_restricao else "#15803D"
        btn_active = "#1E40AF" if self.tem_restricao else "#166534"

        btn_gerar = tk.Button(
            bottom_bar,
            text=btn_text,
            font=("Segoe UI", 10, "bold"),
            bg=btn_bg,
            fg="white",
            activebackground=btn_active,
            activeforeground="white",
            relief="flat",
            bd=0,
            padx=20,
            pady=8,
            cursor="hand2",
            command=self._on_gerar
        )
        btn_gerar.pack(side="right", pady=12)

    def _criar_card(self, titulo: str, row: int, col: int) -> tk.LabelFrame:
        card = tk.LabelFrame(
            self.scrollable_frame,
            text=f" {titulo} ",
            font=("Segoe UI", 10, "bold"),
            fg="#1E3A8A",
            bg="#FFFFFF",
            bd=1,
            relief="solid",
            padx=12,
            pady=10
        )
        card.grid(row=row, column=col, sticky="nsew", padx=8, pady=8)
        card.columnconfigure(1, weight=1)
        return card

    def _add_field(self, parent, label_text: str, default_val: str, row: int) -> tk.Entry:
        lbl = tk.Label(parent, text=label_text, font=("Segoe UI", 9), fg="#334155", bg="#FFFFFF")
        lbl.grid(row=row, column=0, sticky="w", pady=3, padx=6)
        ent = ttk.Entry(parent, font=("Segoe UI", 9))
        ent.grid(row=row, column=1, sticky="ew", pady=3, padx=6)
        ent.insert(0, default_val)
        return ent

    def _add_dir_field(self, parent, label_text: str, default_val: str, row: int) -> ttk.Entry:
        lbl = tk.Label(parent, text=label_text, font=("Segoe UI", 9, "bold"), fg="#1E293B", bg="#FFFFFF")
        lbl.grid(row=row, column=0, sticky="w", pady=3, padx=6)

        f_dir = tk.Frame(parent, bg="#FFFFFF")
        f_dir.grid(row=row, column=1, sticky="ew", pady=3, padx=6)
        f_dir.columnconfigure(0, weight=1)

        ent = ttk.Entry(f_dir, font=("Segoe UI", 9))
        ent.grid(row=0, column=0, sticky="ew")
        ent.insert(0, default_val)

        btn = tk.Button(
            f_dir,
            text="Procurar...",
            font=("Segoe UI", 8),
            bg="#E2E8F0",
            fg="#1E293B",
            relief="groove",
            bd=1,
            padx=6,
            pady=1,
            cursor="hand2",
            command=lambda: self._on_browse_dir(ent)
        )
        btn.grid(row=0, column=1, padx=(4, 0))
        return ent

    def _on_browse_dir(self, entry: ttk.Entry):
        cur = entry.get().strip()
        initial = cur if os.path.isdir(cur) else os.path.expanduser("~")
        chosen = filedialog.askdirectory(parent=self.root, title="Selecionar Pasta de Destino", initialdir=initial)
        if chosen:
            entry.delete(0, tk.END)
            entry.insert(0, os.path.normpath(chosen))

    def _on_gerar(self):
        num_proj = self.ent_projeto.get().strip()
        if not num_proj:
            messagebox.showwarning("Aviso", "Por favor, informe o número do projeto.")
            return

        dados = {
            "NUMERO_PROJETO": num_proj,
            "LOCAL_OBRA": self.ent_local.get().strip(),
            "MUNICIPIO": self.ent_municipio.get().strip(),
            "PROJETISTA_RESPONSAVEL": self.ent_projetista.get().strip(),
            "EMPRESA": self.ent_empresa.get().strip(),
            "ENDERECO_DERIVACAO": self.ent_derivacao.get().strip(),
            "POLO": self.ent_polo.get().strip(),
            "ANO": self.ent_ano.get().strip(),
            "TIPO_AREA": self.ent_tipo_area.get().strip(),
            "TIPO_CLIENTE": self.ent_tipo_area.get().strip().lower(),
            "EXTENSAO_MT": self.ent_ext_mt.get().strip(),
            "TENSAO_MT": self.ent_tensao_mt.get().strip(),
            "CABO_MT": self.ent_cabo_mt.get().strip(),
            "POSTE_INI_MT": self.ent_p_ini_mt.get().strip(),
            "POSTE_INI_MT_X": self.ent_p_ini_mt_x.get().strip(),
            "POSTE_INI_MT_Y": self.ent_p_ini_mt_y.get().strip(),
            "POSTE_FIM_MT": self.ent_p_fim_mt.get().strip(),
            "POSTE_FIM_MT_X": self.ent_p_fim_mt_x.get().strip(),
            "POSTE_FIM_MT_Y": self.ent_p_fim_mt_y.get().strip(),
            "EXTENSAO_BT": self.ent_ext_bt.get().strip(),
            "CABO_BT": self.ent_cabo_bt.get().strip(),
            "POSTE_INI_BT": self.ent_p_ini_bt.get().strip(),
            "POSTE_INI_BT_X": self.ent_p_ini_bt_x.get().strip(),
            "POSTE_INI_BT_Y": self.ent_p_ini_bt_y.get().strip(),
            "POSTE_FIM_BT": self.ent_p_fim_bt.get().strip(),
            "POSTE_FIM_BT_X": self.ent_p_fim_bt_x.get().strip(),
            "POSTE_FIM_BT_Y": self.ent_p_fim_bt_y.get().strip(),
            "TRANSFORMADOR": self.ent_trafo.get().strip(),
            "QTD_POSTES_TOTAL": self.ent_qtd_postes.get().strip(),
            "DETALHAMENTO_POSTES": self.ent_detalhe_postes.get().strip(),
            "QTD_PODAS": self.ent_podas.get().strip(),
            "QTD_SUPRESSOES": self.ent_supressoes.get().strip(),
            "AREAS_AMBIENTAIS": self.txt_areas.get("1.0", "end-1c").strip()
        }

        # Salva o Smart Cache com as preferências digitadas
        salvar_preferencias({
            "projetista": dados["PROJETISTA_RESPONSAVEL"],
            "empresa": dados["EMPRESA"],
            "polo": dados["POLO"],
            "municipio": dados["MUNICIPIO"],
            "tipo_area": dados["TIPO_AREA"]
        })

        safe_num = num_proj.replace('/', '_').replace('\\', '_')
        out_name = f"MEMORIAL_DESCRITIVO_{safe_num}.docx"

        # Resolução e validação robusta da pasta de destino
        out_dir = self.ent_out_dir.get().strip().strip('"').strip("'")
        if not out_dir:
            out_dir = self.default_out_dir

        if not os.path.exists(out_dir):
            try:
                os.makedirs(out_dir, exist_ok=True)
            except Exception:
                pass

        if not is_safe_writable_dir(out_dir):
            fallback_dir = determinar_pasta_saida(None, None)
            messagebox.showwarning(
                "Pasta Sem Permissão de Escrita",
                f"A pasta selecionada:\n{out_dir}\n\n"
                f"não possui permissão de escrita ou é protegida pelo sistema operacional.\n"
                f"Os arquivos serão salvos na pasta alternativa:\n{fallback_dir}"
            )
            out_dir = fallback_dir
            self.ent_out_dir.delete(0, tk.END)
            self.ent_out_dir.insert(0, out_dir)

        out_path = os.path.join(out_dir, out_name)

        # 1. Gera o Croqui / Ortofoto de Satélite Georreferenciado
        croqui_out_name = f"CROQUI_AMBIENTAL_{safe_num}.png"
        croqui_out_path = os.path.join(out_dir, croqui_out_name)
        croqui_ok = False
        try:
            if self.pts_list:
                from croqui_ambiental_generator import gerar_croqui_satelite
                gerar_croqui_satelite(
                    pontos_utm=self.pts_list,
                    fuso_utm=self.fuso,
                    output_path=croqui_out_path,
                    dados_projeto={"SOB": num_proj, "MUNICIPIO": dados["MUNICIPIO"], "PROJETISTA": dados["PROJETISTA_RESPONSAVEL"]}
                )
                croqui_ok = True
        except Exception as e_croqui:
            print(f"[Aviso] Falha ao gerar croqui de satélite: {e_croqui}")

        # 2. Gera sempre o Checklist Ambiental em Excel no mesmo diretório
        checklist_out_name = f"CHECKLIST_{safe_num}.xlsx"
        checklist_out_path = os.path.join(out_dir, checklist_out_name)
        checklist_ok = False

        try:
            from checklist_ambiental_manager import preencher_checklist_ambiental, TEMPLATE_DEFAULT, TEMPLATE_FALLBACK
            chk_tpl = TEMPLATE_DEFAULT if os.path.exists(TEMPLATE_DEFAULT) else TEMPLATE_FALLBACK
            if os.path.exists(chk_tpl):
                try:
                    preencher_checklist_ambiental(chk_tpl, checklist_out_path, dados, boxes_txt_path=self.boxes_path)
                except PermissionError:
                    ts = datetime.now().strftime("%H%M%S")
                    checklist_out_name = f"CHECKLIST_{safe_num}_{ts}.xlsx"
                    checklist_out_path = os.path.join(out_dir, checklist_out_name)
                    preencher_checklist_ambiental(chk_tpl, checklist_out_path, dados, boxes_txt_path=self.boxes_path)
                    messagebox.showinfo(
                        "Aviso - Arquivo Aberto no Excel",
                        f"O checklist anterior estava em uso no Excel.\n"
                        f"Foi gerada uma nova cópia como:\n{checklist_out_name}"
                    )
                checklist_ok = True
        except Exception as e_chk:
            messagebox.showerror("Erro", f"Falha ao gerar o Checklist Ambiental: {e_chk}")
            return

        # 3. Regra de Negócio: Se NÃO houver restrição ambiental e usuário desmarcou o Memorial, dispensa
        deve_gerar_memorial = self.tem_restricao or (hasattr(self, "var_gerar_memorial") and self.var_gerar_memorial.get())
        if not deve_gerar_memorial:
            msg_aviso = (
                f"PROJETO LIVRE DE RESTRIÇÕES AMBIENTAIS!\n\n"
                f"O caminhamento NÃO incide em nenhuma área de restrição ambiental cartográfica.\n"
                f"Conforme o procedimento técnico da concessionária:\n\n"
                f"✔ O CHECKLIST AMBIENTAL foi gerado atestando ausência de restrições:\n"
                f"{checklist_out_path}\n\n"
            )
            if croqui_ok:
                msg_aviso += f"✔ A ORTOFOTO / CROQUI DE SATÉLITE foi gerada:\n{croqui_out_path}\n\n"
            msg_aviso += (
                f"ℹ O MEMORIAL DESCRITIVO AMBIENTAL foi dispensado conforme seleção do usuário.\n\n"
                f"Os arquivos gerados estão prontos para conferência."
            )
            messagebox.showinfo("Checklist Preenchido - Sem Restrições", msg_aviso)
            self.root.destroy()
            if checklist_ok and os.path.exists(checklist_out_path):
                try:
                    os.startfile(checklist_out_path)
                except Exception:
                    pass
            return

        # 4. Caso contrário, gera o Memorial Descritivo
        try:
            if not self.template_path or not os.path.exists(self.template_path):
                candidatos_tpl = [
                    r"G:\Meu Drive\1-HBT\04_MODELOS\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
                    r"G:\Meu Drive\1-HBT\05_PROJETOS_HISTORICOS\SETEMBRO\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
                    os.path.join(os.path.dirname(os.path.abspath(__file__)), "modelos", "MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx"),
                    r"G:\Meu Drive\1-HBT\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
                ]
                for t_cand in candidatos_tpl:
                    if os.path.exists(t_cand):
                        self.template_path = t_cand
                        break

            if not self.template_path or not os.path.exists(self.template_path):
                messagebox.showwarning(
                    "Aviso - Modelo Word Não Localizado",
                    f"O modelo do Memorial Descritivo (.docx) não foi localizado em:\n{self.template_path}\n\n"
                    f"✔ O Checklist Ambiental (.xlsx) foi gerado com sucesso em:\n{checklist_out_path}\n\n"
                    "Para gerar também o arquivo Word (.docx), certifique-se de que o Google Drive ou modelo esteja acessível."
                )
                self.root.destroy()
                if checklist_ok and os.path.exists(checklist_out_path):
                    try:
                        os.startfile(checklist_out_path)
                    except Exception:
                        pass
                return
            try:
                generate_docx_memorial(self.template_path, out_path, dados)
            except PermissionError:
                ts = datetime.now().strftime("%H%M%S")
                out_name = f"MEMORIAL_DESCRITIVO_{safe_num}_{ts}.docx"
                out_path = os.path.join(out_dir, out_name)
                generate_docx_memorial(self.template_path, out_path, dados)
                messagebox.showinfo(
                    "Aviso - Arquivo Word Aberto",
                    f"O arquivo Word anterior estava aberto.\n"
                    f"Foi gerada uma nova cópia como:\n{out_name}"
                )
        except Exception as e:
            messagebox.showerror("Erro", f"Falha ao gerar o Memorial Descritivo: {e}")
            return

        msg_sucesso = (
            f"Área Ambiental Detectada no Projeto!\n\n"
            f"Documentos técnicos e croqui gerados com sucesso na mesma pasta:\n"
            f"Local: {out_dir}\n\n"
            f"1. {out_name} (Word - Memorial Descritivo da Intervenção)\n"
            f"2. {checklist_out_name} (Excel - Checklist com Apontamento de Restrições)\n"
        )
        if croqui_ok:
            msg_sucesso += f"3. {croqui_out_name} (Ortofoto de Satélite Georreferenciada)\n"

        messagebox.showinfo("Sucesso - Documentos Ambientais Gerados", msg_sucesso)
        self.root.destroy()

        try:
            os.startfile(out_path)
        except Exception:
            pass

        if checklist_ok and os.path.exists(checklist_out_path):
            try:
                os.startfile(checklist_out_path)
            except Exception:
                pass

    def run(self):
        self.root.mainloop()


if __name__ == '__main__':
    candidatos_tpl = [
        r"G:\Meu Drive\1-HBT\04_MODELOS\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
        r"G:\Meu Drive\1-HBT\05_PROJETOS_HISTORICOS\SETEMBRO\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "modelos", "MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx"),
        r"G:\Meu Drive\1-HBT\MODELO MEMORIAL DESCRITIVO_AMBIENTAL.docx",
    ]
    template = ""
    for t_cand in candidatos_tpl:
        if os.path.exists(t_cand):
            template = t_cand
            break

    pts = r"C:\Temp\kmz_points_input.json"
    boxes = r"C:\Temp\kmz_boxes.txt"
    fuso = 23
    out_dir = None

    if len(sys.argv) > 1 and sys.argv[1].strip():
        template = sys.argv[1].strip()
    if len(sys.argv) > 2 and sys.argv[2].strip():
        pts = sys.argv[2].strip()
    if len(sys.argv) > 3 and sys.argv[3].strip():
        boxes = sys.argv[3].strip()
    if len(sys.argv) > 4 and sys.argv[4].strip():
        try:
            fuso = int(sys.argv[4])
        except ValueError:
            fuso = 23
    if len(sys.argv) > 5 and sys.argv[5].strip():
        out_dir = sys.argv[5].strip()

    app = MemorialGUI(template, pts, boxes, fuso, candidate_out_dir=out_dir)
    app.run()
