"""
Módulo de Preenchimento Automático do Checklist Ambiental (.xlsx)
Preenche os dados do projeto e as incidências cartográficas identificadas
pelo levantamento georreferenciado (KMZ) de acordo com os padrões da concessionária/HBT.
"""

import os
import re
from typing import Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Alignment

TEMPLATE_DEFAULT = r"G:\Meu Drive\1-HBT\HBT\checklist.xlsx"
_mod_dir = os.path.dirname(os.path.abspath(__file__))
_root_dir = os.path.dirname(_mod_dir)
if os.path.exists(os.path.join(_mod_dir, "modelos", "checklist_modelo.xlsx")):
    TEMPLATE_FALLBACK = os.path.join(_mod_dir, "modelos", "checklist_modelo.xlsx")
else:
    TEMPLATE_FALLBACK = os.path.join(_root_dir, "modelos", "checklist_modelo.xlsx")

# Mapeamento das linhas da tabela de APP (Planilha1)
APP_ROW_MAPPING = {
    "RIO_MENOR_10M": 24,
    "RIO_10M_50M": 25,
    "RIO_50M_200M": 26,
    "RIO_200M_600M": 27,
    "RIO_MAIOR_600M": 28,
    "LAGO_URBANO": 29,
    "LAGO_RURAL_ATE_20HA": 30,
    "LAGO_RURAL_MAIOR_20HA": 31,
    "RESERVATORIO": 32,
    "NASCENTE": 33,
    "ENCOSTA": 34,
    "RESTINGA": 35,
    "MANGUEZAL": 36,
    "TABULEIRO": 37,
    "TOPO_DE_MORRO": 38,
    "ALTITUDE_1800M": 39,
    "VEREDA": 40,
    "INTERESSE_SOCIAL": 41,
}


def extrair_municipio(local_obra: str) -> str:
    """Extrai o nome provável do município a partir da string de localização."""
    if not local_obra:
        return "NÃO INFORMADO"

    texto = local_obra.upper().strip()

    # Se contiver hífen com estado (ex: "ANGRA DOS REIS - RJ")
    match_uf = re.search(r'([A-Z\sÀ-Ú]+)\s*-\s*([A-Z]{2})', texto)
    if match_uf:
        cidade = match_uf.group(1).strip()
        # Se contiver vírgula (ex: "ILHA GRANDE, ANGRA DOS REIS")
        if ',' in cidade:
            cidade = cidade.split(',')[-1].strip()
        return cidade

    # Se contiver vírgulas (ex: "RUA PRINCIPAL, SÃO GONÇALO")
    partes = [p.strip() for p in texto.split(',') if p.strip()]
    if len(partes) > 1:
        return partes[-1]

    return texto


def analisar_incidencias_kmz(boxes_txt_path: str) -> Dict:
    """
    Analisa as linhas de categorias e especificações do kmz_boxes.txt
    e classifica em APP, Unidade de Conservação/Restrição e Área de Exclusão.
    """
    resultado = {
        "em_app": False,
        "app_rows": set(),
        "em_uc": False,
        "ucs": [],
        "em_exclusao": False,
        "exclusoes": [],
        "observacoes": []
    }

    if not boxes_txt_path or not os.path.exists(boxes_txt_path):
        return resultado

    with open(boxes_txt_path, 'r', encoding='utf-8') as f:
        for line in f:
            line_str = line.strip()
            if not line_str:
                continue

            parts = line_str.split('|')
            cat = parts[0].strip().upper()
            spec = parts[1].strip().upper() if len(parts) > 1 else ""
            comb = f"{cat} {spec}".strip()

            # 1. Checagem de APP (Área de Preservação Permanente)
            if "PRESERVAÇÃO PERMANENTE" in comb or "PRESERVACAO PERMANENTE" in comb or "APP" in comb.split():
                resultado["em_app"] = True

                # Identifica qual linha da tabela de APP marcar
                if "MANGUE" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["MANGUEZAL"])
                elif "RESTINGA" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["RESTINGA"])
                elif "NASCENTE" in comb or "OLHO" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["NASCENTE"])
                elif "ENCOSTA" in comb or "DECLIVIDADE" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["ENCOSTA"])
                elif "TOPO" in comb or "MORRO" in comb or "MONTANHA" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["TOPO_DE_MORRO"])
                elif "TABULEIRO" in comb or "CHAPADA" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["TABULEIRO"])
                elif "VEREDA" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["VEREDA"])
                elif "LAGO" in comb or "LAGOA" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["LAGO_URBANO"])
                elif "RESERVAT" in comb or "REPRESA" in comb:
                    resultado["app_rows"].add(APP_ROW_MAPPING["RESERVATORIO"])
                else:
                    # Padrão: Rios (Largura menor que 10m / 30m)
                    resultado["app_rows"].add(APP_ROW_MAPPING["RIO_MENOR_10M"])

            # 2. Checagem de Área de Exclusão
            if any(term in comb for term in ["EXCLUSÃO", "EXCLUSAO", "INTANGÍVEL", "INTANGIVEL", "TERRA INDÍGENA", "INDIGENA"]):
                resultado["em_exclusao"] = True
                nome_exclusao = spec if spec else cat
                if nome_exclusao not in resultado["exclusoes"]:
                    resultado["exclusoes"].append(nome_exclusao)

            # 3. Checagem de Unidade de Conservação / Zona de Amortecimento / Área com Restrição
            termos_uc = [
                "UNIDADE DE CONSERVAÇÃO", "UNIDADE DE CONSERVACAO", "UC",
                "ZONA DE AMORTECIMENTO", "ÁREA COM RESTRIÇÃO", "AREA COM RESTRICAO",
                "PARQUE", "REBIO", "ESTAÇÃO ECOLÓGICA", "ESTACAO ECOLOGICA",
                "APA", "ÁREA DE PROTEÇÃO AMBIENTAL", "AREA DE PROTECAO AMBIENTAL",
                "FLONA", "RESEX", "ARIE", "RPPN"
            ]
            if any(term in comb for term in termos_uc):
                resultado["em_uc"] = True
                # Formata o nome amigável da UC ou Zona
                if "ZONA DE AMORTECIMENTO" in comb:
                    nome_uc = f"ZONA DE AMORTECIMENTO ({spec})" if spec else cat
                else:
                    nome_uc = spec if spec else cat
                if nome_uc not in resultado["ucs"]:
                    resultado["ucs"].append(nome_uc)

    return resultado


def preencher_checklist_ambiental(
    template_path: str,
    output_path: str,
    dados_projeto: Dict[str, str],
    boxes_txt_path: str = ""
) -> str:
    """
    Carrega o modelo do checklist, preenche os dados e salva na mesma pasta do Memorial.
    """
    # Se template_path não existir, tenta o padrão ou fallback local
    if not os.path.exists(template_path):
        if os.path.exists(TEMPLATE_DEFAULT):
            template_path = TEMPLATE_DEFAULT
        elif os.path.exists(TEMPLATE_FALLBACK):
            template_path = TEMPLATE_FALLBACK
        else:
            raise FileNotFoundError(f"Modelo do checklist não encontrado em: {template_path}")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    # Analisa incidências ambientais do KMZ
    analise = analisar_incidencias_kmz(boxes_txt_path)

    wb = openpyxl.load_workbook(template_path)
    if "Planilha1" not in wb.sheetnames:
        raise ValueError("Aba 'Planilha1' não encontrada no modelo de checklist.")

    ws = wb["Planilha1"]

    # 1. Cabeçalho
    num_sob = dados_projeto.get("NUMERO_PROJETO", "").strip().upper()
    municipio = dados_projeto.get("MUNICIPIO", "").strip().upper()
    if not municipio:
        municipio = extrair_municipio(dados_projeto.get("LOCAL_OBRA", ""))

    projetista = dados_projeto.get("PROJETISTA_RESPONSAVEL", "FABIO VARGAS").strip().upper()
    empresa = dados_projeto.get("EMPRESA", "HBT").strip().upper()
    validado_por = dados_projeto.get("VALIDADO_POR", "PROJETO").strip().upper()

    ws["B6"] = num_sob
    ws["B7"] = municipio
    ws["B8"] = validado_por
    ws["B9"] = projetista
    ws["B10"] = empresa

    # 2. Status Geral de Interseção
    # B12: SOB em APP (SIM / NÃO)
    tem_app = dados_projeto.get("SOB_EM_APP")
    if tem_app is None:
        tem_app = "SIM" if analise["em_app"] else "NÃO"
    ws["B12"] = tem_app.strip().upper()

    # B14: SOB em Unidade de Conservação/Zona de Amortecimento/Área com Restrição (SIM / NÃO)
    tem_uc = dados_projeto.get("SOB_EM_UC")
    if tem_uc is None:
        tem_uc = "SIM" if analise["em_uc"] else "NÃO"
    ws["B14"] = tem_uc.strip().upper()

    # B15: Qual/Quais? (Listar UCs)
    qual_uc = dados_projeto.get("QUAIS_UCS")
    if qual_uc is None:
        qual_uc = ", ".join(analise["ucs"]) if analise["ucs"] else ""
    ws["B15"] = qual_uc if ws["B14"].value == "SIM" else None

    # B17: SOB em Área de Exclusão (SIM / NÃO)
    tem_exc = dados_projeto.get("SOB_EM_EXCLUSAO")
    if tem_exc is None:
        tem_exc = "SIM" if analise["em_exclusao"] else "NÃO"
    ws["B17"] = tem_exc.strip().upper()

    # B18: Qual/Quais? (Listar Exclusões)
    qual_exc = dados_projeto.get("QUAIS_EXCLUSOES")
    if qual_exc is None:
        qual_exc = ", ".join(analise["exclusoes"]) if analise["exclusoes"] else ""
    ws["B18"] = qual_exc if ws["B17"].value == "SIM" else None

    # 3. Tabela de Detalhamento de APP (Linhas 24 a 41, Coluna D)
    # Primeiro limpa todas as marcas prévias
    for r in range(24, 42):
        ws.cell(r, 4).value = None

    # Se houver APP, marca o checkmark na(s) linha(s) correspondente(s)
    check_char = "✓"
    align_center = Alignment(horizontal="center", vertical="center")

    if ws["B12"].value == "SIM":
        # Se houver linhas mapeadas da análise
        rows_to_check = analise["app_rows"] if analise["app_rows"] else {24}
        for r in rows_to_check:
            c = ws.cell(r, 4)
            c.value = check_char
            c.alignment = align_center

    # 4. Observações (A44 é a célula mestre do intervalo mesclado A44:D44)
    obs = dados_projeto.get("OBSERVACOES_CHECKLIST", "").strip()
    if not obs and (dados_projeto.get("QTD_PODAS", "0") != "0" or dados_projeto.get("QTD_SUPRESSOES", "0") != "0"):
        obs = f"Podas previstas: {dados_projeto.get('QTD_PODAS', '0')} | Supressões previstas: {dados_projeto.get('QTD_SUPRESSOES', '0')}."
    if obs:
        ws["A44"] = f"Observações (caso necessário): {obs}"
    else:
        ws["A44"] = "Observações (caso necessário):"

    wb.save(output_path)
    return output_path
