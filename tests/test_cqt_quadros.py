import os
import pytest
from cqt_quadros_generator import format_num, extract_quadros_data, generate_cad_lisp, save_quadros_json

REAL_EXCEL = r"G:\Meu Drive\1-HBT\HBT\SETEMBRO\05\A057146257\AN_A057146257_CQT.xlsm"

def test_format_num():
    assert format_num(None) == "-"
    assert format_num(0) == "-"
    assert format_num("-") == "-"
    assert format_num(1.3, 2) == "1,30"
    assert format_num(65, 2) == "65,00"
    assert format_num(0.4416666, is_pct=True) == "44,17%"
    assert format_num(0.433, is_pct=True) == "43,30%"

def test_extract_quadros_data_real():
    if not os.path.exists(REAL_EXCEL):
        pytest.skip("Planilha real de teste não encontrada no caminho especificado.")

    data = extract_quadros_data(REAL_EXCEL, custom_ponto_trafo="77245")
    assert "resumo_trafo" in data
    resumo = data["resumo_trafo"]

    assert "ATUAL+NC" in resumo
    assert "ATUAL" in resumo

    anc = resumo["ATUAL+NC"]
    assert anc["ponto_trafo"] == "77245"
    assert anc["residencial"]["qtde"] == "50"
    assert anc["residencial"]["dmdi"] == "1,30"
    assert anc["residencial"]["total"] == "65,00"
    assert anc["req_resid"]["qtde"] == "1"
    assert anc["req_resid"]["total"] == "1,30"
    assert anc["ilum_publica"]["qtde"] == "10"
    assert anc["ilum_publica"]["total"] == "1,25"
    assert anc["total_clientes"] == "50"
    assert anc["max_prev"] == "66,25kVA"
    assert anc["carregamento"] == "44,17%"

    at = resumo["ATUAL"]
    assert at["residencial"]["qtde"] == "49"
    assert at["residencial"]["dmdi"] == "1,30"
    assert at["residencial"]["total"] == "63,70"
    assert at["n_residencial"]["qtde"] == "-"
    assert at["ilum_publica"]["qtde"] == "10"
    assert at["ilum_publica"]["total"] == "1,25"
    assert at["total_clientes"] == "49"
    assert at["max_prev"] == "64,95kVA"
    assert at["carregamento"] == "43,30%"

    # Valida Queda de Tensão Acumulada por Trecho extraída diretamente do Excel
    trechos = data["trechos_cqt"]
    assert len(trechos) == 3
    assert trechos[0]["trecho"] == "A"
    assert trechos[0]["qt_atual"] == "3,85%"
    assert trechos[0]["qt_atual_nc"] == "3,99%"
    assert trechos[1]["trecho"] == "B"
    assert trechos[1]["qt_atual"] == "3,85%"
    assert trechos[1]["qt_atual_nc"] == "4,02%"
    assert trechos[2]["trecho"] == "C"
    assert trechos[2]["qt_atual"] == "4,64%"
    assert trechos[2]["qt_atual_nc"] == "4,77%"

def test_generate_cad_lisp_and_json(tmp_path):
    if not os.path.exists(REAL_EXCEL):
        pytest.skip("Planilha real não encontrada.")

    json_path = str(tmp_path / "quadros.json")
    save_quadros_json(REAL_EXCEL, json_out_path=json_path, custom_ponto_trafo="77245")
    assert os.path.exists(json_path)

    # Teste de leitura do LSP gerado
    with open(r"C:\Temp\cqt_draw_quadros.lsp", "r", encoding="windows-1252") as f:
        content = f.read()

    assert "c:CQTQUADROS" in content
    assert "77245 - ATUAL+NC" in content
    assert "77245 - ATUAL" in content
    assert "ILUM. PÚBLICA" in content
    assert "MÁX. PREV." in content
    assert "CARREGAMENTO" in content


def test_extract_quadros_excel_existente_xlsx(tmp_path):
    """Testa geração de quadros a partir de um arquivo .xlsx existente sem aba CAPA-RESUMO."""
    import openpyxl

    wb = openpyxl.Workbook()
    # Cria abas de referência diretamente sem usar CAPA-RESUMO
    ws_at = wb.active
    ws_at.title = "ATUAL"
    ws_at["C3"] = "CLASSE C"
    ws_at["C4"] = 49
    ws_at["C5"] = 1.30
    ws_at["D6"] = 63.70
    ws_at["D7"] = 1.25
    ws_at["C10"] = 150
    ws_at["D10"] = 64.95
    ws_at["B10"] = 0.433
    # 10 luminárias IP125W
    ws_at["U3"] = 10

    ws_anc = wb.create_sheet("ATUAL+NC")
    ws_anc["C3"] = "CLASSE C"
    ws_anc["C4"] = 50
    ws_anc["C5"] = 1.30
    ws_anc["D6"] = 65.00
    ws_anc["D7"] = 1.25
    ws_anc["C10"] = 150
    ws_anc["D10"] = 66.25
    ws_anc["B10"] = 0.4417
    ws_anc["U3"] = 10

    ws_cqt_at = wb.create_sheet("CQT ATUAL")
    ws_cqt_at["F4"] = "77245"
    ws_cqt_at["A11"] = "A"
    ws_cqt_at["C11"] = 2.10
    ws_cqt_at["L11"] = 1.50

    ws_cqt_anc = wb.create_sheet("CQT ATUAL+NC")
    ws_cqt_anc["F4"] = "77245"
    ws_cqt_anc["A11"] = "A"
    ws_cqt_anc["C11"] = 2.10
    ws_cqt_anc["L11"] = 2.10

    excel_file = tmp_path / "CQT_PROJETO_TESTE.xlsx"
    wb.save(str(excel_file))
    wb.close()

    # 1. Extração com ponto trafo "auto"
    data = extract_quadros_data(str(excel_file), custom_ponto_trafo="auto")
    assert "resumo_trafo" in data
    assert "ATUAL" in data["resumo_trafo"]
    assert "ATUAL+NC" in data["resumo_trafo"]
    assert data["resumo_trafo"]["ATUAL"]["ponto_trafo"] == "77245"

    # 2. Geração de JSON e LISP
    json_out = str(tmp_path / "saida_quadros.json")
    save_quadros_json(str(excel_file), json_out_path=json_out, custom_ponto_trafo=None)
    assert os.path.exists(json_out)

    # 3. Teste da CLI
    import sys
    import subprocess
    py_exec = sys.executable
    cmd = [py_exec, "cqt/cqt_quadros_generator.py", str(excel_file), "auto"]
    res = subprocess.run(cmd, capture_output=True, text=True)
    assert res.returncode == 0
    assert "[SUCESSO]" in res.stdout or "exportados" in res.stdout


def test_ignora_capa_resumo_e_usa_abas_referencia(tmp_path):
    """Garante que dados defasados na aba CAPA-RESUMO são ignorados em favor das abas ATUAL e CQT ATUAL."""
    import openpyxl

    wb = openpyxl.Workbook()
    # Cria CAPA-RESUMO com dados 'lixo' / defasados de template
    ws_capa = wb.active
    ws_capa.title = "CAPA-RESUMO"
    ws_capa["B16"] = 999  # Residencial fake
    ws_capa["B21"] = 888  # Total clientes fake
    ws_capa["B22"] = 999.9  # Max prev fake
    ws_capa["B23"] = 0.999  # Carregamento fake

    # Cria abas de referência com dados reais da obra
    ws_at = wb.create_sheet("ATUAL")
    ws_at["C3"] = "CLASSE B"
    ws_at["C4"] = 20
    ws_at["C5"] = 1.70
    ws_at["D6"] = 34.00
    ws_at["D7"] = 0.75
    ws_at["C10"] = 112.5
    ws_at["D10"] = 34.75
    ws_at["B10"] = 0.3089

    ws_cqt_at = wb.create_sheet("CQT ATUAL")
    ws_cqt_at["F4"] = "A55323"
    ws_cqt_at["A11"] = "A"
    ws_cqt_at["C11"] = 1.50
    ws_cqt_at["L11"] = 2.45

    fpath = tmp_path / "TESTE_CAPA_IGNORADA.xlsx"
    wb.save(str(fpath))
    wb.close()

    data = extract_quadros_data(str(fpath))
    resumo_at = data["resumo_trafo"]["ATUAL"]

    # Deve extrair os dados reais da aba ATUAL, não da CAPA-RESUMO
    assert resumo_at["ponto_trafo"] == "A55323"
    assert resumo_at["residencial"]["qtde"] == "20"
    assert resumo_at["residencial"]["dmdi"] == "1,70"
    assert resumo_at["residencial"]["total"] == "34,00"
    assert resumo_at["total_clientes"] == "20"
    assert resumo_at["max_prev"] == "34,75kVA"
    assert resumo_at["carregamento"] == "30,89%"


def test_extract_quadros_analytical_fallback(tmp_path):
    """Testa o motor de fallback analítico quando a planilha não possui resumo pré-calculado em células."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws_at = wb.active
    ws_at.title = "ATUAL"
    ws_at["C3"] = "CLASSE C"
    ws_at["C10"] = 45
    # Linha 3 de cargas (trecho A): 10 mono, 2 IP70
    ws_at["L3"] = 10
    ws_at["S3"] = 2

    ws_cqt_at = wb.create_sheet("CQT ATUAL")
    ws_cqt_at["F4"] = "99881"
    # Trecho A sem coluna L de QT preenchida
    ws_cqt_at["A11"] = "A"
    ws_cqt_at["B11"] = "TRAFO"
    ws_cqt_at["C11"] = 0.50  # 50m
    ws_cqt_at["I11"] = "3x70+1x70 - Al - F.P 1"

    fpath = tmp_path / "TESTE_FALLBACK_ANALITICO.xlsx"
    wb.save(str(fpath))
    wb.close()

    data = extract_quadros_data(str(fpath), custom_ponto_trafo="99881")
    assert "resumo_trafo" in data
    assert "ATUAL" in data["resumo_trafo"]
    resumo_at = data["resumo_trafo"]["ATUAL"]
    assert resumo_at["ponto_trafo"] == "99881"
    assert resumo_at["total_clientes"] == "10"
    assert len(data["trechos_cqt"]) >= 1



