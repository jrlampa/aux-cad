"""
Testes Unitários para o Módulo de Orçamento e Levantamento de Materiais (orcamento/).
Valida classificação analítica de postes, condutores, equipamentos, estruturas e geração Excel.
"""

import os
import openpyxl
import pytest

from orcamento_core import (
    carregar_catalogo,
    classificar_cabo,
    classificar_equipamento_ou_estrutura,
    classificar_poste,
    processar_dados_orcamento
)
from orcamento_excel import exportar_planilha_orcamento


def test_catalogo_padrao_existente():
    catalogo = carregar_catalogo()
    assert isinstance(catalogo, dict)
    assert len(catalogo) >= 15
    assert "PST_CONC_12_300" in catalogo
    assert "CABO_BT_3X70_50" in catalogo
    assert "TRAFO_45_KVA" in catalogo


def test_classificacao_postes():
    # 1. Padrão barra: XX/YYY
    chave, alt, esf = classificar_poste("POSTE_12/300_NOVO")
    assert chave == "PST_CONC_12_300"
    assert alt == 12
    assert esf == 300

    # 2. Padrão hífen: DT-11-600
    chave, alt, esf = classificar_poste("DT-11-600")
    assert chave == "PST_CONC_11_600"
    assert alt == 11
    assert esf == 600

    # 3. Padrão por extenso: 10M 300DAN
    chave, alt, esf = classificar_poste("POSTE CONCRETO 10M 300DAN")
    assert chave == "PST_CONC_10_300"
    assert alt == 10
    assert esf == 300

    # 4. Caso não identificado: retorna genérico
    chave, alt, esf = classificar_poste("POSTE_DESCONHECIDO")
    assert chave == "PST_CONC_GENERICO"
    assert alt == 0
    assert esf == 0


def test_classificacao_cabos():
    # BT Multiplexado
    chave, tensao, _ = classificar_cabo("BT_CABO_3X70+50")
    assert chave == "CABO_BT_3X70_50"
    assert tensao == "BT"

    chave, tensao, _ = classificar_cabo("REDE_BT_1X35")
    assert chave == "CABO_BT_1X35_35"
    assert tensao == "BT"

    # MT Nu CAA
    chave, tensao, _ = classificar_cabo("MT_3X1/0_CAA_RAVEN")
    assert chave == "CABO_MT_3X1_0_CAA"
    assert tensao == "MT"

    # MT Compacto
    chave, tensao, _ = classificar_cabo("REDE_COMPACTA_50MM2")
    assert chave == "CABO_MT_COMPACTO_50"
    assert tensao == "MT"


def test_classificacao_equipamentos_e_estruturas():
    # Transformadores
    res = classificar_equipamento_ou_estrutura("TRAFO 45KVA 13.8KV")
    assert res is not None
    assert res[0] == "TRAFO_45_KVA"
    assert res[1] == "Equipamentos"

    # Chaves e Proteção
    res = classificar_equipamento_ou_estrutura("CHAVE FUSIVEL MATHEUS 15KV")
    assert res is not None
    assert res[0] == "CHAVE_FUSIVEL_15KV"

    # Estruturas de MT e BT
    res = classificar_equipamento_ou_estrutura("ESTRUTURA N1 TANGENCIA")
    assert res is not None
    assert res[0] == "ESTRUTURA_N1"

    res = classificar_equipamento_ou_estrutura("ARMACAO SECUNDARIA B1")
    assert res is not None
    assert res[0] == "ESTRUTURA_B1"

    # Estai e Aterramento
    res = classificar_equipamento_ou_estrutura("ESTAI DE ANCORA COMPLETO")
    assert res is not None
    assert res[0] == "ESTAI_ANCORA"


def test_processamento_dados_orcamento_completo():
    dados_cad = {
        "postes": [
            {"id": "P1", "bloco": "POSTE_12_300", "attrs": "NUM:1", "x": 600000.0, "y": 7500000.0, "layer": "POSTES"},
            {"id": "P2", "bloco": "POSTE_12_600", "attrs": "NUM:2", "x": 600040.0, "y": 7500000.0, "layer": "POSTES"}
        ],
        "cabos": [
            {"id": "C1", "layer": "BT_CABO_3X70+50", "comprimento": 40.0, "vaos": 1},
            {"id": "C2", "layer": "MT_CABO_3X1/0_CAA", "comprimento": 100.0, "vaos": 2}
        ],
        "estruturas": [
            {"id": "E1", "tipo": "N1", "bloco": "ESTR_N1"}
        ],
        "equipamentos": [
            {"id": "EQ1", "tipo": "TRAFO 45 KVA", "bloco": "TRAFO"}
        ]
    }

    # Margem padrão de 5%
    resultado = processar_dados_orcamento(dados_cad, margem_cabos_pct=5.0)

    assert resultado["qtd_postes"] == 2
    assert resultado["qtd_trechos_cabo"] == 2
    # Comprimento orçado = 40 * 1.05 (42.0) + 100 * 1.05 (105.0) = 147.0
    assert resultado["extensao_total_m"] == 147.0
    assert resultado["total_geral"] > 0.0

    chaves_geradas = [m.chave for m in resultado["materiais"]]
    assert "PST_CONC_12_300" in chaves_geradas
    assert "PST_CONC_12_600" in chaves_geradas
    assert "CABO_BT_3X70_50" in chaves_geradas
    assert "CABO_MT_3X1_0_CAA" in chaves_geradas
    assert "ESTRUTURA_N1" in chaves_geradas
    assert "TRAFO_45_KVA" in chaves_geradas


def test_exportacao_planilha_orcamento_excel(tmp_path):
    dados_cad = {
        "postes": [
            {"id": "P1", "bloco": "POSTE_12_300", "attrs": "NUM:1", "x": 600000.0, "y": 7500000.0, "layer": "POSTES"}
        ],
        "cabos": [
            {"id": "C1", "layer": "BT_CABO_3X70+50", "comprimento": 50.0, "vaos": 1}
        ]
    }
    processados = processar_dados_orcamento(dados_cad, margem_cabos_pct=5.0)
    arquivo_excel = os.path.join(tmp_path, "ORCAMENTO_TESTE.xlsx")

    caminho_gerado = exportar_planilha_orcamento(
        dados_processados=processados,
        caminho_saida=arquivo_excel,
        num_projeto="OS-2026-TESTE",
        municipio="NITEROI - RJ",
        projetista="ENG. TESTE"
    )

    assert os.path.exists(caminho_gerado)

    # Abre e valida as abas e estrutura
    wb = openpyxl.load_workbook(caminho_gerado)
    abas_esperadas = ["RESUMO ORÇAMENTO", "DETALHAMENTO POSTES", "DETALHAMENTO CABOS", "BANCO DE PREÇOS"]
    for aba in abas_esperadas:
        assert aba in wb.sheetnames

    ws_resumo = wb["RESUMO ORÇAMENTO"]
    assert "LEVANTAMENTO QUANTITATIVO E ORÇAMENTO DE MATERIAIS" in str(ws_resumo["A1"].value)
    assert ws_resumo["B2"].value == "OS-2026-TESTE"


def test_carregar_dados_json(tmp_path):
    import json
    from orcamento_manager import carregar_dados_json

    # Teste com arquivo inexistente
    dados_vazio = carregar_dados_json("caminho_inexistente.json")
    assert dados_vazio["postes"] == []

    # Teste com arquivo existente
    arq = tmp_path / "dados.json"
    conteudo = {"postes": [{"id": "P1"}], "cabos": [], "estruturas": [], "equipamentos": []}
    arq.write_text(json.dumps(conteudo), encoding="utf-8")

    dados = carregar_dados_json(str(arq))
    assert len(dados["postes"]) == 1
    assert dados["postes"][0]["id"] == "P1"

    # Teste com JSON corrompido / inválido
    arq_corrompido = tmp_path / "corrompido.json"
    arq_corrompido.write_text("{json invalido:", encoding="utf-8")
    dados_inv = carregar_dados_json(str(arq_corrompido))
    assert dados_inv == {"postes": [], "cabos": [], "estruturas": [], "equipamentos": []}



def test_orcamento_manager_cli(tmp_path, monkeypatch):
    import json
    import sys
    from orcamento_manager import main

    arq_json = tmp_path / "dados_cli.json"
    arq_xlsx = tmp_path / "saida_cli.xlsx"

    conteudo = {
        "postes": [{"id": "P1", "bloco": "12/300"}],
        "cabos": [{"id": "C1", "layer": "BT_CABO", "comprimento": 30.0}],
        "estruturas": [],
        "equipamentos": []
    }
    arq_json.write_text(json.dumps(conteudo), encoding="utf-8")

    test_args = [
        "orcamento_manager.py",
        "--json", str(arq_json),
        "--saida", str(arq_xlsx),
        "--no-gui",
        "--projeto", "CLI-TESTE-2026"
    ]
    monkeypatch.setattr(sys, "argv", test_args)

    main()

    assert os.path.exists(str(arq_xlsx))


def test_orcamento_database_consultas():
    import orcamento_database as db

    # 1. Contratos
    contratos = db.listar_contratos()
    assert len(contratos) >= 4
    nomes_emp = [c["empreiteira"] for c in contratos]
    assert "ENGELMIG" in nomes_emp
    assert "INDICA" in nomes_emp

    # 2. Material
    mat = db.obter_material("DT11600")
    assert mat is not None
    assert mat["preco_unitario"] > 0
    assert "11M" in mat["descricao"] or "600" in mat["descricao"]

    # 3. Busca textual
    postes = db.buscar_materiais("POSTE", limite=5)
    assert len(postes) > 0

    # 4. Mão de Obra
    mo = db.obter_servico_mo("AHO110", empreiteira="ENGELMIG", territorio="A")
    assert mo is not None
    assert mo["custo_unitario"] > 0
    assert mo["contrato"] == "JA10114811"


def test_orcamento_kits_explosao():
    import orcamento_kits as kits

    # 1. Normalização
    assert "EST.M1" in kits.normalizar_codigo_kit("M1", esforco_poste=300)
    assert "EST.B1" in kits.normalizar_codigo_kit("B1", esforco_poste=300)
    assert "EST.M2" in kits.normalizar_codigo_kit("M2", esforco_poste=300)
    assert "EST.B2" in kits.normalizar_codigo_kit("B2", esforco_poste=300)
    assert "TR-TRI-45KVA" in kits.normalizar_codigo_kit("TRAFO 45 KVA")
    assert "ATER_EQUIP" in kits.normalizar_codigo_kit("ATERRAMENTO")
    assert "ESTAI-7,9" in kits.normalizar_codigo_kit("ESTAI DE ANCORA")
    assert "CXDERDAT" == kits.normalizar_codigo_kit("DAE")

    # 2. Explosão de Kit de Estrutura
    itens_m1 = kits.explodir_kit("M1", "NA", esforco_poste=300)
    assert len(itens_m1) >= 5
    codigos_m1 = [it.codigo_material for it in itens_m1]
    assert any(c for c in codigos_m1 if c.isdigit())

    # 3. Explosão de Trafo
    itens_trafo = kits.explodir_kit("TRAFO 45KVA", "NA")
    assert len(itens_trafo) >= 5

    # 4. Explosão de Aterramento e Estai
    itens_ater = kits.explodir_kit("ATERRAMENTO")
    assert len(itens_ater) >= 1

    itens_estai = kits.explodir_kit("ESTAI")
    assert len(itens_estai) >= 1

    # 5. Explosão completa de poste
    comp = kits.explodir_poste_completo(["M1", "B1"], ["TRAFO 45KVA", "ATERRAMENTO"], "NA", 600)
    assert len(comp) >= 12


def test_orcamento_core_mao_de_obra_e_bdi():
    from orcamento_core import processar_dados_orcamento

    dados = {
        "postes": [{"id": "P1", "bloco": "POSTE_11_600", "attrs": "11/600", "estruturas": ["M1"], "equipamentos": ["TRAFO 45KVA"]}],
        "cabos": [{"id": "C1", "layer": "CABO_MT_1/0_RAVEN", "comprimento": 100.0, "vaos": 2}]
    }

    # Teste para ENGELMIG Classe A
    res_a = processar_dados_orcamento(dados, empreiteira="ENGELMIG", territorio="A", bdi_pct=10.0)
    assert res_a["total_materiais"] > 0
    assert res_a["total_mao_de_obra"] > 0
    assert res_a["total_bdi"] > 0
    assert res_a["total_geral"] == round(res_a["total_materiais"] + res_a["total_mao_de_obra"] + res_a["total_bdi"], 2)
    assert len(res_a["mao_de_obra"]) >= 3

    # Teste para INDICA Classe B
    res_b = processar_dados_orcamento(dados, empreiteira="INDICA", territorio="B", bdi_pct=0.0)
    assert res_b["total_mao_de_obra"] > 0
    assert res_b["empreiteira"] == "INDICA"
    assert res_b["territorio"] == "B"


def test_orcamento_excel_abas_novas(tmp_path):
    from orcamento_core import processar_dados_orcamento
    from orcamento_excel import exportar_planilha_orcamento

    dados = {
        "postes": [{"id": "P1", "bloco": "POSTE_11_600", "attrs": "11/600", "estruturas": ["M1"], "equipamentos": ["TRAFO 45KVA"]}],
        "cabos": [{"id": "C1", "layer": "CABO_MT_1/0_RAVEN", "comprimento": 50.0, "vaos": 1}]
    }
    res = processar_dados_orcamento(dados, empreiteira="PROGEN", territorio="C")
    out_file = os.path.join(tmp_path, "TESTE_OFICIAL_ABAS.xlsx")

    caminho = exportar_planilha_orcamento(res, out_file)
    assert os.path.exists(caminho)

    wb = openpyxl.load_workbook(caminho)
    assert "DETALHAMENTO MÃO DE OBRA" in wb.sheetnames
    assert "FORMATO sisORÇA" in wb.sheetnames

    ws_so = wb["FORMATO sisORÇA"]
    assert ws_so["A2"].value == "P1"
    assert ws_so["B2"].value == "DT11600"
    assert ws_so["D2"].value == "M1"


