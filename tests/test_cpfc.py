"""
test_cpfc.py - Testes Unitários e de Integração do Módulo CPFC
Módulo CPFC - Cálculo de Participação Financeira do Consumidor (REN 1.059 ANEEL)
"""

import os
import json
import tempfile
import pytest
import openpyxl

from cpfc.cpfc_core import (
    DadosCadastrais,
    OrcamentoObra,
    ParametrosDemanda,
    ResultadoCPFC,
    calcular_cpfc,
    obter_codigo_subgrupo,
    obter_fatores_k,
    SUBGRUPOS_LOOKUP,
    FATORES_K,
)
from cpfc.cpfc_excel import preencher_e_salvar_cpfc, localizar_modelo_cpfc
from cpfc.cpfc_manager import carregar_dados_intercambio


def test_subgrupos_e_fatores_k():
    """Valida o mapeamento de subgrupos e a integridade da tabela de Fatores K."""
    assert obter_codigo_subgrupo("Subgrupo B1") == "B1"
    assert obter_codigo_subgrupo("Subgrupo A4") == "A4"
    assert obter_codigo_subgrupo("Subgrupo B4a") == "B4a"
    assert obter_codigo_subgrupo("Iluminação Publica - Bulbo de lampada") == "B4b"
    assert obter_codigo_subgrupo("B1") == "B1"
    assert obter_codigo_subgrupo("Inexistente") == "B1"  # Fallback seguro

    kc_b1, kg_b1 = obter_fatores_k("B1")
    assert kc_b1 == 1711.94
    assert kg_b1 == 0.0

    kc_a4, kg_a4 = obter_fatores_k("A4")
    assert kc_a4 == 1668.09
    assert kg_a4 == 1176.18


def test_calculo_cpfc_caso_isento():
    """Valida o cálculo do caso oficial de referência (PE_A057032667): Isento de PF."""
    cad = DadosCadastrais(
        projeto="A057032667",
        tipo_ligacao="Acréscimo de carga",
        subgrupo="Subgrupo B1",
        endereco="Rua Jose Timoteo Caldara 834",
        municipio="PETRÓPOLIS",
    )
    orc = OrcamentoObra(
        origem="GOM",
        servico_terceiros=17172.51,
        material=29988.28,
    )
    dem = ParametrosDemanda(
        demanda_carga=40.0,
        demanda_geracao=0.0,
        demanda_cliente_erd=39.52,
        demanda_disponibilizada=69.0,
    )

    res: ResultadoCPFC = calcular_cpfc(cad, orc, dem)

    # Valores esperados conforme a planilha oficial
    assert res.codigo_subgrupo == "B1"
    assert res.fator_kc == 1711.94
    assert res.fator_kg == 0.0
    assert res.erd == pytest.approx(67655.87, abs=0.01)
    assert res.custo_total_obra == pytest.approx(47160.79, abs=0.01)
    assert res.fator_proporcionalizacao == pytest.approx(40.0 / 69.0, abs=0.0001)
    assert res.cop == pytest.approx(27339.59, abs=0.01)
    assert res.pf == 0.0
    assert res.possui_participacao is False
    assert "SEM PARTICIPAÇÃO FINANCEIRA" in res.status_pf
    assert res.alerta is None


def test_calculo_cpfc_com_participacao():
    """Valida caso em que o COP supera o ERD, gerando Participação Financeira devida."""
    cad = DadosCadastrais(
        projeto="A999999999",
        tipo_ligacao="Extensão de Rede",
        subgrupo="Subgrupo B4a",  # Kc menor = 941.39
        endereco="Estrada Geral",
        municipio="Teresópolis",
    )
    orc = OrcamentoObra(
        origem="OPD",
        servico_terceiros=40000.0,
        material=80000.0,  # Total = 120.000,00
    )
    dem = ParametrosDemanda(
        demanda_carga=50.0,
        demanda_geracao=0.0,
        demanda_cliente_erd=50.0,
        demanda_disponibilizada=50.0,  # Fator prop = 1.0 -> COP = 120.000,00
    )

    res: ResultadoCPFC = calcular_cpfc(cad, orc, dem)

    # ERD = 50 * 941.39 = 47.069,50
    # COP = 120.000,00
    # PF = 120.000 - 47.069,50 = 72.930,50
    assert res.erd == pytest.approx(47069.50, abs=0.01)
    assert res.cop == pytest.approx(120000.00, abs=0.01)
    assert res.pf == pytest.approx(72930.50, abs=0.01)
    assert res.possui_participacao is True
    assert "COM PARTICIPAÇÃO FINANCEIRA" in res.status_pf


def test_calculo_cpfc_com_geracao_distribuida():
    """Valida o cálculo de ERD quando há geração superior à demanda de carga."""
    cad = DadosCadastrais(
        projeto="A057000111",
        subgrupo="Subgrupo A4",
    )
    orc = OrcamentoObra(
        servico_terceiros=10000.0,
        material=20000.0,
    )
    dem = ParametrosDemanda(
        demanda_carga=30.0,
        demanda_geracao=75.0,  # Geração excede carga
        demanda_disponibilizada=100.0,
    )

    res: ResultadoCPFC = calcular_cpfc(cad, orc, dem)

    # A4: Kc = 1668.09, Kg = 1176.18
    # d_erd = 30.0
    # Parcela carga: 30 * 1668.09 = 50042.70
    # Parcela geração: (75 - 30) * 1176.18 = 45 * 1176.18 = 52928.10
    # ERD esperado = 50042.70 + 52928.10 = 102970.80
    assert res.erd == pytest.approx(102970.80, abs=0.01)
    assert res.demanda_cliente == 75.0
    assert res.fator_proporcionalizacao == pytest.approx(0.75, abs=0.001)


def test_calculo_cpfc_alerta_demanda_insuficiente():
    """Valida emissão de alerta se a demanda disponibilizada for inferior à demanda do cliente."""
    cad = DadosCadastrais()
    orc = OrcamentoObra(servico_terceiros=10000.0, material=10000.0)
    dem = ParametrosDemanda(
        demanda_carga=50.0,
        demanda_disponibilizada=30.0,  # Menor que a demanda do cliente!
    )

    res: ResultadoCPFC = calcular_cpfc(cad, orc, dem)
    assert res.alerta is not None
    assert "MENOR" in res.alerta
    assert res.cop == 0.0


def test_sanitizacao_e_propriedades_orcamento():
    """Testa sanitização e cálculo de transporte no orçamento."""
    orc = OrcamentoObra(
        origem="gom",
        servico_terceiros=1000.0,
        material=2000.0,
        mao_de_obra=500.0,
    )
    san = orc.sanitizar()
    assert san.origem == "GOM"
    assert san.transporte == pytest.approx(55.0, abs=0.01)  # 11% de 500
    assert san.custo_total == 3000.0


def test_preencher_e_salvar_cpfc_excel():
    """Testa a geração e persistência no arquivo Excel real mantendo fórmulas."""
    cad = DadosCadastrais(
        projeto="TESTE999",
        tipo_ligacao="Extensão de Rede",
        subgrupo="Subgrupo B1",
        endereco="Rua das Flores 123",
        municipio="Petrópolis",
    )
    orc = OrcamentoObra(
        origem="GOM",
        servico_terceiros=15000.0,
        material=25000.0,
    )
    dem = ParametrosDemanda(
        demanda_carga=45.0,
        demanda_geracao=0.0,
        demanda_cliente_erd=45.0,
        demanda_disponibilizada=75.0,
    )

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        arquivo_gerado = preencher_e_salvar_cpfc(cad, orc, dem, tmp_path)
        assert os.path.exists(arquivo_gerado)

        # Abre o arquivo gravado e inspeciona células e fórmulas
        wb = openpyxl.load_workbook(arquivo_gerado, data_only=False)
        ws = wb["BASE PFC MODELO"]

        assert ws["C11"].value == "TESTE999"
        assert ws["C13"].value == "Extensão de Rede"
        assert ws["C15"].value == "Subgrupo B1"
        assert ws["C17"].value == "Rua das Flores 123"
        assert ws["C19"].value == "PETRÓPOLIS"
        assert ws["C24"].value == "GOM"
        assert ws["F30"].value == 15000.0
        assert ws["F31"].value == 25000.0
        assert ws["C41"].value == 45.0
        assert ws["G66"].value == 45.0
        assert ws["G80"].value == 75.0

        # Fórmulas nativas devem permanecer intactas
        assert ws["F32"].value == "=SUM(F30:F31)"
        assert str(ws["G76"].value).startswith("=IFERROR(")
        assert str(ws["G84"].value).startswith("=IFERROR(")
        assert str(ws["G88"].value).startswith("=IF(")
        assert str(ws["G93"].value).startswith("=IFERROR(")

        # Imagens originais devem ser preservadas
        assert len(ws._images) == 2

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_carregar_dados_intercambio():
    """Testa a leitura de arquivo de intercâmbio JSON gerado pelo AutoCAD."""
    dados_exemplo = {
        "projeto": "A057032667",
        "pasta_destino": "C:/Projetos",
        "desenho": "A057032667.dwg",
    }

    with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False, encoding="utf-8") as tmp:
        json.dump(dados_exemplo, tmp)
        tmp_path = tmp.name

    try:
        lidos = carregar_dados_intercambio(tmp_path)
        assert lidos["projeto"] == "A057032667"
        assert lidos["pasta_destino"] == "C:/Projetos"
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_cpfc_gui_lifecycle(monkeypatch):
    """Testa o ciclo de vida da interface gráfica CPFCManagerGUI sem bloquear o loop de eventos."""
    from cpfc.cpfc_gui import CPFCManagerGUI
    import tkinter as tk

    app = CPFCManagerGUI(
        projeto_inicial="A057032667",
        dados_iniciais={
            "endereco": "Rua Teste 10",
            "municipio": "PETRÓPOLIS",
            "tipo_ligacao": "Extensão de Rede",
            "subgrupo": "Subgrupo B1",
            "servico_terceiros": 20000.0,
            "material": 30000.0,
            "demanda_carga": 50.0,
            "demanda_geracao": 0.0,
            "demanda_erd": 50.0,
            "demanda_disponibilizada": 80.0,
        },
    )

    try:
        # Verifica preenchimento dos campos
        assert app.ent_projeto.get() == "A057032667"
        assert app.ent_endereco.get() == "Rua Teste 10"
        assert app.ent_municipio.get() == "PETRÓPOLIS"
        assert app.cb_tipo_ligacao.get() == "Extensão de Rede"
        assert app.cb_subgrupo.get() == "Subgrupo B1"

        # Simula alteração e cálculo em tempo real
        app.ent_terceiros.delete(0, tk.END)
        app.ent_terceiros.insert(0, "15000")
        app.ent_material.delete(0, tk.END)
        app.ent_material.insert(0, "25000")
        app._calcular_em_tempo_real()

        assert "40.000,00" in app.lbl_total_obra.cget("text")

        # Testa sincronização automática de demanda de carga -> ERD
        app.ent_dem_carga.delete(0, tk.END)
        app.ent_dem_carga.insert(0, "60")
        app._on_dem_carga_change()
        assert app.ent_dem_erd.get() == "60"

        # Simula diálogo de salvamento cancelado
        monkeypatch.setattr("tkinter.filedialog.asksaveasfilename", lambda **kwargs: "")
        app._salvar_planilha_dialogo()
        assert app.ultimo_arquivo_salvo is None

        # Simula diálogo de salvamento confirmado
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_save_path = tmp.name

        monkeypatch.setattr("tkinter.filedialog.asksaveasfilename", lambda **kwargs: tmp_save_path)
        monkeypatch.setattr(app, "_exibir_dialogo_sucesso", lambda f: None)
        app._salvar_planilha_dialogo()
        assert app.ultimo_arquivo_salvo == tmp_save_path
        assert os.path.exists(tmp_save_path)

        if os.path.exists(tmp_save_path):
            os.remove(tmp_save_path)

    finally:
        app.destroy()


def test_cpfc_manager_main(monkeypatch):
    """Testa a execução do orquestrador CLI/GUI cpfc_manager."""
    import sys
    from cpfc import cpfc_manager

    class MockGUI:
        def __init__(self, projeto_inicial, pasta_padrao, dados_iniciais):
            self.projeto = projeto_inicial
            self.pasta = pasta_padrao
            self.dados = dados_iniciais

        def mainloop(self):
            pass

    monkeypatch.setattr(cpfc_manager, "CPFCManagerGUI", MockGUI)

    # Chamada com argumentos de diretório e projeto
    monkeypatch.setattr(sys, "argv", ["cpfc_manager.py", "C:/Projetos", "A057000123"])
    cpfc_manager.main()

    # Chamada com argumento de arquivo JSON
    with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False) as tmp:
        json.dump({"projeto": "A123", "pasta_destino": "C:/Temp"}, tmp)
        tmp_name = tmp.name
    try:
        monkeypatch.setattr(sys, "argv", ["cpfc_manager.py", tmp_name])
        cpfc_manager.main()
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)
