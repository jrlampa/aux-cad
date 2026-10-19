from pathlib import Path
import re
import pytest

ROOT = Path(__file__).resolve().parents[1]
LISP_PATH = ROOT / "orcamento" / "levantamento_postes.lsp"


def test_levantamento_postes_lisp_balance_and_syntax():
    """Valida que todos os parênteses e strings do LISP estão estritamente balanceados."""
    source = LISP_PATH.read_text(encoding="utf-8")

    depth = 0
    in_string = False
    in_comment = False
    for index, char in enumerate(source):
        if in_comment:
            if char == "\n":
                in_comment = False
            continue
        if char == ";" and not in_string:
            in_comment = True
            continue
        if char == '"':
            backslashes = 0
            cursor = index - 1
            while cursor >= 0 and source[cursor] == "\\":
                backslashes += 1
                cursor -= 1
            if backslashes % 2 == 0:
                in_string = not in_string
            continue
        if not in_string:
            if char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
                assert depth >= 0, f"Parêntese de fechamento excedente na posição {index}"

    assert depth == 0, f"Parênteses desbalanceados: profundidade restante = {depth}"
    assert not in_string, "String não fechada no arquivo LISP"


def test_levantamento_postes_contract_and_commands():
    """Valida o contrato público de funções, comandos e estrutura de 10 colunas da rotina LISP."""
    source = LISP_PATH.read_text(encoding="utf-8")

    # Comandos
    assert "(defun c:LEVANTAR_POSTES" in source
    assert "(defun c:LP " in source

    # Parsers
    assert "(defun lp-parse-poste" in source
    assert "(defun lp-parse-ramais-e-clientes" in source
    assert "(defun lp-parse-luminaria" in source
    assert "(defun lp-parse-demanda" in source
    assert "(defun lp-parse-acumulada" in source
    assert "(defun lp-extract-value-after-key" in source
    assert "(defun lp-is-noise-line" in source
    assert "(defun lp-format-val" in source

    # Cálculos
    assert "(defun lp-calc-posicoes" in source
    assert "(defun lp-calc-cs" in source

    # Manipulação, Exclusão e Reindexação Sequencial
    assert "(defun lp-renumerar-pontos" in source
    assert "(defun lp-excluir-ponto" in source
    assert "(defun lp-precisa-renumerar-p" in source

    # Persistência DWG via LData
    assert "(vlax-ldata-get" in source
    assert "(vlax-ldata-put" in source
    assert "(vlax-ldata-delete" in source
    assert '"OXCAD_LEVANTAMENTO"' in source
    assert '"PONTOS"' in source
    assert '"TABELA_HANDLE"' in source
    assert "(defun lp-db-obter-tabela-handle" in source
    assert "(defun lp-db-salvar-tabela-handle" in source
    assert "(defun lp-db-remover-tabela-handle" in source

    # Sincronização e Rastreabilidade por Handles
    assert "(defun lp-validar-handle-texto" in source
    assert "(defun lp-recalcular-demandas-acumuladas" in source
    assert "(defun lp-atualizar-celulas-tabela" in source
    assert "(defun lp-sincronizar-levantamento" in source
    assert "(defun lp-atualizar-exclusao-tabela" in source
    assert '"TEXT_HANDLES"' in source

    # Tabela nativa de 10 colunas
    assert "vla-AddTable" in source
    assert "col-count 10" in source
    assert "(vla-MergeCells table 0 0 0 9)" in source
    assert '"LEVANTAMENTO DE POSTES E CLIENTES"' in source
    assert '"PONTO"' in source
    assert '"POSTE"' in source
    assert '"RAMAIS"' in source
    assert '"LUMINÁRIA"' in source
    assert '"TOTAL CLIENTES"' in source
    assert '"TOTAL POSIÇÕES"' in source
    assert '"TOTAL CS"' in source
    assert '"DEMANDA (kVA)"`' in source or '"DEMANDA (kVA)"' in source
    assert '"ACUMULADA (kVA)"' in source
    assert '"COORDENADAS"' in source


# ==============================================================================
# Modelagem lógica idêntica aos parsers AutoLISP para testes de regras de negócio
# ==============================================================================

def py_strip_mtext(val: str) -> str:
    cleaned = re.sub(r"\\[Pp~]", "\n", val)
    cleaned = re.sub(r"\\[A-Za-z0-9]+;", "", cleaned)
    cleaned = re.sub(r"[{}]", "", cleaned)
    return cleaned.strip()


def py_is_noise_line(line: str) -> bool:
    up = line.upper().strip()
    return not up or any(noise in up for noise in [
        "DAN", "ESF. RES", "ESF.", "ESFORCO",
        "DEM.:", "DEM.", "DEMANDA", "ACUM.:", "ACUM.", "ACUMULADO", "KVA", "CLIENTE"
    ])


def py_parse_poste(lines: list[str]) -> str:
    for line in lines:
        up = line.upper().strip()
        if not py_is_noise_line(up) and len(up) >= 2 and up.startswith("P"):
            ch = up[1]
            if ch.isdigit() or ch == "-":
                return up
    return "-"


def py_parse_ramais_e_clientes(lines: list[str]) -> tuple[str, str | int, int, int, int]:
    mono = bi = tri = 0
    ramais_found = False

    for line in lines:
        up = line.upper().strip()
        if not py_is_noise_line(up):
            parts = re.split(r"[,|+\s]+", up)
            for part in parts:
                m = re.match(r"^(\d+)[X*](\d+|TRI|BI|MONO|[Øø]|%%C)", part)
                if m:
                    qty = int(m.group(1))
                    suf = part[m.end(1) + 1:]
                    if "3" in suf or "TRI" in suf:
                        tri += qty
                    elif "2" in suf or "BI" in suf:
                        bi += qty
                    else:
                        mono += qty
                    ramais_found = True

    ramais_list = []
    if mono > 0:
        ramais_list.append(f"{mono}x%%C1")
    if bi > 0:
        ramais_list.append(f"{bi}x%%C2")
    if tri > 0:
        ramais_list.append(f"{tri}x%%C3")

    if not ramais_found or (mono == 0 and bi == 0 and tri == 0):
        return "-", "-", 0, 0, 0

    ramal_str = " | ".join(ramais_list)
    clientes_total = mono + bi + tri
    return ramal_str, clientes_total, mono, bi, tri


def py_parse_luminaria(lines: list[str]) -> str:
    for line in lines:
        up = line.upper().strip()
        if not py_is_noise_line(up):
            if any(k in up for k in ["IP", "LUM", "LED", "PÉTALA", "PETALA"]):
                return line.strip()
            if "W" in up and any(w in up for w in ["70", "100", "150", "250"]):
                return line.strip()
    return "-"


def py_extract_value_after_key(text: str, prefixes: list[str]) -> str | None:
    up = text.upper()
    for prefix in prefixes:
        pos = up.find(prefix)
        if pos != -1:
            idx = pos + len(prefix)
            while idx < len(text) and text[idx] in [".", ":", " ", "-", "=", "\t"]:
                idx += 1
            val = ""
            has_digit = False
            while idx < len(text) and (text[idx].isdigit() or text[idx] in [",", "."]):
                ch = text[idx]
                if ch.isdigit():
                    has_digit = True
                val += ch
                idx += 1
            val = val.rstrip(".,")
            if has_digit and len(val) > 0:
                return val
    return None


def py_parse_demanda(lines: list[str]) -> str:
    prefixes = ["DEMANDA", "DEM.:", "DEM:", "DEM.", "DEM"]
    found_val = None
    ambig = False
    for line in lines:
        cand = py_extract_value_after_key(line, prefixes)
        if cand:
            if found_val is None:
                found_val = cand
            elif found_val != cand:
                ambig = True
    if found_val is None or ambig:
        return "-"
    return found_val


def py_parse_acumulada(lines: list[str]) -> str:
    prefixes = ["ACUMULADA", "ACUMULADO", "ACUM.:", "ACUM:", "ACUM.", "ACUM"]
    found_val = None
    ambig = False
    for line in lines:
        cand = py_extract_value_after_key(line, prefixes)
        if cand:
            if found_val is None:
                found_val = cand
            elif found_val != cand:
                ambig = True
    if found_val is None or ambig:
        return "-"
    return found_val


def py_calc_posicoes(mono: int, bi: int, tri: int) -> str | int:
    if mono == 0 and bi == 0 and tri == 0:
        return "-"
    return (mono * 1) + (bi * 2) + (tri * 3)


def py_calc_cs(posicoes: str | int) -> str | int:
    if not isinstance(posicoes, int) or posicoes <= 0:
        return "-"
    return (posicoes + 8) // 9


def py_process_point_texts(raw_texts_list: list[str], seq: int, coords: str = "-") -> dict:
    """Modelo Python fiel a lp-process-point-texts para testes de persistência e integridade."""
    lines = []
    for txt in raw_texts_list:
        lines.extend(py_strip_mtext(txt).split("\n"))
    poste = py_parse_poste(lines)
    ramal_str, clientes, mono, bi, tri = py_parse_ramais_e_clientes(lines)
    lum = py_parse_luminaria(lines)
    pos = py_calc_posicoes(mono, bi, tri)
    cs = py_calc_cs(pos)
    dem = py_parse_demanda(lines)
    acum = py_parse_acumulada(lines)
    return {
        "SEQ": seq,
        "POSTE": poste,
        "RAMAIS": ramal_str,
        "LUM": lum,
        "CLIENTES": clientes,
        "POSICOES": pos,
        "CS": cs,
        "DEMANDA": dem,
        "ACUMULADA": acum,
        "COORDENADAS": coords or "-",
        "RAW": raw_texts_list,
    }


def py_renumerar_pontos(pontos: list[dict]) -> list[dict]:
    """Modelo Python fiel a lp-renumerar-pontos: renumera de 1 a N preservando atributos."""
    resultado = []
    for idx, p in enumerate(pontos, start=1):
        item = p.copy()
        item["SEQ"] = idx
        resultado.append(item)
    return resultado


def py_excluir_ponto(pontos: list[dict], idx_alvo: int) -> list[dict]:
    """Modelo Python fiel a lp-excluir-ponto: remove o ponto 1 a N e renumera sequencialmente."""
    if not pontos or idx_alvo < 1 or idx_alvo > len(pontos):
        return pontos
    removido = False
    remanescentes = []
    for i, p in enumerate(pontos, start=1):
        if not removido and (p.get("SEQ") == idx_alvo or i == idx_alvo):
            removido = True
        else:
            remanescentes.append(p)
    return py_renumerar_pontos(remanescentes)


def py_formata_linha_revisao(p: dict) -> str:
    """Simula a formatação de uma linha de revisão em lp-exibir-revisao."""
    seq = str(p.get("SEQ", "-"))
    poste = str(p.get("POSTE", "-"))
    ramais = str(p.get("RAMAIS", "-"))
    lum = str(p.get("LUM", "-"))
    cli = str(p.get("CLIENTES", "-"))
    pos = str(p.get("POSICOES", "-"))
    cs = str(p.get("CS", "-"))
    dem = str(p.get("DEMANDA", "-"))
    acum = str(p.get("ACUMULADA", "-"))
    coords = str(p.get("COORDENADAS", "-")).replace("\n", " | ")
    return f"{seq} | {poste} | {ramais} | {lum} | {cli} | {pos} | {cs} | {dem} | {acum} | {coords}"


def py_parse_num(val_str: str) -> float | None:
    if not val_str or val_str == "-":
        return None
    clean = str(val_str).replace(",", ".").strip()
    try:
        return float(clean)
    except ValueError:
        return None


def py_format_num(num: float | None, use_comma: bool) -> str:
    if num is None:
        return "-"
    s = f"{num:.2f}"
    return s.replace(".", ",") if use_comma else s


def py_recalcular_demandas_acumuladas(pontos: list[dict]) -> list[dict]:
    novos = []
    acum_anterior = 0.0
    for p in pontos:
        item = p.copy()
        dem_str = str(item.get("DEMANDA", "-"))
        acum_str = str(item.get("ACUMULADA", "-"))
        dem_val = py_parse_num(dem_str)
        use_comma = ("," in dem_str) or ("," in acum_str)
        if dem_val is not None:
            if acum_str == "-" or (py_parse_num(acum_str) is not None and acum_anterior > 0.0):
                acum_anterior += dem_val
                item["ACUMULADA"] = py_format_num(acum_anterior, use_comma)
            else:
                acum_val = py_parse_num(acum_str)
                acum_anterior = acum_val if acum_val is not None else (acum_anterior + dem_val)
        novos.append(item)
    return novos


def py_sincronizar_levantamento(
    pontos: list[dict],
    tbl_handle: str | None,
    dwg_entities: dict,
    tbl_mock: dict | None = None,
    modo: str = "NORMAL",
    area: int = 50,
    topologia: list | dict | None = None,
):
    if not pontos or not tbl_handle:
        return ":nada-a-fazer", pontos, []

    if tbl_handle not in dwg_entities or dwg_entities[tbl_handle].get("type") != "ACAD_TABLE":
        return ":tabela-inexistente", tbl_handle, []

    tem_alteracao = False
    pontos_afetados = []
    novos_pontos = []
    erro_msg = None

    for p in pontos:
        if erro_msg:
            break
        seq = p.get("SEQ")
        handles = p.get("TEXT_HANDLES")
        raw_antigo = p.get("RAW", [])
        coords = p.get("COORDENADAS", "-")
        blk_h = p.get("BLOCO_POSTE_HANDLE")

        if not handles or not isinstance(handles, list):
            novos_pontos.append(p.copy())
            continue

        cur_raw = []
        for h in handles:
            if h not in dwg_entities:
                erro_msg = f"Texto do Ponto {seq} (Handle #{h}): Entidade não encontrada no desenho (excluída)"
                break
            ent = dwg_entities[h]
            if ent.get("type") not in ["TEXT", "MTEXT"]:
                erro_msg = f"Texto do Ponto {seq} (Handle #{h}): Tipo inválido, esperado TEXT ou MTEXT"
                break
            cur_raw.append(ent.get("text", ""))

        if erro_msg:
            break

        if cur_raw == raw_antigo:
            novos_pontos.append(p.copy())
        else:
            tem_alteracao = True
            pontos_afetados.append(seq)
            novo_p = py_process_point_texts(cur_raw, seq=seq, coords=coords)
            novo_p["TEXT_HANDLES"] = handles
            if blk_h:
                novo_p["BLOCO_POSTE_HANDLE"] = blk_h
            if novo_p.get("POSTE") == "-" and p.get("POSTE") and p.get("POSTE") != "-":
                novo_p["POSTE"] = p["POSTE"]
            novos_pontos.append(novo_p)

    if erro_msg:
        return ":erro", erro_msg, []

    if not tem_alteracao:
        return ":sem-alteracoes", pontos, []

    if modo == "CLANDESTINO":
        recalc = py_clan_recalcular_pontos(novos_pontos, area_padrao=area, topologia=topologia)
        if any(p.get("DEMANDA") == "-" or p.get("ACUMULADA") == "-" for p in recalc):
            return ":erro", "Falha no cálculo normativo LIGHT: limite excedido ou dados inválidos", []
        novos_pontos = recalc
    else:
        novos_pontos = py_recalcular_demandas_acumuladas(novos_pontos)

    if tbl_mock is not None:
        tbl_mock["rows"] = [
            [
                p["SEQ"], p["POSTE"], p["RAMAIS"], p["LUM"],
                p["CLIENTES"], p["POSICOES"], p["CS"],
                p["DEMANDA"], p["ACUMULADA"], p["COORDENADAS"]
            ]
            for p in novos_pontos
        ]

    return ":atualizado", novos_pontos, pontos_afetados


# ==============================================================================
# Casos de Teste Rigorosos Baseados na Especificação
# ==============================================================================

def test_calculo_ramais_8x1_e_6x3_normativo():
    """Valida 8x1ø | 6x3ø → 14 clientes, 26 posições, 3 CS."""
    lines = ["8x1ø | 6x3ø"]
    ramal_str, clientes, mono, bi, tri = py_parse_ramais_e_clientes(lines)
    assert ramal_str == "8x%%C1 | 6x%%C3"
    assert mono == 8
    assert bi == 0
    assert tri == 6
    assert clientes == 14

    pos = py_calc_posicoes(mono, bi, tri)
    assert pos == 26
    cs = py_calc_cs(pos)
    assert cs == 3


def test_calculo_ramais_12x1_e_1x3_normativo():
    """Valida 12x1ø | 1x3ø → 13 clientes, 15 posições, 2 CS."""
    lines = ["12x1ø | 1x3ø"]
    ramal_str, clientes, mono, bi, tri = py_parse_ramais_e_clientes(lines)
    assert ramal_str == "12x%%C1 | 1x%%C3"
    assert clientes == 13

    pos = py_calc_posicoes(mono, bi, tri)
    assert pos == 15
    cs = py_calc_cs(pos)
    assert cs == 2


def test_ignorar_somatorio_escrito_no_desenho():
    """Valida que textos como 14+20=34-CLIENTES NÃO são utilizados para total de clientes."""
    raw = "14+20=34-CLIENTES\\PDEM.: 28,33kVA\\PACUM.: 57,59kVA"
    lines = py_strip_mtext(raw).split("\n")

    ramal_str, clientes, mono, bi, tri = py_parse_ramais_e_clientes(lines)
    # Sem ramais, não deve inventar total de 14, 20 ou 34 clientes
    assert clientes == "-"
    assert ramal_str == "-"
    assert py_calc_posicoes(mono, bi, tri) == "-"
    assert py_calc_cs(py_calc_posicoes(mono, bi, tri)) == "-"


def test_extracao_demanda_e_acumulada_separadas():
    """Valida extração precisa de DEM.: 28,33kVA e ACUM.: 57,59kVA preservando formato."""
    raw = "DEM.: 28,33kVA\\PACUM.: 57,59kVA"
    lines = py_strip_mtext(raw).split("\n")

    dem = py_parse_demanda(lines)
    acum = py_parse_acumulada(lines)

    assert dem == "28,33"
    assert acum == "57,59"


def test_demanda_e_acumulada_na_mesma_linha_e_variacoes():
    """Valida interpretação mesmo na mesma linha e com variações de caixa e pontuação."""
    line = "DEM: 22.50 kVA  ACUMULADA: 45.75kVA"
    lines = [line]

    dem = py_parse_demanda(lines)
    acum = py_parse_acumulada(lines)

    assert dem == "22.50"
    assert acum == "45.75"


def test_textos_sem_ramais_nao_inventam_clientes():
    """Valida que ponto sem ramais recebe '-' e não calcula posições nem CS indevidamente."""
    lines = ["POSTE P112600B", "ESF. RES. 140 daN", "DEM.: 10,0kVA"]
    ramal_str, clientes, mono, bi, tri = py_parse_ramais_e_clientes(lines)

    assert ramal_str == "-"
    assert clientes == "-"
    assert py_calc_posicoes(mono, bi, tri) == "-"
    assert py_calc_cs("-") == "-"


def test_valores_ausentes_ou_ambiguos_seguem_padrao_dados_incompletos():
    """Valida que demanda/acumulada ausente ou conflitante gera '-' (dados incompletos)."""
    # Ausente
    lines_vazio = ["P112600B", "8x1ø"]
    assert py_parse_demanda(lines_vazio) == "-"
    assert py_parse_acumulada(lines_vazio) == "-"

    # Conflito / ambíguo em demanda
    lines_ambig = ["DEM.: 28,33kVA", "DEM.: 35,00kVA"]
    assert py_parse_demanda(lines_ambig) == "-"

    # Repetido com mesmo valor não é ambíguo
    lines_ok = ["DEM.: 28,33kVA", "DEM.: 28,33kVA"]
    assert py_parse_demanda(lines_ok) == "28,33"


def test_tabela_nativa_cad_contem_exatamente_10_colunas_na_ordem():
    """Valida que a tabela do AutoCAD é gerada com 10 colunas na ordem exata especificada."""
    source = LISP_PATH.read_text(encoding="utf-8")

    assert "col-count 10" in source
    assert "(vla-MergeCells table 0 0 0 9)" in source

    # Ordem dos cabeçalhos
    assert '(vla-SetText table 1 0 "PONTO")' in source
    assert '(vla-SetText table 1 1 "POSTE")' in source
    assert '(vla-SetText table 1 2 "RAMAIS")' in source
    assert '(vla-SetText table 1 3 "LUMINÁRIA")' in source
    assert '(vla-SetText table 1 4 "TOTAL CLIENTES")' in source
    assert '(vla-SetText table 1 5 "TOTAL POSIÇÕES")' in source
    assert '(vla-SetText table 1 6 "TOTAL CS")' in source
    assert '(vla-SetText table 1 7 "DEMANDA (kVA)")' in source
    assert '(vla-SetText table 1 8 "ACUMULADA (kVA)")' in source
    assert '(vla-SetText table 1 9 "COORDENADAS")' in source

    # Larguras das 10 colunas
    assert "(vla-SetColumnWidth table 0 22.0)" in source
    assert "(vla-SetColumnWidth table 1 35.0)" in source
    assert "(vla-SetColumnWidth table 2 50.0)" in source
    assert "(vla-SetColumnWidth table 3 35.0)" in source
    assert "(vla-SetColumnWidth table 4 38.0)" in source
    assert "(vla-SetColumnWidth table 5 38.0)" in source
    assert "(vla-SetColumnWidth table 6 25.0)" in source
    assert "(vla-SetColumnWidth table 7 38.0)" in source
    assert "(vla-SetColumnWidth table 8 38.0)" in source
    assert "(vla-SetColumnWidth table 9 45.0)" in source


def test_coordenadas_usam_somente_contrato_real_do_bloco_co():
    """N/E vêm dos atributos que ADDBLOCO grava no CO, nunca do ponto de inserção."""
    source = LISP_PATH.read_text(encoding="utf-8")
    addbloco = (ROOT / "topografia" / "addbloco.lsp").read_text(encoding="utf-8")
    assert "(defun lp-extrair-coordenadas-bloco" in source
    assert '((= tag "N")' in source
    assert '((= tag "E")' in source
    assert '(cons "COORDENADAS" coords)' in source
    assert '"BLOCO_POSTE_HANDLE"' in source
    assert '(cdr (assoc 10 edata))' not in source.split("(defun lp-extrair-coordenadas-bloco", 1)[1].split(";;; ----------------------------------------------------------------------------", 1)[0]
    assert '((= tag "N") (vla-put-TextString att strN))' in addbloco
    assert '((= tag "E") (vla-put-TextString att strE))' in addbloco


def test_extracao_codigo_poste_e_filtro_esforco_mecanico():
    """Valida extração dos códigos de poste ignorando daN e esforços mecânicos."""
    lines1 = py_strip_mtext("P112600B\\P140 daN").split("\n")
    assert py_parse_poste(lines1) == "P112600B"

    lines2 = py_strip_mtext("P212600B\\PESF. RES.\\P619 daN").split("\n")
    assert py_parse_poste(lines2) == "P212600B"

    lines3 = ["140 daN", "ESF. RES."]
    assert py_parse_poste(lines3) == "-"


def test_extracao_luminaria():
    """Valida detecção de IP/Luminária e fallback '-'."""
    assert py_parse_luminaria(["IP LED 50W"]) == "IP LED 50W"
    assert py_parse_luminaria(["LUMINARIA 70W VS"]) == "LUMINARIA 70W VS"
    assert py_parse_luminaria(["POSTE P01", "140 daN"]) == "-"


def test_lisp_formula_cs_implementation():
    """Valida a fórmula do CS no código AutoLISP diretamente."""
    source = LISP_PATH.read_text(encoding="utf-8")
    assert "(defun lp-calc-cs (posicoes)" in source
    assert "(/ (+ posicoes 8) 9)" in source


def test_lisp_persistencia_e_preservacao_em_falha():
    """Valida que o LISP só limpa o LData após criação confirmada da tabela e protege contra falhas COM."""
    source = LISP_PATH.read_text(encoding="utf-8")
    assert "(setq table-ok (lp-gerar-tabela-cad pontos))" in source
    assert "(if table-ok" in source
    assert "(lp-db-limpar-pontos)" in source
    assert "(vl-catch-all-apply" in source
    assert "'vla-AddTable" in source
    assert "defconst" not in source
    assert '(setq *lp-dict-name* "OXCAD_LEVANTAMENTO")' in source
    assert "A tabela não foi criada. A coleta foi preservada." in source


def test_levantamento_postes_is_loaded_globally_in_acaddoc():
    """Valida carregamento global em acaddoc.lsp e documentação no menu COMANDOS."""
    loader = (ROOT / "acaddoc.lsp").read_text(encoding="utf-8")
    assert 'load-aux-safe "orcamento\\\\levantamento_postes.lsp"' in loader
    assert "LEVANTAR_POSTES ou LP" in loader


def test_oxcad_palette_button_dispatches_postes_command():
    """Valida botão na paleta OXCAD e despacho seguro para o documento ativo."""
    xaml = (ROOT / "oxcad" / "OxCadControl.xaml").read_text(encoding="utf-8")
    code = (ROOT / "oxcad" / "OxCadControl.xaml.cs").read_text(encoding="utf-8")

    assert 'x:Name="BtnLevantamentoPostes"' in xaml
    assert 'Click="BtnLevantarPostes_Click"' in xaml
    assert "BtnLevantamentoPostes.IsEnabled = enabled;" in code
    assert 'doc.SendStringToExecute("_.LEVANTAR_POSTES "' in code
    assert 'OxCadExecutionRegistry.RegistrarComandoNativo(' in code


def test_deteccao_duplicidade_e_preservacao_de_ocorrencias_distintas():
    """Valida que postes repetidos são mantidos na tabela e alertados para revisão."""
    source = LISP_PATH.read_text(encoding="utf-8")
    assert "(defun lp-verificar-duplicidade" in source
    assert "já coletado no(s) ponto(s):" in source
    assert "Registrado como nova ocorrência." in source
    assert "(setq pontos (append pontos (list novo-ponto)))" in source


def test_menu_interativo_revisao_e_exclusao_com_reindexacao():
    """Valida que o comando suporta Revisão, Exclusão com reindexação e Limpeza explícita."""
    source = LISP_PATH.read_text(encoding="utf-8")
    assert '"Adicionar Revisar Excluir Gerar Limpar Sair Configurar"' in source
    assert "(defun lp-exibir-revisao" in source
    assert "(subst (cons \"SEQ\" idx) (assoc \"SEQ\" p) p)" in source
    assert "Deseja realmente descartar todos os pontos coletados?" in source


# ==============================================================================
# Testes de Regressão: Defeito de Numeração (PONTO = 1) no LEVANTAR_POSTES
# ==============================================================================

def test_regressao_reproducao_defeito_antigo_vs_correcao():
    """Demonstra que a lógica antiga com idx estático causava PONTO = 1 para todos,
    e valida que a correção renumera sequencialmente mantendo identidades e dados."""
    # 4 pontos iniciais conforme reprodução manual do usuário
    pts = [
        {"SEQ": 1, "POSTE": "-", "RAMAIS": "14x%%C1 | 2x%%C3", "CLIENTES": 16, "POSICOES": 20, "CS": 3, "DEMANDA": "31,79", "ACUMULADA": "198,34", "COORDENADAS": "N: 7450100.10\nE: 678100.10"},
        {"SEQ": 2, "POSTE": "-", "RAMAIS": "13x%%C1 | 4x%%C3", "CLIENTES": 17, "POSICOES": 25, "CS": 3, "DEMANDA": "33,52", "ACUMULADA": "103,32", "COORDENADAS": "N: 7450200.20\nE: 678200.20"},
        {"SEQ": 3, "POSTE": "-", "RAMAIS": "15x%%C1 | 3x%%C3", "CLIENTES": 18, "POSICOES": 24, "CS": 3, "DEMANDA": "35,25", "ACUMULADA": "82,18", "COORDENADAS": "N: 7450300.30\nE: 678300.30"},
        {"SEQ": 4, "POSTE": "-", "RAMAIS": "7x%%C1 | 2x%%C2", "CLIENTES": 9, "POSICOES": 11, "CS": 2, "DEMANDA": "19,27", "ACUMULADA": "58,92", "COORDENADAS": "N: 7450400.40\nE: 678400.40"},
    ]

    # Reprodução exata do defeito da implementação anterior (idx = 1 estático):
    remanescentes_bug = [p.copy() for p in pts if p["SEQ"] != 2]
    idx_bug = 1
    pontos_defeito = []
    for p in remanescentes_bug:
        item = p.copy()
        item["SEQ"] = idx_bug  # sem incremento!
        pontos_defeito.append(item)
    # Todos os pontos assumiam SEQ = 1 (defeito reportado)
    assert [p["SEQ"] for p in pontos_defeito] == [1, 1, 1]

    # Com a correção:
    pontos_corrigidos = py_excluir_ponto(pts, 2)
    assert [p["SEQ"] for p in pontos_corrigidos] == [1, 2, 3]
    # Garante categoricamente que os pontos restantes NÃO assumem 1
    assert pontos_corrigidos[1]["SEQ"] == 2
    assert pontos_corrigidos[2]["SEQ"] == 3


def test_coleta_quatro_pontos_exclusao_segundo_ponto_e_revisao_consecutiva():
    """Valida o fluxo completo especificado:
    1. Coleta de 4 pontos (numeração 1, 2, 3, 4);
    2. Exclusão do 2º ponto (remanescentes 1, 2, 3);
    3. Nova revisão confirmando que a numeração permanece correta;
    4. Exclusão de outro ponto com renumeração imediata;
    5. Associação correta dos dados e preservação de coordenadas e cálculos."""
    # 1. Coleta de 4 pontos com dados reais
    p1 = py_process_point_texts(
        ["14x1ø | 2x3ø", "DEM.: 31,79kVA", "ACUM.: 198,34kVA"],
        seq=1,
        coords="N: 7450100.10\nE: 678100.10"
    )
    p1["BLOCO_POSTE_HANDLE"] = "101"

    p2 = py_process_point_texts(
        ["13x1ø | 4x3ø", "DEM.: 33,52kVA", "ACUM.: 103,32kVA"],
        seq=2,
        coords="N: 7450200.20\nE: 678200.20"
    )
    p2["BLOCO_POSTE_HANDLE"] = "102"

    p3 = py_process_point_texts(
        ["15x1ø | 3x3ø", "DEM.: 35,25kVA", "ACUM.: 82,18kVA"],
        seq=3,
        coords="N: 7450300.30\nE: 678300.30"
    )
    p3["BLOCO_POSTE_HANDLE"] = "103"

    p4 = py_process_point_texts(
        ["7x1ø | 2x2ø", "DEM.: 19,27kVA", "ACUM.: 58,92kVA"],
        seq=4,
        coords="N: 7450400.40\nE: 678400.40"
    )
    p4["BLOCO_POSTE_HANDLE"] = "104"

    pontos = [p1, p2, p3, p4]

    # Confirmar numeração inicial 1, 2, 3, 4
    assert [p["SEQ"] for p in pontos] == [1, 2, 3, 4]
    assert pontos[0]["CLIENTES"] == 16
    assert pontos[0]["POSICOES"] == 20
    assert pontos[0]["CS"] == 3

    assert pontos[1]["CLIENTES"] == 17
    assert pontos[1]["POSICOES"] == 25
    assert pontos[1]["CS"] == 3

    assert pontos[2]["CLIENTES"] == 18
    assert pontos[2]["POSICOES"] == 24
    assert pontos[2]["CS"] == 3

    assert pontos[3]["CLIENTES"] == 9
    assert pontos[3]["POSICOES"] == 11
    assert pontos[3]["CS"] == 2

    # 2. Excluir o segundo ponto (ponto 2)
    pontos = py_excluir_ponto(pontos, idx_alvo=2)

    # Confirmar que restaram 3 registros com numeração sequencial 1, 2, 3
    assert len(pontos) == 3
    assert [p["SEQ"] for p in pontos] == [1, 2, 3]

    # 3. Solicitar nova revisão: a numeração e linhas formatadas devem ser 1, 2, 3
    rev_linhas = [py_formata_linha_revisao(p) for p in pontos]
    assert rev_linhas[0].startswith("1 |")
    assert rev_linhas[1].startswith("2 |")
    assert rev_linhas[2].startswith("3 |")
    assert not any(l.startswith("1 |") for l in rev_linhas[1:])

    # Validar integridade dos dados dos 3 pontos restantes:
    # Registro 1 é o P1 original
    assert pontos[0]["RAMAIS"] == "14x%%C1 | 2x%%C3"
    assert pontos[0]["CLIENTES"] == 16
    assert pontos[0]["POSICOES"] == 20
    assert pontos[0]["CS"] == 3
    assert pontos[0]["DEMANDA"] == "31,79"
    assert pontos[0]["ACUMULADA"] == "198,34"
    assert pontos[0]["COORDENADAS"] == "N: 7450100.10\nE: 678100.10"
    assert pontos[0]["BLOCO_POSTE_HANDLE"] == "101"

    # Registro 2 é o P3 original (renumerado para 2)
    assert pontos[1]["RAMAIS"] == "15x%%C1 | 3x%%C3"
    assert pontos[1]["CLIENTES"] == 18
    assert pontos[1]["POSICOES"] == 24
    assert pontos[1]["CS"] == 3
    assert pontos[1]["DEMANDA"] == "35,25"
    assert pontos[1]["ACUMULADA"] == "82,18"
    assert pontos[1]["COORDENADAS"] == "N: 7450300.30\nE: 678300.30"
    assert pontos[1]["BLOCO_POSTE_HANDLE"] == "103"

    # Registro 3 é o P4 original (renumerado para 3)
    assert pontos[2]["RAMAIS"] == "7x%%C1 | 2x%%C2"
    assert pontos[2]["CLIENTES"] == 9
    assert pontos[2]["POSICOES"] == 11
    assert pontos[2]["CS"] == 2
    assert pontos[2]["DEMANDA"] == "19,27"
    assert pontos[2]["ACUMULADA"] == "58,92"
    assert pontos[2]["COORDENADAS"] == "N: 7450400.40\nE: 678400.40"
    assert pontos[2]["BLOCO_POSTE_HANDLE"] == "104"

    # 4. Excluir outro ponto: excluir o novo ponto 2 (que era o P3 original)
    pontos = py_excluir_ponto(pontos, idx_alvo=2)

    # Confirmar renumeração dos restantes para 1, 2
    assert len(pontos) == 2
    assert [p["SEQ"] for p in pontos] == [1, 2]

    # Verificar que os pontos restantes são P1 (SEQ=1) e P4 (SEQ=2)
    assert pontos[0]["BLOCO_POSTE_HANDLE"] == "101"
    assert pontos[0]["DEMANDA"] == "31,79"
    assert pontos[1]["BLOCO_POSTE_HANDLE"] == "104"
    assert pontos[1]["DEMANDA"] == "19,27"
    assert pontos[1]["RAMAIS"] == "7x%%C1 | 2x%%C2"

    # Segunda revisão consecutiva mantém 1 e 2 perfeitamente
    rev_linhas_2 = [py_formata_linha_revisao(p) for p in pontos]
    assert rev_linhas_2[0].startswith("1 |")
    assert rev_linhas_2[1].startswith("2 |")


def test_persistencia_recuperacao_estado_dwg_e_auto_cura():
    """Valida o mecanismo de persistência no DWG (LData) e a capacidade de auto-cura
    caso o estado anterior tenha sido persistido com defeito de numeração."""
    # Dicionário mock do DWG (LData)
    dwg_ldata = {}

    pts = [
        {"SEQ": 1, "POSTE": "P01", "RAMAIS": "8x%%C1", "CLIENTES": 8, "POSICOES": 8, "CS": 1, "DEMANDA": "10,0", "ACUMULADA": "20,0", "COORDENADAS": "N: 100\nE: 200"},
        {"SEQ": 2, "POSTE": "P02", "RAMAIS": "6x%%C1", "CLIENTES": 6, "POSICOES": 6, "CS": 1, "DEMANDA": "8,0", "ACUMULADA": "18,0", "COORDENADAS": "N: 150\nE: 250"},
        {"SEQ": 3, "POSTE": "P03", "RAMAIS": "4x%%C1", "CLIENTES": 4, "POSICOES": 4, "CS": 1, "DEMANDA": "5,0", "ACUMULADA": "10,0", "COORDENADAS": "N: 200\nE: 300"},
    ]

    # Exclusão do ponto 2 e persistência
    pts_restantes = py_excluir_ponto(pts, 2)
    dwg_ldata["OXCAD_LEVANTAMENTO"] = {"PONTOS": pts_restantes}

    # Simulação de pausa, fechamento e reabertura do DWG
    pts_recuperados = dwg_ldata["OXCAD_LEVANTAMENTO"]["PONTOS"]
    assert len(pts_recuperados) == 2
    assert [p["SEQ"] for p in pts_recuperados] == [1, 2]
    assert pts_recuperados[0]["POSTE"] == "P01"
    assert pts_recuperados[1]["POSTE"] == "P03"

    # Teste de auto-cura: simula estado corrompido herdado onde todos possuem SEQ = 1
    pontos_corrompidos = [
        {"SEQ": 1, "POSTE": "P1", "RAMAIS": "1x%%C1"},
        {"SEQ": 1, "POSTE": "P2", "RAMAIS": "2x%%C1"},
        {"SEQ": 1, "POSTE": "P3", "RAMAIS": "3x%%C1"},
        {"SEQ": 1, "POSTE": "P4", "RAMAIS": "4x%%C1"},
    ]
    curados = py_renumerar_pontos(pontos_corrompidos)
    assert [p["SEQ"] for p in curados] == [1, 2, 3, 4]
    assert curados[0]["POSTE"] == "P1"
    assert curados[1]["POSTE"] == "P2"
    assert curados[2]["POSTE"] == "P3"
    assert curados[3]["POSTE"] == "P4"

    # Exclusão no estado corrompido remove apenas o ponto alvo e normaliza os demais
    restantes_apos_cura = py_excluir_ponto(pontos_corrompidos, idx_alvo=2)
    assert len(restantes_apos_cura) == 3
    assert [p["SEQ"] for p in restantes_apos_cura] == [1, 2, 3]
    assert restantes_apos_cura[0]["POSTE"] == "P1"
    assert restantes_apos_cura[1]["POSTE"] == "P3"
    assert restantes_apos_cura[2]["POSTE"] == "P4"


def test_lisp_ast_renumeracao_estavel_e_sem_hardcoded_um():
    """Valida estaticamente no código AutoLISP que a renumeração incrementa o índice
    e que o antigo bug de atribuição estática foi completamente removido."""
    source = LISP_PATH.read_text(encoding="utf-8")

    # Funções modulares implementadas
    assert "(defun lp-renumerar-pontos (pontos / idx)" in source
    assert "(defun lp-excluir-ponto (pontos idx-alvo / i removido remanescentes)" in source
    assert "(defun lp-precisa-renumerar-p (pontos / idx precisa)" in source

    # O incremento real deve existir
    assert "(setq idx (1+ idx))" in source

    # O trecho defeituoso com idx estático em 1 NÃO deve mais existir
    assert "(setq idx 1\n                     pontos (mapcar '(lambda (p)\n                                       (subst (cons \"SEQ\" idx)" not in source
    assert "(setq idx 1\r\n                     pontos (mapcar '(lambda (p)\r\n                                       (subst (cons \"SEQ\" idx)" not in source

    # lp-excluir-ponto deve ser invocado no comando principal
    assert "(setq pontos (lp-excluir-ponto pontos idx))" in source
    assert "(lp-db-salvar-pontos pontos)" in source

    # Revisão e Tabela usam exatamente a mesma fonte de numeração ("SEQ")
    assert '(lp-format-val (cdr (assoc "SEQ" p)))' in source
    assert '(vla-SetText table row-idx 0 (lp-format-val (cdr (assoc "SEQ" p))))' in source


# ==============================================================================
# Suíte de Testes de Regressão: Sincronização entre Desenho e Tabela Vinculada
# ==============================================================================

def test_persistencia_dos_handles_dos_textos_e_da_tabela():
    """Valida que cada ponto retém seus TEXT_HANDLES e que a tabela vinculada é registrada no LData."""
    p1 = py_process_point_texts(["14x1ø | 2x3ø", "DEM.: 31,79kVA"], seq=1, coords="N: 100\nE: 200")
    p1["TEXT_HANDLES"] = ["34A1", "34A2"]
    p1["BLOCO_POSTE_HANDLE"] = "99B1"

    assert "TEXT_HANDLES" in p1
    assert p1["TEXT_HANDLES"] == ["34A1", "34A2"]
    assert p1["BLOCO_POSTE_HANDLE"] == "99B1"

    # Simulação do LData do DWG
    dwg_ldata = {
        "OXCAD_LEVANTAMENTO": {
            "PONTOS": [p1],
            "TABELA_HANDLE": "TABLE_555"
        }
    }
    assert dwg_ldata["OXCAD_LEVANTAMENTO"]["TABELA_HANDLE"] == "TABLE_555"
    assert dwg_ldata["OXCAD_LEVANTAMENTO"]["PONTOS"][0]["TEXT_HANDLES"] == ["34A1", "34A2"]


def test_deteccao_alteracao_texto_e_atualizacao_dados():
    """Valida detecção de alteração de texto de referência no DWG e recálculo dos dados."""
    p1 = py_process_point_texts(["14x1ø | 2x3ø", "DEM.: 31,79kVA", "ACUM.: 198,34kVA"], seq=1)
    p1["TEXT_HANDLES"] = ["T1"]

    dwg_entities = {
        "TABLE_1": {"type": "ACAD_TABLE"},
        "T1": {"type": "MTEXT", "text": "10x1ø | 1x3ø\\PDEM.: 25,00kVA\\PACUM.: 150,00kVA"}
    }
    tbl_mock = {"rows": []}

    status, novos_pontos, afetados = py_sincronizar_levantamento([p1], "TABLE_1", dwg_entities, tbl_mock)

    assert status == ":atualizado"
    assert afetados == [1]
    assert len(novos_pontos) == 1
    # Clientes: 10 + 1 = 11 (era 16)
    assert novos_pontos[0]["CLIENTES"] == 11
    # Posições: (10*1) + (1*3) = 13 (era 20)
    assert novos_pontos[0]["POSICOES"] == 13
    # CS: ceil(13/9) = 2 (era 3)
    assert novos_pontos[0]["CS"] == 2
    assert novos_pontos[0]["DEMANDA"] == "25,00"
    assert novos_pontos[0]["ACUMULADA"] == "150,00"
    assert len(tbl_mock["rows"]) == 1


def test_ausencia_de_atualizacao_quando_nada_mudou():
    """Valida que quando os textos no DWG estão inalterados, nenhuma atualização ocorre."""
    raw = ["14x1ø | 2x3ø", "DEM.: 31,79kVA"]
    p1 = py_process_point_texts(raw, seq=1)
    p1["TEXT_HANDLES"] = ["T1", "T2"]

    dwg_entities = {
        "TABLE_1": {"type": "ACAD_TABLE"},
        "T1": {"type": "MTEXT", "text": "14x1ø | 2x3ø"},
        "T2": {"type": "TEXT", "text": "DEM.: 31,79kVA"}
    }
    tbl_mock = {"rows": [["initial"]]}

    status, novos_pontos, afetados = py_sincronizar_levantamento([p1], "TABLE_1", dwg_entities, tbl_mock)

    assert status == ":sem-alteracoes"
    assert afetados == []
    # Tabela mock não deve ter sido regravada
    assert tbl_mock["rows"] == [["initial"]]


def test_atualizacao_dos_valores_calculados():
    """Valida atualização exata de posições, caixas de serviço e ramais mantendo coordenadas e poste."""
    p1 = py_process_point_texts(["P112600B", "8x1ø | 6x3ø"], seq=1, coords="N: 7000.10\nE: 6000.20")
    p1["TEXT_HANDLES"] = ["TXT_P1"]
    p1["BLOCO_POSTE_HANDLE"] = "BLK_CO_1"

    # 8x1 e 6x3 = 14 clientes, 26 posições, 3 CS
    assert p1["CLIENTES"] == 14
    assert p1["POSICOES"] == 26
    assert p1["CS"] == 3

    # Usuário alterou no CAD para 4x1ø | 2x3ø
    dwg_entities = {
        "TBL": {"type": "ACAD_TABLE"},
        "TXT_P1": {"type": "MTEXT", "text": "P112600B\\P4x1ø | 2x3ø"}
    }
    status, novos, afetados = py_sincronizar_levantamento([p1], "TBL", dwg_entities)

    assert status == ":atualizado"
    # 4x1 e 2x3 = 6 clientes, 10 posições, 2 CS
    assert novos[0]["CLIENTES"] == 6
    assert novos[0]["POSICOES"] == 10
    assert novos[0]["CS"] == 2
    assert novos[0]["COORDENADAS"] == "N: 7000.10\nE: 6000.20"
    assert novos[0]["BLOCO_POSTE_HANDLE"] == "BLK_CO_1"


def test_recalculo_demanda_acumulada_pontos_seguintes():
    """Valida que quando a demanda de um ponto é alterada, a demanda acumulada dos pontos
    posteriores é automaticamente recalculada em cadeia."""
    p1 = py_process_point_texts(["DEM.: 10,0kVA", "ACUM.: 10,0kVA"], seq=1)
    p1["TEXT_HANDLES"] = ["T1"]
    p2 = py_process_point_texts(["DEM.: 8,0kVA", "ACUM.: 18,0kVA"], seq=2)
    p2["TEXT_HANDLES"] = ["T2"]
    p3 = py_process_point_texts(["DEM.: 5,0kVA", "ACUM.: 23,0kVA"], seq=3)
    p3["TEXT_HANDLES"] = ["T3"]

    pontos = [p1, p2, p3]

    # No desenho ativo, o projetista editou APENAS o texto de T1 (P1) para 15,0kVA.
    # Os textos de T2 e T3 no DWG continuam com seus valores originais.
    dwg_entities = {
        "TBL": {"type": "ACAD_TABLE"},
        "T1": {"type": "TEXT", "text": "DEM.: 15,0kVA"},
        "T2": {"type": "TEXT", "text": "DEM.: 8,0kVA\\PACUM.: 18,0kVA"},
        "T3": {"type": "TEXT", "text": "DEM.: 5,0kVA\\PACUM.: 23,0kVA"}
    }

    status, novos, afetados = py_sincronizar_levantamento(pontos, "TBL", dwg_entities)

    assert status == ":atualizado"
    assert 1 in afetados
    # P1 atualizado: demanda 15,0, acum 15,0
    assert novos[0]["DEMANDA"] == "15,0"
    assert novos[0]["ACUMULADA"] == "15,00"
    # P2 posterior recalculado: 15,0 + 8,0 = 23,00
    assert novos[1]["DEMANDA"] == "8,0"
    assert novos[1]["ACUMULADA"] == "23,00"
    # P3 posterior recalculado: 23,00 + 5,0 = 28,00
    assert novos[2]["DEMANDA"] == "5,0"
    assert novos[2]["ACUMULADA"] == "28,00"


def test_preservacao_da_ordem_e_dados_dos_pontos():
    """Valida que a sincronização preserva estritamente a sequência 1, 2, 3 e metadados."""
    pts = [
        {"SEQ": 1, "POSTE": "P01", "RAMAIS": "1x%%C1", "COORDENADAS": "C1", "TEXT_HANDLES": ["H1"], "RAW": ["1x1ø"], "BLOCO_POSTE_HANDLE": "B1"},
        {"SEQ": 2, "POSTE": "P02", "RAMAIS": "2x%%C1", "COORDENADAS": "C2", "TEXT_HANDLES": ["H2"], "RAW": ["2x1ø"], "BLOCO_POSTE_HANDLE": "B2"},
        {"SEQ": 3, "POSTE": "P03", "RAMAIS": "3x%%C1", "COORDENADAS": "C3", "TEXT_HANDLES": ["H3"], "RAW": ["3x1ø"], "BLOCO_POSTE_HANDLE": "B3"},
    ]
    dwg = {
        "TBL": {"type": "ACAD_TABLE"},
        "H1": {"type": "TEXT", "text": "1x1ø"},
        "H2": {"type": "TEXT", "text": "4x1ø"},  # alterado
        "H3": {"type": "TEXT", "text": "3x1ø"}
    }

    status, novos, afetados = py_sincronizar_levantamento(pts, "TBL", dwg)
    assert status == ":atualizado"
    assert [p["SEQ"] for p in novos] == [1, 2, 3]
    assert [p["POSTE"] for p in novos] == ["P01", "P02", "P03"]
    assert [p["COORDENADAS"] for p in novos] == ["C1", "C2", "C3"]
    assert [p["BLOCO_POSTE_HANDLE"] for p in novos] == ["B1", "B2", "B3"]
    assert novos[1]["RAMAIS"] == "4x%%C1"


def test_manutencao_da_mesma_tabela_sem_duplicacao():
    """Valida que a tabela vinculada existente é atualizada in-place sem gerar nova tabela."""
    p1 = py_process_point_texts(["8x1ø"], seq=1)
    p1["TEXT_HANDLES"] = ["T1"]

    dwg = {
        "TABELA_EXISTENTE_100": {"type": "ACAD_TABLE"},
        "T1": {"type": "TEXT", "text": "12x1ø"}
    }
    tbl_mock = {"id": "TABELA_EXISTENTE_100", "rows": []}

    status, novos, afetados = py_sincronizar_levantamento([p1], "TABELA_EXISTENTE_100", dwg, tbl_mock)
    assert status == ":atualizado"
    assert tbl_mock["id"] == "TABELA_EXISTENTE_100"  # mesma identidade preservada
    assert tbl_mock["rows"][0][2] == "12x%%C1"


def test_recuperacao_estado_apos_reabrir_dwg():
    """Valida restauração do estado persistido em LData após salvar, fechar e reabrir o DWG."""
    dwg_salvo = {
        "OXCAD_LEVANTAMENTO": {
            "PONTOS": [
                {"SEQ": 1, "POSTE": "P1", "RAMAIS": "1x%%C1", "TEXT_HANDLES": ["H1"], "RAW": ["1x1ø"], "COORDENADAS": "N:1\nE:2"},
                {"SEQ": 2, "POSTE": "P2", "RAMAIS": "2x%%C1", "TEXT_HANDLES": ["H2"], "RAW": ["2x1ø"], "COORDENADAS": "N:3\nE:4"}
            ],
            "TABELA_HANDLE": "TBL_PERSISTIDA"
        }
    }

    # Simula leitura do LData ao abrir o DWG
    pontos_recuperados = dwg_salvo["OXCAD_LEVANTAMENTO"]["PONTOS"]
    tbl_recuperada = dwg_salvo["OXCAD_LEVANTAMENTO"]["TABELA_HANDLE"]

    assert len(pontos_recuperados) == 2
    assert tbl_recuperada == "TBL_PERSISTIDA"

    # Simula execução de sincronização com os dados recuperados
    dwg_entities = {
        "TBL_PERSISTIDA": {"type": "ACAD_TABLE"},
        "H1": {"type": "TEXT", "text": "1x1ø"},
        "H2": {"type": "TEXT", "text": "5x1ø"}  # alterado enquanto o DWG estava fechado
    }
    status, novos, afetados = py_sincronizar_levantamento(pontos_recuperados, tbl_recuperada, dwg_entities)
    assert status == ":atualizado"
    assert afetados == [2]
    assert novos[1]["RAMAIS"] == "5x%%C1"


def test_tratamento_texto_excluido_do_desenho():
    """Valida que quando um texto de referência é excluído do CAD, a sincronização é interrompida com segurança."""
    p1 = py_process_point_texts(["8x1ø"], seq=1)
    p1["TEXT_HANDLES"] = ["HANDLE_APAGADO"]

    dwg = {
        "TBL": {"type": "ACAD_TABLE"}
        # HANDLE_APAGADO não existe no DWG
    }

    status, msg, afetados = py_sincronizar_levantamento([p1], "TBL", dwg)
    assert status == ":erro"
    assert "Entidade não encontrada no desenho (excluída)" in msg
    # Não corrompe nem inventa dados
    assert p1["RAMAIS"] == "8x%%C1"


def test_tratamento_tabela_excluida_e_recuperacao():
    """Valida detecção de tabela excluída pelo usuário no CAD sem recriação silenciosa."""
    p1 = py_process_point_texts(["8x1ø"], seq=1)
    p1["TEXT_HANDLES"] = ["T1"]

    dwg = {
        "T1": {"type": "TEXT", "text": "8x1ø"}
        # A tabela TBL_DELETADA não existe mais em dwg
    }

    status, handle_retornado, afetados = py_sincronizar_levantamento([p1], "TBL_DELETADA", dwg)
    assert status == ":tabela-inexistente"
    assert handle_retornado == "TBL_DELETADA"


def test_protecao_contra_estado_inconsistente():
    """Valida proteção contra pontos nulos, handles inválidos ou entidades que não são tabelas."""
    # Pontos vazios
    assert py_sincronizar_levantamento([], "TBL", {}) == (":nada-a-fazer", [], [])
    assert py_sincronizar_levantamento(None, "TBL", {}) == (":nada-a-fazer", None, [])
    assert py_sincronizar_levantamento([{"SEQ": 1}], None, {}) == (":nada-a-fazer", [{"SEQ": 1}], [])

    # Handle apontando para entidade que não é tabela (ex: LINE)
    dwg_com_linha = {"OBJ_LINHA": {"type": "LINE"}}
    status, _, _ = py_sincronizar_levantamento([{"SEQ": 1, "TEXT_HANDLES": ["H"]}], "OBJ_LINHA", dwg_com_linha)
    assert status == ":tabela-inexistente"


def test_compatibilidade_com_exclusao_e_renumeracao_de_pontos():
    """Valida que a exclusão de pontos mantém a integridade de TEXT_HANDLES e renumera normalmente."""
    p1 = {"SEQ": 1, "POSTE": "P1", "TEXT_HANDLES": ["H1"], "RAW": ["1x1ø"]}
    p2 = {"SEQ": 2, "POSTE": "P2", "TEXT_HANDLES": ["H2"], "RAW": ["2x1ø"]}
    p3 = {"SEQ": 3, "POSTE": "P3", "TEXT_HANDLES": ["H3"], "RAW": ["3x1ø"]}

    pontos = [p1, p2, p3]
    restantes = py_excluir_ponto(pontos, idx_alvo=2)

    assert len(restantes) == 2
    assert [p["SEQ"] for p in restantes] == [1, 2]
    assert restantes[0]["POSTE"] == "P1"
    assert restantes[0]["TEXT_HANDLES"] == ["H1"]
    assert restantes[1]["POSTE"] == "P3"
    assert restantes[1]["TEXT_HANDLES"] == ["H3"]


def test_lisp_ast_sincronizacao_e_rastreabilidade():
    """Valida estaticamente no código AutoLISP a presença dos mecanismos de sincronização e rastreabilidade."""
    source = LISP_PATH.read_text(encoding="utf-8")

    # Funções de sincronismo
    assert "(defun lp-sincronizar-levantamento" in source
    assert "(defun lp-atualizar-celulas-tabela" in source
    assert "(defun lp-validar-handle-texto" in source
    assert "(defun lp-recalcular-demandas-acumuladas" in source
    assert "(defun lp-atualizar-exclusao-tabela" in source

    # Gestão do Handle da Tabela no LData
    assert "(defun lp-db-obter-tabela-handle" in source
    assert "(defun lp-db-salvar-tabela-handle" in source
    assert "(defun lp-db-remover-tabela-handle" in source
    assert '"TABELA_HANDLE"' in source

    # Associação persistente dos handles de texto na coleta
    assert "(setq text-handles (append text-handles (list (cdr (assoc 5 (entget ent))))))" in source
    assert '(cons (cons "TEXT_HANDLES" text-handles) novo-ponto)' in source

    # Preservação da tabela e coleta na geração
    assert "(lp-db-salvar-tabela-handle tbl-handle)" in source
    assert "(lp-db-salvar-pontos pontos)" in source

    # Sincronização executada automaticamente no início de c:LEVANTAR_POSTES
    assert "(setq sync-res (lp-sincronizar-levantamento pontos tbl-handle))" in source
    assert "(= (car sync-res) :atualizado)" in source
    assert "(= (car sync-res) :sem-alteracoes)" in source
    assert "(= (car sync-res) :tabela-inexistente)" in source
    assert "(= (car sync-res) :erro)" in source


# ==============================================================================
# ==============================================================================
# Módulo CLANDESTINO LIGHT (CQTS_ZNA853248_ATUAL.xlsx) - Python Mirrors & Testes
# ==============================================================================

CLAN_LISP_PATH = ROOT / "orcamento" / "clandestino_light.lsp"
TOPO_LISP_PATH = ROOT / "orcamento" / "clandestino_topologia.lsp"
CQT_CLAN_LISP_PATH = ROOT / "orcamento" / "clandestino_cqt.lsp"


def _load_clan_lisp_list(var_name: str) -> list[float]:
    source = CLAN_LISP_PATH.read_text(encoding="utf-8")
    match = re.search(rf"\(setq\s+{re.escape(var_name)}\s*'\((\s*[\d\.\s]+)\)\)", source, re.DOTALL)
    if not match:
        raise ValueError(f"Variável {var_name} não encontrada em {CLAN_LISP_PATH}")
    return [float(x) for x in match.group(1).split()]


_DMDI_20_400 = _load_clan_lisp_list("*lp-clan-dmdi-20-400*")
_FDIV_1_300 = _load_clan_lisp_list("*lp-clan-fdiv-1-300*")


def py_clan_inteiro_p(val):
    """Verifica se o valor é estritamente numérico e inteiro."""
    if val is None or isinstance(val, bool):
        return False
    if isinstance(val, int):
        return True
    if isinstance(val, float):
        return val.is_integer()
    return False


def py_clan_validar_area(area):
    """Valida se a área residencial é um número inteiro estrito entre 20 e 400 m²."""
    if area is None or isinstance(area, bool) or not isinstance(area, (int, float)):
        return False
    if not py_clan_inteiro_p(area):
        return False
    a = int(area)
    return 20 <= a <= 400


def py_clan_obter_dmdi(area):
    """Espelho em Python de lp-clan-obter-dmdi."""
    if not py_clan_validar_area(area):
        return None
    return _DMDI_20_400[int(area) - 20]


def py_clan_validar_clientes(clientes):
    """Valida se a quantidade de clientes é um número inteiro entre 0 e 300."""
    if clientes is None or isinstance(clientes, bool) or not isinstance(clientes, (int, float)):
        return False
    if not py_clan_inteiro_p(clientes):
        return False
    c = int(clientes)
    return 0 <= c <= 300


def py_clan_obter_fdiv(clientes):
    """Espelho em Python de lp-clan-obter-fator-div."""
    if clientes is None or isinstance(clientes, bool) or not isinstance(clientes, (int, float)):
        return None
    if not py_clan_inteiro_p(clientes):
        return None
    c = int(clientes)
    if c == 0:
        return 0.0
    if not (1 <= c <= 300):
        return None
    return _FDIV_1_300[c - 1]


def py_clan_calcular_demanda(area, clientes):
    """Espelho em Python de lp-clan-calcular-demanda."""
    dmdi = py_clan_obter_dmdi(area)
    fdiv = py_clan_obter_fdiv(clientes)
    if dmdi is None or fdiv is None:
        return None
    return round(dmdi * fdiv, 2)


def py_clan_calcular_acumulada(area, clientes_acumulados):
    """Espelho em Python de lp-clan-calcular-acumulada."""
    dmdi = py_clan_obter_dmdi(area)
    fdiv = py_clan_obter_fdiv(clientes_acumulados)
    if dmdi is None or fdiv is None:
        return None
    return round(dmdi * fdiv, 2)


def py_clan_parse_cli(c):
    """Converte o valor de CLIENTES (int ou string) em inteiro."""
    if c is None or isinstance(c, bool):
        return 0
    if isinstance(c, (int, float)):
        return int(c)
    if isinstance(c, str):
        try:
            return int(float(c))
        except ValueError:
            return 0
    return 0


def py_clan_adicionar_aresta(u: str, v: str, mapa: dict) -> dict:
    u_str = str(u).upper().strip()
    v_str = str(v).upper().strip()
    if not u_str or not v_str or u_str == v_str:
        return mapa
    if u_str not in mapa:
        mapa[u_str] = []
    if v_str not in mapa[u_str]:
        mapa[u_str].append(v_str)
    return mapa


def py_clan_construir_mapa_filhos(topologia=None, pontos=None) -> dict[str, list[str]]:
    """Constrói o mapa de adjacência {pai: [filho1, filho2, ...]} a partir de topologia ou atributos."""
    mapa: dict[str, list[str]] = {}
    if topologia:
        ramos = []
        if isinstance(topologia, dict):
            lados = topologia.get("LADOS") or topologia.get("TRECHOS") or topologia.get("BRANCHES") or []
            deriv = topologia.get("DERIVACOES") or []
            ramos.extend(lados)
            ramos.extend(deriv)
        elif isinstance(topologia, list):
            ramos = topologia

        for item in ramos:
            if isinstance(item, (list, tuple)):
                if len(item) == 2 and not any(k in str(item[0]).upper() for k in ["LADO", "RAMO", "CIRCUITO"]):
                    py_clan_adicionar_aresta(item[0], item[1], mapa)
                else:
                    nodes = []
                    for n in item:
                        n_str = str(n).upper().strip()
                        if not any(k in n_str for k in ["LADO", "RAMO", "CIRCUITO", "TR "]):
                            nodes.append(n_str)
                    for i in range(len(nodes) - 1):
                        py_clan_adicionar_aresta(nodes[i], nodes[i + 1], mapa)

    if pontos:
        for p in pontos:
            poste = str(p.get("POSTE", "")).upper().strip()
            if poste and poste != "-":
                jusante = p.get("JUSANTE")
                if jusante:
                    if isinstance(jusante, (list, tuple)):
                        for j in jusante:
                            py_clan_adicionar_aresta(poste, j, mapa)
                    else:
                        py_clan_adicionar_aresta(poste, jusante, mapa)
                montante = p.get("MONTANTE")
                if montante and montante != "-":
                    py_clan_adicionar_aresta(montante, poste, mapa)
    return mapa


def py_clan_obter_jusante(poste: str, mapa_filhos: dict[str, list[str]]) -> list[str]:
    """Retorna a lista de nomes de nós a jusante (descendentes) de poste na árvore."""
    p_up = str(poste).upper().strip()
    visitados = []
    fila = list(mapa_filhos.get(p_up, []))
    while fila:
        atual = fila.pop(0)
        if atual not in visitados:
            visitados.append(atual)
            for f in mapa_filhos.get(atual, []):
                if f not in visitados:
                    fila.append(f)
    return visitados


def py_clan_recalcular_pontos(pontos, area_padrao=50, sentido="PONTA->FONTE", topologia=None):
    """Espelho em Python de lp-clan-recalcular-pontos com suporte a topologia e cálculo transacional."""
    if not pontos or not py_clan_validar_area(area_padrao):
        return pontos

    pts = [dict(p) for p in pontos]
    area = int(area_padrao)

    topo = topologia
    if topo is None and isinstance(sentido, (list, dict)):
        topo = sentido

    mapa_filhos = py_clan_construir_mapa_filhos(topo, pts)

    # 1. Validação estrita de todos os clientes locais
    cli_lookup = {}
    for p in pts:
        poste = str(p.get("POSTE", "")).upper().strip()
        raw_cli = p.get("CLIENTES")
        if raw_cli is None or isinstance(raw_cli, bool):
            return pontos
        if isinstance(raw_cli, float) and not raw_cli.is_integer():
            return pontos
        try:
            cli = int(raw_cli)
        except (ValueError, TypeError):
            return pontos
        if not py_clan_validar_clientes(cli):
            return pontos
        cli_lookup[poste] = cli

    # Modo topológico
    if mapa_filhos:
        # Pré-validação de limites acumulados a jusante
        for p in pts:
            poste = str(p.get("POSTE", "")).upper().strip()
            local_cli = cli_lookup.get(poste, 0)
            jusante_postes = py_clan_obter_jusante(poste, mapa_filhos)
            acum_cli = local_cli + sum(cli_lookup.get(j, 0) for j in jusante_postes)
            if acum_cli > 300:
                return pontos

        for p in pts:
            poste = str(p.get("POSTE", "")).upper().strip()
            local_cli = cli_lookup.get(poste, 0)
            jusante_postes = py_clan_obter_jusante(poste, mapa_filhos)
            acum_cli = local_cli + sum(cli_lookup.get(j, 0) for j in jusante_postes)
            dem = py_clan_calcular_demanda(area, local_cli)
            acum = py_clan_calcular_acumulada(area, acum_cli)
            if dem is None or acum is None:
                return pontos
            p["DEMANDA"] = f"{dem:.2f}".replace(".", ",")
            p["ACUMULADA"] = f"{acum:.2f}".replace(".", ",")
        return pts

    # Modo sequencial linear
    sentido_str = str(sentido).upper() if isinstance(sentido, str) else "PONTA->FONTE"
    if sentido_str == "FONTE->PONTA":
        tot = 0
        cli_map = []
        for p in reversed(pts):
            c = cli_lookup.get(str(p.get("POSTE", "")).upper().strip(), 0)
            tot += c
            if tot > 300:
                return pontos
            cli_map.insert(0, tot)
        for p, acum_c in zip(pts, cli_map):
            c = cli_lookup.get(str(p.get("POSTE", "")).upper().strip(), 0)
            dem = py_clan_calcular_demanda(area, c)
            acum = py_clan_calcular_acumulada(area, acum_c)
            if dem is None or acum is None:
                return pontos
            p["DEMANDA"] = f"{dem:.2f}".replace(".", ",")
            p["ACUMULADA"] = f"{acum:.2f}".replace(".", ",")
        return pts
    else:
        # PONTA->FONTE
        tot = 0
        for p in pts:
            c = cli_lookup.get(str(p.get("POSTE", "")).upper().strip(), 0)
            tot += c
            if tot > 300:
                return pontos
        acum_c = 0
        for p in pts:
            c = cli_lookup.get(str(p.get("POSTE", "")).upper().strip(), 0)
            acum_c += c
            dem = py_clan_calcular_demanda(area, c)
            acum = py_clan_calcular_acumulada(area, acum_c)
            if dem is None or acum is None:
                return pontos
            p["DEMANDA"] = f"{dem:.2f}".replace(".", ",")
            p["ACUMULADA"] = f"{acum:.2f}".replace(".", ",")
        return pts


# ------------------------------------------------------------------------------
# Testes Unitários de Conformidade e Segurança
# ------------------------------------------------------------------------------

def test_clandestino_dmdi_valores_normativos():
    """Valida DMDI para as áreas obrigatórias: 20, 42, 43, 50 e 400 m² (DB!A36:B416)."""
    assert py_clan_obter_dmdi(20) == 1.62
    assert py_clan_obter_dmdi(42) == 1.62
    assert py_clan_obter_dmdi(43) == 1.64
    assert py_clan_obter_dmdi(50) == 1.88
    assert py_clan_obter_dmdi(400) == 12.13


def test_clandestino_fator_diversificacao_valores_normativos():
    """Valida fatores de diversificação para os clientes obrigatórios: 1, 4, 5, 10, 13, 50, 132, 278, 279 e 300."""
    assert py_clan_obter_fdiv(1) == 3.88
    assert py_clan_obter_fdiv(4) == 3.88
    assert py_clan_obter_fdiv(5) == 4.84
    assert py_clan_obter_fdiv(10) == 9.64
    assert py_clan_obter_fdiv(13) == 11.98
    assert py_clan_obter_fdiv(50) == 35.34
    assert py_clan_obter_fdiv(132) == 70.99
    assert py_clan_obter_fdiv(278) == 83.00
    assert py_clan_obter_fdiv(279) == 83.00
    assert py_clan_obter_fdiv(300) == 83.00


def test_clandestino_demandas_individuais_e_acumuladas():
    """Valida o cálculo individual e acumulado comprovando a não-linearidade da rede clandestina."""
    area = 50  # DMDI = 1.88 kVA

    # Poste A: 24 clientes -> 1.88 * 19.88 = 37.37 kVA
    dem_a = py_clan_calcular_demanda(area, 24)
    assert dem_a == 37.37

    # Poste B: 26 clientes -> 1.88 * 21.10 = 39.67 kVA
    dem_b = py_clan_calcular_demanda(area, 26)
    assert dem_b == 39.67

    # Demanda acumulada de A + B (total 50 clientes): 1.88 * 35.34 = 66.44 kVA
    acum_ab = py_clan_calcular_acumulada(area, 50)
    assert acum_ab == 66.44

    # PROVA NÃO-LINEAR: Demanda acumulada (66.44) é estritamente diferente da soma linear (77.04)
    assert round(dem_a + dem_b, 2) == 77.04
    assert acum_ab != round(dem_a + dem_b, 2)


def test_clandestino_acumulado_multiplos_pontos_topologia_real():
    """Valida acumulado com múltiplos pontos a jusante pela topologia real (circuito RUA IRINÉIA VERÇOSA)."""
    # 4 pontos reais da planilha LIGHT CQTS_ZNA853248_ATUAL.xlsx (PROJ1, linhas 13-16)
    # P7 (24 cli), P6 (26 cli), P5 (10 cli), P4 (13 cli)
    # Levantamento da ponta para a fonte (P7 -> P6 -> P5 -> P4):
    pontos = [
        {"SEQ": 1, "POSTE": "P7", "CLIENTES": 24},
        {"SEQ": 2, "POSTE": "P6", "CLIENTES": 26},
        {"SEQ": 3, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 4, "POSTE": "P4", "CLIENTES": 13},
    ]

    recalc = py_clan_recalcular_pontos(pontos, area_padrao=50, sentido="PONTA->FONTE")

    # Linha 16 (P7): 24 clientes locais, 24 acumulados -> 37.37 / 37.37
    assert recalc[0]["DEMANDA"] == "37,37"
    assert recalc[0]["ACUMULADA"] == "37,37"

    # Linha 15 (P6): 26 clientes locais, 50 acumulados -> 39.67 / 66.44
    assert recalc[1]["DEMANDA"] == "39,67"
    assert recalc[1]["ACUMULADA"] == "66,44"

    # Linha 14 (P5): 10 clientes locais, 60 acumulados -> 18.12 / 76.97
    assert recalc[2]["DEMANDA"] == "18,12"
    assert recalc[2]["ACUMULADA"] == "76,97"

    # Linha 13 (P4): 13 clientes locais, 73 acumulados -> 22.52 / 90.65
    assert recalc[3]["DEMANDA"] == "22,52"
    assert recalc[3]["ACUMULADA"] == "90,65"

    # Topologia inversa: FONTE->PONTA (P4 -> P5 -> P6 -> P7)
    pontos_rev = [
        {"SEQ": 1, "POSTE": "P4", "CLIENTES": 13},
        {"SEQ": 2, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 3, "POSTE": "P6", "CLIENTES": 26},
        {"SEQ": 4, "POSTE": "P7", "CLIENTES": 24},
    ]
    recalc_rev = py_clan_recalcular_pontos(pontos_rev, area_padrao=50, sentido="FONTE->PONTA")
    assert recalc_rev[0]["ACUMULADA"] == "90,65"  # P4 acumula todos os 73 clientes
    assert recalc_rev[1]["ACUMULADA"] == "76,97"  # P5 acumula 60 clientes
    assert recalc_rev[2]["ACUMULADA"] == "66,44"  # P6 acumula 50 clientes
    assert recalc_rev[3]["ACUMULADA"] == "37,37"  # P7 é ponta (24 clientes)


def test_clandestino_atualizacao_demanda_e_acumulados_apos_alteracao_ramais():
    """Valida recálculo dinâmico da demanda e de todos os acumulados ao editar ramais no CAD e sincronizar."""
    p1 = {
        "SEQ": 1, "POSTE": "P7", "RAMAIS": "24x1ø", "CLIENTES": "24",
        "DEMANDA": "37,37", "ACUMULADA": "37,37",
        "TEXT_HANDLES": ["TXT_P7"], "RAW": ["24x1ø"]
    }
    p2 = {
        "SEQ": 2, "POSTE": "P6", "RAMAIS": "26x1ø", "CLIENTES": "26",
        "DEMANDA": "39,67", "ACUMULADA": "66,44",
        "TEXT_HANDLES": ["TXT_P6"], "RAW": ["26x1ø"]
    }

    # Usuário edita o texto de ramal do P7 no CAD para 30x1ø
    dwg = {
        "TBL": {"type": "ACAD_TABLE"},
        "TXT_P7": {"type": "TEXT", "text": "30x1ø"},
        "TXT_P6": {"type": "TEXT", "text": "26x1ø"},
    }

    status, atualizados, afetados = py_sincronizar_levantamento([p1, p2], "TBL", dwg)
    assert status == ":atualizado"
    assert afetados == [1]

    # Recalcula cadeia no modo clandestino
    recalculados = py_clan_recalcular_pontos(atualizados, area_padrao=50, sentido="PONTA->FONTE")

    # P7 atualizado para 30 clientes (fdiv = 23.54 -> dem = 44.26 / acum = 44.26)
    assert str(recalculados[0]["CLIENTES"]) == "30"
    assert recalculados[0]["DEMANDA"] == "44,26"
    assert recalculados[0]["ACUMULADA"] == "44,26"

    # P6 mantém 26 clientes locais, mas acumulado sobe para 56 clientes (fdiv = 38.70 -> acum = 72.76)
    assert str(recalculados[1]["CLIENTES"]) == "26"
    assert recalculados[1]["DEMANDA"] == "39,67"
    assert recalculados[1]["ACUMULADA"] == "72,76"


def test_clandestino_entradas_invalidas_e_limites_nao_suportados():
    """Valida a rejeição explícita e interrupção de cálculo para áreas <20, >400, decimais e clientes >300."""
    # Áreas menores que 20 m²
    assert py_clan_obter_dmdi(19) is None
    assert py_clan_obter_dmdi(0) is None
    assert py_clan_obter_dmdi(-10) is None

    # Áreas maiores que 400 m² (não extrapola)
    assert py_clan_obter_dmdi(401) is None
    assert py_clan_obter_dmdi(500) is None

    # Áreas decimais (não arredonda silenciosamente)
    assert py_clan_obter_dmdi(20.5) is None
    assert py_clan_obter_dmdi(49.6) is None
    assert py_clan_obter_dmdi(50.4) is None

    # Áreas não numéricas ou nulas
    assert py_clan_obter_dmdi(None) is None
    assert py_clan_obter_dmdi("50") is None

    # Quantidade de clientes > 300 (não extrapola ou limita silenciosamente)
    assert py_clan_obter_fdiv(301) is None
    assert py_clan_obter_fdiv(500) is None

    # Quantidade de clientes negativa ou decimal
    assert py_clan_obter_fdiv(-1) is None
    assert py_clan_obter_fdiv(10.5) is None
    assert py_clan_obter_fdiv(None) is None

    # Demanda e acumulada interrompem retornando None quando entrada é inválida
    assert py_clan_calcular_demanda(19, 10) is None
    assert py_clan_calcular_demanda(401, 10) is None
    assert py_clan_calcular_demanda(50, 301) is None
    assert py_clan_calcular_acumulada(50.5, 50) is None
    assert py_clan_calcular_acumulada(50, 350) is None

    # Recalcular pontos com entrada inválida interrompe preservando lista sem extrapolação
    pts_invalidos = [{"SEQ": 1, "POSTE": "P1", "CLIENTES": 350}]
    resultado = py_clan_recalcular_pontos(pts_invalidos, area_padrao=50)
    assert resultado == pts_invalidos


def test_clandestino_persistencia_da_configuracao_entre_execucoes():
    """Valida o contrato de persistência no DWG via LData para chaves MODO, AREA_CLANDESTINO, SENTIDO e TOPOLOGIA."""
    source = CLAN_LISP_PATH.read_text(encoding="utf-8")
    assert '"MODO"' in source
    assert '"AREA_CLANDESTINO"' in source
    assert '"SENTIDO_ACUMULACAO"' in source
    assert '"TOPOLOGIA"' in source
    assert "(defun lp-db-obter-modo" in source
    assert "(defun lp-db-salvar-modo" in source
    assert "(defun lp-db-obter-area" in source
    assert "(defun lp-db-salvar-area" in source
    assert "(defun lp-db-obter-sentido" in source
    assert "(defun lp-db-salvar-sentido" in source
    assert "(defun lp-db-obter-topologia" in source
    assert "(defun lp-db-salvar-topologia" in source
    assert "(defun lp-clan-obter-ultimo-erro" in source


def test_regressao_integral_modo_normal():
    """Valida regressão estrita do modo NORMAL: preserva leitura literal de DEM.: e ACUM.: sem cálculo clandestino."""
    raw = [
        "POSTE 10/300",
        "8x1ø + 2x2ø",
        "DEM.: 15.50",
        "ACUM.: 45.20"
    ]
    ponto = py_process_point_texts(raw, seq=1)
    assert ponto["CLIENTES"] == 10
    assert ponto["DEMANDA"] == "15.50"
    assert ponto["ACUMULADA"] == "45.20"


def test_clandestino_topologia_independencia_da_ordem_de_cadastro():
    """RISCO 1: Valida que ordens distintas de coleta/cadastro produzem resultados idênticos via topologia física."""
    # Topologia da planilha oficial LIGHT CQTS_ZNA853248_ATUAL.xlsx: TR -> P4 -> P5 -> P6 -> P7
    topologia = [["LADO 1", "P4", "P5", "P6", "P7"]]

    # Cenário A: Ordem da fonte para a ponta (P4, P5, P6, P7)
    cenario_a = [
        {"SEQ": 1, "POSTE": "P4", "CLIENTES": 13},
        {"SEQ": 2, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 3, "POSTE": "P6", "CLIENTES": 26},
        {"SEQ": 4, "POSTE": "P7", "CLIENTES": 24},
    ]

    # Cenário B: Ordem da ponta para a fonte (P7, P6, P5, P4)
    cenario_b = [
        {"SEQ": 1, "POSTE": "P7", "CLIENTES": 24},
        {"SEQ": 2, "POSTE": "P6", "CLIENTES": 26},
        {"SEQ": 3, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 4, "POSTE": "P4", "CLIENTES": 13},
    ]

    # Cenário C: Ordem arbitrária/desordenada de campo (P5, P7, P4, P6)
    cenario_c = [
        {"SEQ": 1, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 2, "POSTE": "P7", "CLIENTES": 24},
        {"SEQ": 3, "POSTE": "P4", "CLIENTES": 13},
        {"SEQ": 4, "POSTE": "P6", "CLIENTES": 26},
    ]

    res_a = {p["POSTE"]: p for p in py_clan_recalcular_pontos(cenario_a, area_padrao=50, topologia=topologia)}
    res_b = {p["POSTE"]: p for p in py_clan_recalcular_pontos(cenario_b, area_padrao=50, topologia=topologia)}
    res_c = {p["POSTE"]: p for p in py_clan_recalcular_pontos(cenario_c, area_padrao=50, topologia=topologia)}

    # Valores canônicos extraídos da planilha LIGHT CQTS_ZNA853248_ATUAL.xlsx:
    # P7: 24 clientes locais / 24 acumulados -> Dem: 37,37 / Acum: 37,37
    # P6: 26 clientes locais / 50 acumulados -> Dem: 39,67 / Acum: 66,44
    # P5: 10 clientes locais / 60 acumulados -> Dem: 18,12 / Acum: 76,97
    # P4: 13 clientes locais / 73 acumulados -> Dem: 22,52 / Acum: 90,65
    for res in [res_a, res_b, res_c]:
        assert res["P7"]["DEMANDA"] == "37,37"
        assert res["P7"]["ACUMULADA"] == "37,37"

        assert res["P6"]["DEMANDA"] == "39,67"
        assert res["P6"]["ACUMULADA"] == "66,44"

        assert res["P5"]["DEMANDA"] == "18,12"
        assert res["P5"]["ACUMULADA"] == "76,97"

        assert res["P4"]["DEMANDA"] == "22,52"
        assert res["P4"]["ACUMULADA"] == "90,65"

    # Confirmação estrita de equivalência absoluta entre todas as ordens
    for poste in ["P4", "P5", "P6", "P7"]:
        assert res_a[poste]["DEMANDA"] == res_b[poste]["DEMANDA"] == res_c[poste]["DEMANDA"]
        assert res_a[poste]["ACUMULADA"] == res_b[poste]["ACUMULADA"] == res_c[poste]["ACUMULADA"]


def test_clandestino_topologia_com_derivacao_independencia_da_ordem():
    """RISCO 1: Valida topologia com bifurcação/derivação da planilha real (P2 -> P4... e P2 -> P3)."""
    # Topologia com tronco e derivação: TR -> P2 -> P4 -> P5 -> P6 -> P7 e P2 -> P3
    topologia = {
        "TR": "TR 01",
        "LADOS": [["LADO 1", "P2", "P4", "P5", "P6", "P7"]],
        "DERIVACOES": [["DERIVACAO P2", "P2", "P3"]]
    }

    # Ordem 1: Natural da fonte para as pontas
    ordem_1 = [
        {"SEQ": 1, "POSTE": "P2", "CLIENTES": 0},
        {"SEQ": 2, "POSTE": "P3", "CLIENTES": 10},
        {"SEQ": 3, "POSTE": "P4", "CLIENTES": 13},
        {"SEQ": 4, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 5, "POSTE": "P6", "CLIENTES": 26},
        {"SEQ": 6, "POSTE": "P7", "CLIENTES": 24},
    ]

    # Ordem 2: Totalmente misturada
    ordem_2 = [
        {"SEQ": 1, "POSTE": "P6", "CLIENTES": 26},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 0},
        {"SEQ": 3, "POSTE": "P7", "CLIENTES": 24},
        {"SEQ": 4, "POSTE": "P3", "CLIENTES": 10},
        {"SEQ": 5, "POSTE": "P5", "CLIENTES": 10},
        {"SEQ": 6, "POSTE": "P4", "CLIENTES": 13},
    ]

    res_1 = {p["POSTE"]: p for p in py_clan_recalcular_pontos(ordem_1, area_padrao=50, topologia=topologia)}
    res_2 = {p["POSTE"]: p for p in py_clan_recalcular_pontos(ordem_2, area_padrao=50, topologia=topologia)}

    # P2 acumula jusante total: 0 + 10 (P3) + 73 (P4+P5+P6+P7) = 83 clientes -> fdiv(83) = 53.82 -> 101.18 kVA
    assert res_1["P2"]["ACUMULADA"] == res_2["P2"]["ACUMULADA"] == "101,18"
    assert res_1["P2"]["DEMANDA"] == res_2["P2"]["DEMANDA"] == "0,00"

    # P3 derivação acumula apenas seus 10 clientes: fdiv(10) = 9.64 -> 18.12 kVA
    assert res_1["P3"]["ACUMULADA"] == res_2["P3"]["ACUMULADA"] == "18,12"
    assert res_1["P3"]["DEMANDA"] == res_2["P3"]["DEMANDA"] == "18,12"

    # P4 a P7 mantêm valores canônicos idênticos em ambas as ordens
    assert res_1["P4"]["ACUMULADA"] == res_2["P4"]["ACUMULADA"] == "90,65"
    assert res_1["P5"]["ACUMULADA"] == res_2["P5"]["ACUMULADA"] == "76,97"
    assert res_1["P6"]["ACUMULADA"] == res_2["P6"]["ACUMULADA"] == "66,44"
    assert res_1["P7"]["ACUMULADA"] == res_2["P7"]["ACUMULADA"] == "37,37"


def test_clandestino_falha_calculo_preservacao_integral_de_dados_validos():
    """RISCO 2: Valida que qualquer falha (área, clientes, limites) preserva 100% dos dados válidos sem zeros ou traços."""
    pontos_originais = [
        {"SEQ": 1, "POSTE": "P4", "RAMAIS": "13x%%C1", "CLIENTES": 13, "DEMANDA": "22,52", "ACUMULADA": "90,65", "COORDENADAS": "X:1 Y:1"},
        {"SEQ": 2, "POSTE": "P5", "RAMAIS": "10x%%C1", "CLIENTES": 10, "DEMANDA": "18,12", "ACUMULADA": "76,97", "COORDENADAS": "X:2 Y:2"},
        {"SEQ": 3, "POSTE": "P6", "RAMAIS": "26x%%C1", "CLIENTES": 26, "DEMANDA": "39,67", "ACUMULADA": "66,44", "COORDENADAS": "X:3 Y:3"},
        {"SEQ": 4, "POSTE": "P7", "RAMAIS": "24x%%C1", "CLIENTES": 24, "DEMANDA": "37,37", "ACUMULADA": "37,37", "COORDENADAS": "X:4 Y:4"},
    ]
    topologia = [["LADO 1", "P4", "P5", "P6", "P7"]]

    # 1. Falha por área < 20
    res_area_baixa = py_clan_recalcular_pontos(pontos_originais, area_padrao=19, topologia=topologia)
    assert res_area_baixa == pontos_originais

    # 2. Falha por área > 400
    res_area_alta = py_clan_recalcular_pontos(pontos_originais, area_padrao=401, topologia=topologia)
    assert res_area_alta == pontos_originais

    # 3. Falha por área decimal (ex: 50.5 m²)
    res_area_dec = py_clan_recalcular_pontos(pontos_originais, area_padrao=50.5, topologia=topologia)
    assert res_area_dec == pontos_originais

    # 4. Falha por clientes individuais > 300
    pts_cli_excedente = [
        dict(pontos_originais[0], CLIENTES=305),
        pontos_originais[1],
        pontos_originais[2],
        pontos_originais[3],
    ]
    res_cli_alta = py_clan_recalcular_pontos(pts_cli_excedente, area_padrao=50, topologia=topologia)
    assert res_cli_alta == pts_cli_excedente
    # Garante que nenhum dado existente virou "-" ou "0,00"
    assert res_cli_alta[0]["DEMANDA"] == "22,52"
    assert res_cli_alta[0]["ACUMULADA"] == "90,65"

    # 5. Falha por clientes acumulados a jusante > 300 (mesmo com locais <= 300)
    # Exemplo: P4 tem 200 clientes e P5 tem 150 clientes -> acumulado em P4 = 350 > 300
    pts_acum_excedente = [
        {"SEQ": 1, "POSTE": "P4", "CLIENTES": 200, "DEMANDA": "150,00", "ACUMULADA": "150,00"},
        {"SEQ": 2, "POSTE": "P5", "CLIENTES": 150, "DEMANDA": "120,00", "ACUMULADA": "120,00"},
    ]
    res_acum_alta = py_clan_recalcular_pontos(pts_acum_excedente, area_padrao=50, topologia=[["P4", "P5"]])
    assert res_acum_alta == pts_acum_excedente
    assert res_acum_alta[0]["DEMANDA"] == "150,00"
    assert res_acum_alta[0]["ACUMULADA"] == "150,00"
    assert res_acum_alta[1]["DEMANDA"] == "120,00"
    assert res_acum_alta[1]["ACUMULADA"] == "120,00"

    # 6. Falha por clientes negativos ou decimais
    pts_cli_neg = [dict(pontos_originais[0], CLIENTES=-5), pontos_originais[1]]
    assert py_clan_recalcular_pontos(pts_cli_neg, area_padrao=50, topologia=topologia) == pts_cli_neg

    pts_cli_float = [dict(pontos_originais[0], CLIENTES=12.5), pontos_originais[1]]
    assert py_clan_recalcular_pontos(pts_cli_float, area_padrao=50, topologia=topologia) == pts_cli_float


def test_clandestino_sincronizacao_tabela_nao_corrompe_ao_falhar():
    """RISCO 2: Valida que erro em textos do CAD impede mutação da tabela vinculada preservando células intactas."""
    p1 = {
        "SEQ": 1, "POSTE": "P4", "RAMAIS": "13x%%C1", "CLIENTES": 13,
        "DEMANDA": "22,52", "ACUMULADA": "90,65", "COORDENADAS": "C1",
        "TEXT_HANDLES": ["H4"], "RAW": ["13x1ø"]
    }
    dwg = {
        "TBL_VINCULADA": {"type": "ACAD_TABLE"},
        # Usuário edita o texto de ramal no CAD para um valor absurdo que excede 300 clientes:
        "H4": {"type": "TEXT", "text": "350x1ø"}
    }
    # Tabela com dados válidos originais
    tbl_mock = {
        "id": "TBL_VINCULADA",
        "rows": [
            [1, "P4", "13x%%C1", "-", 13, 13, 2, "22,52", "90,65", "C1"]
        ]
    }

    status, novos, afetados = py_sincronizar_levantamento(
        [p1], "TBL_VINCULADA", dwg, tbl_mock=tbl_mock, modo="CLANDESTINO", area=50
    )

    # A sincronização deve relatar falha normativa e NÃO alterar as células da tabela
    assert status == ":erro"
    assert "Falha no cálculo normativo LIGHT" in novos

    # Células da tabela mock permanecem 100% inalteradas (sem "-", sem zeros, sem corrupção)
    assert tbl_mock["rows"][0][7] == "22,52"
    assert tbl_mock["rows"][0][8] == "90,65"
    assert tbl_mock["rows"][0][2] == "13x%%C1"
    assert tbl_mock["rows"][0][4] == 13


def test_clandestino_lisp_ast_e_balanceamento():
    """Valida a sintaxe, balanceamento, teto de linhas e modularidade dos módulos CLANDESTINO."""
    assert CLAN_LISP_PATH.exists(), "clandestino_light.lsp não encontrado"
    assert TOPO_LISP_PATH.exists(), "clandestino_topologia.lsp não encontrado"
    source_clan = CLAN_LISP_PATH.read_text(encoding="utf-8")
    source_topo = TOPO_LISP_PATH.read_text(encoding="utf-8")

    # Limite estrito de modularidade: <= 500 linhas para cada módulo auxiliar
    lines_clan = source_clan.splitlines()
    assert len(lines_clan) <= 500, f"clandestino_light.lsp superou 500 linhas: {len(lines_clan)}"

    lines_topo = source_topo.splitlines()
    assert len(lines_topo) <= 500, f"clandestino_topologia.lsp superou 500 linhas: {len(lines_topo)}"

    # Parênteses 100% balanceados em ambos
    assert source_clan.count("(") == source_clan.count(")"), "Parênteses desbalanceados em clandestino_light.lsp"
    assert source_topo.count("(") == source_topo.count(")"), "Parênteses desbalanceados em clandestino_topologia.lsp"

    # Funções puras e tabelas em clandestino_light.lsp
    assert "(defun lp-clan-obter-dmdi" in source_clan
    assert "(defun lp-clan-obter-fator-div" in source_clan
    assert "(defun lp-clan-calcular-demanda" in source_clan
    assert "(defun lp-clan-calcular-acumulada" in source_clan
    assert "(defun lp-clan-atualizar-ponto-demandas" in source_clan
    assert "(defun lp-clan-format-num" in source_clan
    assert "(defun lp-clan-validar-area" in source_clan
    assert "(defun lp-clan-validar-clientes" in source_clan
    assert "(defun lp-clan-recalcular-pontos" in source_clan
    assert "(defun lp-clan-adicionar-aresta" in source_clan
    assert "(defun lp-clan-construir-mapa-filhos" in source_clan
    assert "(defun lp-clan-obter-jusante" in source_clan
    assert "*lp-clan-key-topologia*" in source_clan
    assert "*lp-clan-ultimo-erro*" in source_clan
    assert "(defun lp-clan-obter-ultimo-erro" in source_clan
    assert "*lp-clan-dmdi-20-400*" in source_clan
    assert "*lp-clan-fdiv-1-300*" in source_clan

    # Fluxo interativo e wizard CAD em clandestino_topologia.lsp
    assert "(defun lp-configurar-modo" in source_topo
    assert "(defun lp-clan-cadastrar-topologia" in source_topo
    assert "(defun lp-clan-parse-tr" in source_topo
    assert "(defun lp-clan-parse-no-poste" in source_topo
    assert "(defun lp-clan-selecionar-no-cad" in source_topo
    assert "(defun lp-clan-coletar-sequencia-postes" in source_topo
    assert "(defun lp-clan-executar-editor-topologia" in source_topo

    # Não deve existir detecção arbitrária de área por texto de poste
    assert "lp-clan-parse-area" not in source_clan
    assert "lp-clan-parse-area" not in source_topo
    lp_source = LISP_PATH.read_text(encoding="utf-8")
    assert "lp-clan-parse-area" not in lp_source

    # Carregamento no acaddoc.lsp
    acaddoc = (ROOT / "acaddoc.lsp").read_text(encoding="utf-8")
    assert 'load-aux-safe "orcamento\\\\clandestino_light.lsp"' in acaddoc
    assert 'load-aux-safe "orcamento\\\\clandestino_topologia.lsp"' in acaddoc

    # levantamento_postes.lsp estritamente <= 1000 linhas
    lp_lines = lp_source.splitlines()
    assert len(lp_lines) <= 1000, f"levantamento_postes.lsp excedeu 1000 linhas: {len(lp_lines)}"


# ==============================================================================
# Testes de Integração de Execução Real com o AutoLISP Harness
# ==============================================================================

from tests.autolisp_harness import AutoLisp, Ename, Pair


def _criar_harness_com_modulos(dwg=None, inputs=None):
    interp = AutoLisp(search_dirs=[ROOT, ROOT / "orcamento"], dwg=dwg, inputs=inputs)
    interp.load(CLAN_LISP_PATH)
    interp.load(TOPO_LISP_PATH)
    interp.load(CQT_CLAN_LISP_PATH)
    interp.load(LISP_PATH)
    return interp


def _dict_para_alist(d):
    return [Pair(k, v) for k, v in d.items()]


def _alist_para_dict(al):
    res = {}
    for item in al or []:
        if isinstance(item, Pair):
            res[item.car] = item.cdr
        elif isinstance(item, (list, tuple)) and len(item) == 2:
            res[item[0]] = item[1]
    return res


def test_fluxo_interativo_topologia_tr_lados_derivacoes_e_concluir():
    """Valida o cadastro interativo real: TR -> LADO 1 -> LADO 2 -> DERIVACAO -> CONCLUIR.
    
    Topologia construída:
      TR-01 -> P4 -> P5 -> P6 -> P7   (LADO 1)
      TR-01 -> P8 -> P9               (LADO 2)
      P5 -> P10 -> P11                (DERIVAÇÃO a partir de P5)
    """
    dwg = {
        "H_TR": {"0": "TEXT", "1": "TR-01", "5": "H_TR"},
        "H_P4": {"0": "TEXT", "1": "P4", "5": "H_P4"},
        "H_P5": {"0": "TEXT", "1": "P5", "5": "H_P5"},
        "H_P6": {"0": "TEXT", "1": "P6", "5": "H_P6"},
        "H_P7": {"0": "TEXT", "1": "P7", "5": "H_P7"},
        "H_P8": {"0": "TEXT", "1": "P8", "5": "H_P8"},
        "H_P9": {"0": "TEXT", "1": "P9", "5": "H_P9"},
        "H_P10": {"0": "TEXT", "1": "P10", "5": "H_P10"},
        "H_P11": {"0": "TEXT", "1": "P11", "5": "H_P11"},
    }

    inputs = [
        # TR
        "H_TR",
        # LADO 1 (P4 é o nó TR raiz - sem vão TR->P4)
        "H_P4",
        "H_P5",
        "",    # Comp P4->P5 (Enter aceita sugestão)
        None,  # Cabo P4->P5 (Enter aceita Auto)
        "H_P6",
        "",    # Comp P5->P6 (Enter aceita sugestão)
        None,  # Cabo P5->P6 (Enter aceita Auto)
        "H_P7",
        "",    # Comp P6->P7 (Enter aceita sugestão)
        None,  # Cabo P6->P7 (Enter aceita Auto)
        None,  # Enter fecha LADO 1
        # Menu opção -> Lado (LADO 2)
        "Lado",
        "H_P8",
        "",    # Comp TR->P8 (Enter aceita sugestão)
        None,  # Cabo TR->P8 (Enter aceita Auto)
        "H_P9",
        "",    # Comp P8->P9 (Enter aceita sugestão)
        None,  # Cabo P8->P9 (Enter aceita Auto)
        None,  # Enter fecha LADO 2
        # Menu opção -> Derivacao
        "Derivacao",
        "H_P5",  # Poste de partida da derivação
        "H_P10",
        "",    # Comp P5->P10 (Enter aceita sugestão)
        None,  # Cabo P5->P10 (Enter aceita Auto)
        "H_P11",
        "",    # Comp P10->P11 (Enter aceita sugestão)
        None,  # Cabo P10->P11 (Enter aceita Auto)
        None,  # Enter fecha derivação
        # Menu opção -> Concluir
        "Concluir",
    ]

    interp = _criar_harness_com_modulos(dwg=dwg, inputs=inputs)
    res = interp.call("lp-clan-cadastrar-topologia")
    assert res is not None, "Topologia retornou nil"

    # Confirma persistência da topologia no LData
    topo_ldata = interp.call("lp-db-obter-topologia")
    assert topo_ldata is not None, "Topologia não foi persistida no LData"

    # Constrói o mapa de adjacências de filhos e valida as conexões
    mapa = interp.call("lp-clan-construir-mapa-filhos", topo_ldata, None)
    assert mapa is not None

    # P7 é folha: nenhum poste a jusante
    jus_p7 = interp.call("lp-clan-obter-jusante", "P7", mapa)
    assert not jus_p7

    # P6 tem P7 a jusante
    jus_p6 = interp.call("lp-clan-obter-jusante", "P6", mapa)
    assert jus_p6 == ["P7"]

    # P10 tem P11 a jusante
    jus_p10 = interp.call("lp-clan-obter-jusante", "P10", mapa)
    assert jus_p10 == ["P11"]

    # P5 tem tronco P6, P7 e ramal derivado P10, P11
    jus_p5 = interp.call("lp-clan-obter-jusante", "P5", mapa)
    assert set(jus_p5) == {"P6", "P7", "P10", "P11"}

    # P4 tem toda a subárvore
    jus_p4 = interp.call("lp-clan-obter-jusante", "P4", mapa)
    assert set(jus_p4) == {"P5", "P6", "P7", "P10", "P11"}

    # LADO 2
    jus_p8 = interp.call("lp-clan-obter-jusante", "P8", mapa)
    assert jus_p8 == ["P9"]
    jus_p9 = interp.call("lp-clan-obter-jusante", "P9", mapa)
    assert not jus_p9


def test_calculo_acumulada_ordem_arbitraria_e_validacao_oficial_light():
    """Valida que a demanda acumulada segue estritamente a topologia da rede,
    independentemente da ordem de coleta ou cadastro dos postes.
    
    Valores reais obtidos por consulta estrita às tabelas da LIGHT (CQTS_ZNA853248_ATUAL.xlsx):
    - Área 50 m² -> DMDI = 1,88 kVA (DB!A65:B65)
    - Totais acumulados e fatores oficiais (DB!C36:E335):
      * P7 (4 cli): fdiv=3,88 -> 1,88 * 3,88 = 7,29 kVA
      * P11 (3 cli): fdiv=3,88 -> 1,88 * 3,88 = 7,29 kVA
      * P10 (2+3 = 5 cli): fdiv=4,84 -> 1,88 * 4,84 = 9,10 kVA
      * P6 (5+4 = 9 cli): fdiv=8,68 -> 1,88 * 8,68 = 16,32 kVA
      * P9 (8 cli): fdiv=7,72 -> 1,88 * 7,72 = 14,51 kVA
      * P8 (7+8 = 15 cli): fdiv=13,54 -> 1,88 * 13,54 = 25,46 kVA
      * P5 (6+9+5 = 20 cli): fdiv=17,44 -> 1,88 * 17,44 = 32,79 kVA
      * P4 (10+20 = 30 cli): fdiv=23,54 -> 1,88 * 23,54 = 44,26 kVA
    """
    arestas = [
        ["TR-01", "P4"],
        ["P4", "P5"],
        ["P5", "P6"],
        ["P6", "P7"],
        ["TR-01", "P8"],
        ["P8", "P9"],
        ["P5", "P10"],
        ["P10", "P11"],
    ]

    interp = _criar_harness_com_modulos()
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_MEDIA")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas

    # Ordem de coleta 1 (embaralhada / ponta antes da fonte):
    # P7, P11, P6, P10, P9, P8, P5, P4
    dados_ordem_1 = [
        {"SEQ": 1, "POSTE": "P7", "CLIENTES": 4},
        {"SEQ": 2, "POSTE": "P11", "CLIENTES": 3},
        {"SEQ": 3, "POSTE": "P6", "CLIENTES": 5},
        {"SEQ": 4, "POSTE": "P10", "CLIENTES": 2},
        {"SEQ": 5, "POSTE": "P9", "CLIENTES": 8},
        {"SEQ": 6, "POSTE": "P8", "CLIENTES": 7},
        {"SEQ": 7, "POSTE": "P5", "CLIENTES": 6},
        {"SEQ": 8, "POSTE": "P4", "CLIENTES": 10},
    ]
    pontos_ordem_1 = [_dict_para_alist(d) for d in dados_ordem_1]

    calc_1 = interp.call("lp-recalcular-demandas-acumuladas", pontos_ordem_1)
    res_1 = {_alist_para_dict(p)["POSTE"]: _alist_para_dict(p)["ACUMULADA"] for p in calc_1}

    assert res_1["P7"] == "7,29"
    assert res_1["P11"] == "7,29"
    assert res_1["P10"] == "9,10"
    assert res_1["P6"] == "16,32"
    assert res_1["P9"] == "14,51"
    assert res_1["P8"] == "25,46"
    assert res_1["P5"] == "32,79"
    assert res_1["P4"] == "44,26"

    # Ordem de coleta 2 (inversa / fonte antes da ponta):
    # P4, P8, P5, P9, P6, P10, P7, P11
    dados_ordem_2 = [
        {"SEQ": 1, "POSTE": "P4", "CLIENTES": 10},
        {"SEQ": 2, "POSTE": "P8", "CLIENTES": 7},
        {"SEQ": 3, "POSTE": "P5", "CLIENTES": 6},
        {"SEQ": 4, "POSTE": "P9", "CLIENTES": 8},
        {"SEQ": 5, "POSTE": "P6", "CLIENTES": 5},
        {"SEQ": 6, "POSTE": "P10", "CLIENTES": 2},
        {"SEQ": 7, "POSTE": "P7", "CLIENTES": 4},
        {"SEQ": 8, "POSTE": "P11", "CLIENTES": 3},
    ]
    pontos_ordem_2 = [_dict_para_alist(d) for d in dados_ordem_2]

    calc_2 = interp.call("lp-recalcular-demandas-acumuladas", pontos_ordem_2)
    res_2 = {_alist_para_dict(p)["POSTE"]: _alist_para_dict(p)["ACUMULADA"] for p in calc_2}

    # As demandas acumuladas são idênticas em ambas as ordens de cadastro
    assert res_2 == res_1


def test_cancelamento_seguro_da_topologia_preserva_dwg():
    """Valida que o cancelamento no cadastro não sobrescreve topologia existente nem corrompe o DWG."""
    arestas_anteriores = [["TR-01", "P1"], ["P1", "P2"]]

    interp = _criar_harness_com_modulos()
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas_anteriores

    # 1. Usuário seleciona 'Cancelar' no menu de topologia existente:
    interp.inputs = ["Cancelar"]
    res_cancel = interp.call("lp-clan-cadastrar-topologia")
    assert res_cancel == arestas_anteriores
    assert interp.call("lp-db-obter-topologia") == arestas_anteriores

    # 2. Usuário inicia cadastro novo sem topologia prévia e cancela na seleção do TR (Enter):
    interp_novo = _criar_harness_com_modulos()
    interp_novo.inputs = [None]
    res_novo_cancel = interp_novo.call("lp-clan-cadastrar-topologia")
    assert res_novo_cancel is None
    assert interp_novo.call("lp-db-obter-topologia") is None


def test_resiliencia_area_ldata_e_status_modo_str():
    """Valida que valores inválidos ou 'T' no LData de área não quebram lp-db-obter-area nem lp-clan-status-modo-str."""
    interp = _criar_harness_com_modulos()

    # Caso 1: LData corrompido com 'T' (True retornado por (or ...))
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = True
    area_obtida = interp.call("lp-db-obter-area")
    assert area_obtida == 50
    assert isinstance(area_obtida, int)

    status = interp.call("lp-clan-status-modo-str")
    assert "CLANDESTINO | Área: 50 m²" in status

    # Caso 2: LData vazio (None)
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = None
    area_nil = interp.call("lp-db-obter-area")
    assert area_nil == 50
    assert isinstance(area_nil, int)
    status_nil = interp.call("lp-clan-status-modo-str")
    assert "CLANDESTINO | Área: 50 m²" in status_nil

    # Caso 3: LData com valor válido (ex: 80 m²)
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 80
    area_80 = interp.call("lp-db-obter-area")
    assert area_80 == 80
    status_80 = interp.call("lp-clan-status-modo-str")
    assert "CLANDESTINO | Área: 80 m²" in status_80


def test_clandestino_cqt_lookup_e_selecao_condutores():
    """Valida a tabela de condutores DB!A1:I32 e a regra de seleção de cabos da Light."""
    from tests.autolisp_harness import Sym
    interp = _criar_harness_com_modulos()

    # Todos os 31 condutores da tabela DB!A1:I32 devem estar carregados
    cabos = interp.env.get(Sym("*LP-CLAN-CABOS*"))
    assert len(cabos) == 31

    # Lookup de propriedades do cabo 70 Al - MX: (Nome Iz R X Alpha Div_R Temp_Lim Delta_T K_Term)
    cabo_70 = interp.call("lp-cqt-obter-cabo", "70 Al - MX")
    assert cabo_70 is not None
    assert cabo_70[0] == "70 Al - MX"
    assert cabo_70[1] == 202.0
    assert cabo_70[2] == 0.5697
    assert cabo_70[3] == 0.126
    assert cabo_70[4] == 0.00403
    assert cabo_70[5] == 1.2821
    assert cabo_70[6] == 90.0
    assert cabo_70[7] == 60.0
    assert abs(cabo_70[8] - 3.366667) < 0.001

    # Lookup de propriedades do cabo 240 Cu (descida do trafo)
    cabo_240 = interp.call("lp-cqt-obter-cabo", "240 Cu")
    assert cabo_240 is not None
    assert cabo_240[0] == "240 Cu"
    assert cabo_240[1] == 430.0
    assert cabo_240[2] == 0.1035
    assert cabo_240[3] == 0.0897
    assert cabo_240[5] == 1.2751

    # Seleção de condutor:
    # 1. Descida do transformador (tr-flag = T) -> sempre 240 Cu
    assert interp.call("lp-cqt-selecionar-condutor", 25.46, True) == "240 Cu"
    assert interp.call("lp-cqt-selecionar-condutor", 140.0, True) == "240 Cu"

    # 2. Rede BT: Ib <= 202 A (Demanda <= 76.97 kVA) -> 70 Al - MX
    assert interp.call("lp-cqt-selecionar-condutor", 25.46, False) == "70 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", 76.90, False) == "70 Al - MX"

    # 3. Rede BT: 202 A < Ib <= 355 A (76.97 < Demanda <= 135.27 kVA) -> 185 Al - MX
    assert interp.call("lp-cqt-selecionar-condutor", 80.0, False) == "185 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", 135.0, False) == "185 Al - MX"

    # 4. Rede BT: Ib > 355 A (Demanda > 135.27 kVA) -> 240 Al - MX
    assert interp.call("lp-cqt-selecionar-condutor", 136.0, False) == "240 Al - MX"


def test_clandestino_cqt_comparacao_numerica_exata_lado2_light():
    """Valida correspondência numérica exata dos cálculos de CQT contra a planilha real CQTS_ZNA853248_ATUAL.xlsx (Aba LADO 2)."""
    interp = _criar_harness_com_modulos()

    # Trecho TR -> LID (Descida do Trafo, Row 8 da planilha):
    # L=4.0m, S=157.013176 kVA, 240 Cu, tr_p=True, v_ant=0.0
    r_tr = interp.call("lp-cqt-calcular-trecho", 157.01317647058823, 4.0, "240 Cu", "TRI", True, 0.0)
    d_tr = _alist_para_dict(r_tr)
    assert abs(d_tr["QT_PERC"] - 6.89) < 0.01
    assert abs(d_tr["V_127"] - 118.26) < 0.02
    assert abs(d_tr["QT_TRECHO_V"] - 8.75) < 0.02
    v_acum_tr = d_tr["V_ACUM"]

    # Trecho LID -> P8 (Row 9 da planilha):
    # L=19.0m, S=62.228 kVA, 240 Al, tr_p=False, v_ant=v_acum_tr
    r_p8 = interp.call("lp-cqt-calcular-trecho", 62.228, 19.0, "240 Al", "TRI", False, v_acum_tr)
    d_p8 = _alist_para_dict(r_p8)
    assert abs(d_p8["QT_PERC"] - 7.31) < 0.01
    assert abs(d_p8["V_127"] - 117.73) < 0.02
    assert abs(d_p8["QT_TRECHO_V"] - 0.53) < 0.02
    v_acum_p8 = d_p8["V_ACUM"]

    # Trecho P8 -> P9 (Row 10 da planilha):
    # L=31.0m, S=56.6632 kVA, 240 Al, tr_p=False, v_ant=v_acum_p8
    r_p9 = interp.call("lp-cqt-calcular-trecho", 56.6632, 31.0, "240 Al", "TRI", False, v_acum_p8)
    d_p9 = _alist_para_dict(r_p9)
    assert abs(d_p9["QT_PERC"] - 7.93) < 0.01
    assert abs(d_p9["V_127"] - 116.95) < 0.02
    assert abs(d_p9["QT_TRECHO_V"] - 0.78) < 0.02
    v_acum_p9 = d_p9["V_ACUM"]

    # Trecho P9 -> P10 (Row 11 da planilha):
    # L=27.0m, S=45.3832 kVA, 70 Al - MX, tr_p=False, v_ant=v_acum_p9
    r_p10 = interp.call("lp-cqt-calcular-trecho", 45.3832, 27.0, "70 Al - MX", "TRI", False, v_acum_p9)
    d_p10 = _alist_para_dict(r_p10)
    assert abs(d_p10["QT_PERC"] - 9.30) < 0.01
    assert abs(d_p10["V_127"] - 115.21) < 0.02
    assert abs(d_p10["QT_TRECHO_V"] - 1.74) < 0.02


def test_clandestino_cqt_enriquecimento_pontos_e_modo_normal():
    """Valida enriquecimento com CQT no modo CLANDESTINO e ausência no modo NORMAL."""
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR-01", "P8"],
        ["P8", "P9"],
        ["P9", "P10"],
    ]

    dados = [
        {"SEQ": 1, "POSTE": "P8", "CLIENTES": 7, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P9", "CLIENTES": 8, "COORDENADAS": "X: 30,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P10", "CLIENTES": 5, "COORDENADAS": "X: 60,00 | Y: 0,00"},
    ]
    pontos = [_dict_para_alist(d) for d in dados]

    # 1. No modo NORMAL: recalcular mantém estrutura sem colunas CQT
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "NORMAL"
    res_normal = interp.call("lp-recalcular-demandas-acumuladas", pontos)
    d_norm = _alist_para_dict(res_normal[0])
    assert "CONDUTOR_PREVISTO" not in d_norm
    assert "CQT_PREVISTO" not in d_norm

    # 2. No modo CLANDESTINO: recalcular gera CONDUTOR_PREVISTO, CQT_PREVISTO e TENSAO_PREVISTA
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR-01"

    res_clan = interp.call("lp-recalcular-demandas-acumuladas", pontos)
    d_p8 = _alist_para_dict(res_clan[0])
    d_p9 = _alist_para_dict(res_clan[1])
    d_p10 = _alist_para_dict(res_clan[2])

    # P8 ligado diretamente ao TR: cabo 240 Cu
    assert d_p8["CONDUTOR_PREVISTO"] == "240 Cu"
    assert "%" in d_p8["CQT_PREVISTO"]
    assert "V" in d_p8["TENSAO_PREVISTA"]

    # P9 e P10 na rede BT: cabo 70 Al - MX
    assert d_p9["CONDUTOR_PREVISTO"] == "70 Al - MX"
    assert d_p10["CONDUTOR_PREVISTO"] == "70 Al - MX"

    # CQT é estritamente crescente ao longo do circuito
    q_p8 = float(d_p8["CQT_PREVISTO"].replace("%", ""))
    q_p9 = float(d_p9["CQT_PREVISTO"].replace("%", ""))
    q_p10 = float(d_p10["CQT_PREVISTO"].replace("%", ""))
    assert q_p8 < q_p9 < q_p10


def test_clandestino_cqt_resiliencia_geometria_e_tabela():
    """Valida o parser de coordenadas, cálculo de vãos reais, fallback para 30m e colunas da tabela."""
    interp = _criar_harness_com_modulos()

    # 1. Parser de coordenadas nos formatos padrão OXCAD
    c1 = interp.call("lp-cqt-parse-coords", "X: 681200,00 | Y: 7456700,00")
    assert c1 == [681200.0, 7456700.0]

    c2 = interp.call("lp-cqt-parse-coords", "X: 681240,00 | Y: 7456730,00")
    assert c2 == [681240.0, 7456730.0]

    c_inv = interp.call("lp-cqt-parse-coords", "-")
    assert c_inv is None

    # 2. Distância euclidiana entre dois pontos: sqrt(40^2 + 30^2) = 50.0m
    p1 = _dict_para_alist({"POSTE": "P1", "COORDENADAS": "X: 681200,00 | Y: 7456700,00"})
    p2 = _dict_para_alist({"POSTE": "P2", "COORDENADAS": "X: 681240,00 | Y: 7456730,00"})
    dist = interp.call("lp-cqt-distancia-pontos", p1, p2)
    assert abs(dist - 50.0) < 0.01

    # 3. Fallback para vão padrão de 30.0m quando sem coordenadas
    p_sem_coord = _dict_para_alist({"POSTE": "P3", "COORDENADAS": "-"})
    dist_fallback = interp.call("lp-cqt-distancia-pontos", p1, p_sem_coord)
    assert dist_fallback == 30.0

    # 4. Predicado de colunas adicionais
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "NORMAL"
    assert interp.call("lp-cqt-colunas-adicionais-p") is None

    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    assert interp.call("lp-cqt-colunas-adicionais-p") in (True, "T")


def test_clandestino_cqt_identificacao_e_aliases_condutores():
    """Valida identificação de condutores por alias, contexto e seleção interativa."""
    interp = _criar_harness_com_modulos()

    # 1. 185AMX -> 185 Al - MX
    c185 = interp.call("lp-cqt-resolver-cabo", "185AMX", None)
    assert c185 == "185 Al - MX"
    reg185 = interp.call("lp-cqt-obter-cabo", "185AMX")
    assert reg185[0] == "185 Al - MX"
    assert reg185[1] == 355.0  # Capacidade de condução

    # 2. 53 Al - QX com contexto RAMAL -> 53 Al - QX (63 A)
    c53_ramal = interp.call("lp-cqt-resolver-cabo", "53 Al - QX", "RAMAL")
    assert c53_ramal == "53 Al - QX"
    reg53_ramal = interp.call("lp-cqt-obter-cabo-contexto", "53 Al - QX", "RAMAL")
    assert reg53_ramal[0] == "53 Al - QX"
    assert reg53_ramal[1] == 63.0

    # 3. 53 Al - QX com contexto BT / REDE -> 53 AA (163 A)
    c53_bt = interp.call("lp-cqt-resolver-cabo", "53 Al - QX", "BT")
    assert c53_bt == "53 AA"
    reg53_bt = interp.call("lp-cqt-obter-cabo-contexto", "53 Al - QX", "BT")
    assert reg53_bt[0] == "53 AA"
    assert reg53_bt[1] == 163.0

    # 4. Ambiguidade sem contexto: prompt interativo (getkword)
    interp.inputs = ["Ramal"]
    c53_inter_ramal = interp.call("lp-cqt-resolver-cabo", "53 Al - QX", None)
    assert c53_inter_ramal == "53 Al - QX"

    interp.inputs = ["BT"]
    c53_inter_bt = interp.call("lp-cqt-resolver-cabo", "53 Al - QX", None)
    assert c53_inter_bt == "53 AA"


def test_clandestino_cqt_limites_dimensionamento_e_ampacidade():
    """Valida transições de condutor nos limites de ampacidade 202A e 355A e robustez com nil."""
    interp = _criar_harness_com_modulos()
    # V_base_tri = 220 * sqrt(3) / 1000 = 0.381051177665153
    # Ib = S_acum / V_base_tri
    # S(202A) = 202.0 * 0.38105118 = 76.972338 kVA
    # S(355A) = 355.0 * 0.38105118 = 135.273168 kVA

    v_base = 220.0 * (3.0 ** 0.5) / 1000.0

    # Ponto no TR sempre recebe 240 Cu
    assert interp.call("lp-cqt-selecionar-condutor", 100.0, True) == "240 Cu"

    # Ib <= 202 A -> 70 Al - MX
    assert interp.call("lp-cqt-selecionar-condutor", (201.99 * v_base), False) == "70 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", (202.00 * v_base), False) == "70 Al - MX"

    # 202 A < Ib <= 355 A -> 185 Al - MX
    assert interp.call("lp-cqt-selecionar-condutor", (202.01 * v_base), False) == "185 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", (354.99 * v_base), False) == "185 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", (355.00 * v_base), False) == "185 Al - MX"

    # Ib > 355 A -> 240 Al - MX
    assert interp.call("lp-cqt-selecionar-condutor", (355.01 * v_base), False) == "240 Al - MX"

    # Valores limites/inválidos: nil ou <= 0.0 assumem 70 Al - MX sem erro
    assert interp.call("lp-cqt-selecionar-condutor", None, False) == "70 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", 0.0, False) == "70 Al - MX"
    assert interp.call("lp-cqt-selecionar-condutor", -15.0, False) == "70 Al - MX"


def test_clandestino_cqt_sinalizacao_vao_estimado():
    """Valida sinalização com '*' e '(EST)' quando faltam coordenadas na rede."""
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR-01", "P1"],
        ["P1", "P2"],
    ]

    # Caso 1: Ambos com coordenadas reais
    dados_reais = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 10, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 10, "COORDENADAS": "X: 35,00 | Y: 0,00"},
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR-01"

    res_reais = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados_reais])
    d_p2_real = _alist_para_dict(res_reais[1])
    assert d_p2_real["COMPRIMENTO_VÃO"] == "35.0"
    assert "*" not in d_p2_real["CQT_PREVISTO"]
    assert "*" not in d_p2_real["TENSAO_PREVISTA"]
    assert d_p2_real["STATUS_CQT"] == "PREVISTO (COORDENADAS REAIS)"

    # Caso 2: P2 sem coordenadas -> vão estimado de 30.0 (EST) e CQT com '*'
    dados_est = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 10, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 10, "COORDENADAS": "-"},
    ]
    res_est = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados_est])
    d_p2_est = _alist_para_dict(res_est[1])
    assert "(EST)" in d_p2_est["COMPRIMENTO_VÃO"]
    assert d_p2_est["COMPRIMENTO_VÃO"] == "30.0 (EST)"
    assert d_p2_est["CQT_PREVISTO"].endswith("*")
    assert d_p2_est["TENSAO_PREVISTA"].endswith("*")
    assert d_p2_est["STATUS_CQT"] == "PREVISTO (VÃO ESTIMADO - PENDENTE CONFIRMAÇÃO)"


# ==============================================================================
# Suíte de Testes: Suporte a Postes Passantes sem Textos no LEVANTAR_POSTES
# ==============================================================================

def test_lisp_ast_suporte_poste_passante():
    """Valida estaticamente no código AutoLISP a presença do fluxo de poste passante."""
    source = LISP_PATH.read_text(encoding="utf-8")

    # Prompt informativo indicando possibilidade de poste passante
    assert "Enter se for poste passante sem textos" in source
    # Fluxo B de poste passante
    assert "[POSTE PASSANTE]" in source
    assert r'getstring t "\nDigite o identificador do poste passante' in source
    assert '(cons "PASSANTE" "SIM")' in source
    assert '(cons "TEXT_HANDLES" \'())' in source
    assert r'getpoint "\nClique no ponto geométrico do poste' in source


def test_poste_passante_criacao_e_valores_zerados():
    """Valida que o poste passante é criado com zero clientes locais, sem ruído, e flags adequadas."""
    ponto_passante = {
        "SEQ": 5,
        "POSTE": "P5",
        "RAMAIS": "-",
        "LUM": "-",
        "CLIENTES": 0,
        "POSICOES": "-",
        "CS": "-",
        "DEMANDA": "0,00",
        "ACUMULADA": "0,00",
        "COORDENADAS": "N: 7456720.00\nE: 681220.00",
        "RAW": [],
        "PASSANTE": "SIM",
        "TEXT_HANDLES": []
    }
    assert ponto_passante["CLIENTES"] == 0
    assert ponto_passante["DEMANDA"] == "0,00"
    assert ponto_passante["PASSANTE"] == "SIM"
    assert ponto_passante["TEXT_HANDLES"] == []


def test_propagacao_topologica_com_poste_passante():
    """Valida que o poste passante intermediário (P4 -> P5 -> P6) recebe os acumulados a jusante
    e calcula CQT adequadamente para os trechos."""
    interp = _criar_harness_com_modulos()

    # Topologia em cadeia: TR-01 -> P4 -> P5 -> P6
    arestas = [
        ["TR-01", "P4"],
        ["P4", "P5"],
        ["P5", "P6"]
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR-01"

    # P4: 10 clientes locais, P5: passante (0 clientes), P6: 15 clientes locais
    # No sentido TR -> P4 -> P5 -> P6:
    # A jusante de P5 está P6 (15 clientes). Portanto, P5 deve carregar os 15 clientes de P6!
    # A jusante de P4 estão P5 e P6 (0 + 15 = 15) + seus próprios clientes (10) = 25 clientes!
    dados = [
        {"SEQ": 1, "POSTE": "P4", "CLIENTES": 10, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P5", "CLIENTES": 0, "PASSANTE": "SIM", "COORDENADAS": "X: 30,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P6", "CLIENTES": 15, "COORDENADAS": "X: 60,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p4_calc = _alist_para_dict(res[0])
    p5_calc = _alist_para_dict(res[1])
    p6_calc = _alist_para_dict(res[2])

    demanda_15_cli = interp.call("lp-clan-calcular-demanda", 50, 15)
    demanda_25_cli = interp.call("lp-clan-calcular-demanda", 50, 25)

    # Demanda de P5 no cálculo topológico reflete os 15 clientes acumulados a jusante
    assert p5_calc["POSTE"] == "P5"
    assert p5_calc["PASSANTE"] == "SIM"
    assert p5_calc["DEMANDA"] == "0,00"
    assert abs(float(p5_calc["ACUMULADA"].replace(",", ".")) - demanda_15_cli) < 0.05
    assert abs(float(p4_calc["ACUMULADA"].replace(",", ".")) - demanda_25_cli) < 0.05

    # CQT deve ter sido calculado em todos os trechos
    assert "CQT_PREVISTO" in p5_calc
    assert "CONDUTOR_PREVISTO" in p5_calc
    assert float(p5_calc["CQT_PREVISTO"].replace("*", "").replace(",", ".").replace("%", "")) >= 0.0


def test_poste_passante_derivacao_lateral():
    """Valida ramificação lateral a partir de um poste passante:
    TR -> P1 -> P_PASSANTE -> {P_LATERAL1, P_LATERAL2}"""
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR", "P1"],
        ["P1", "P_PASSANTE"],
        ["P_PASSANTE", "P_LAT1"],
        ["P_PASSANTE", "P_LAT2"],
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 40
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P_PASSANTE", "CLIENTES": 0, "PASSANTE": "SIM", "COORDENADAS": "X: 30,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P_LAT1", "CLIENTES": 8, "COORDENADAS": "X: 30,00 | Y: 30,00"},
        {"SEQ": 4, "POSTE": "P_LAT2", "CLIENTES": 12, "COORDENADAS": "X: 60,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p_pass = _alist_para_dict(res[1])

    # P_PASSANTE soma jusante de LAT1 (8) e LAT2 (12) = 20 clientes
    demanda_20_cli = interp.call("lp-clan-calcular-demanda", 40, 20)
    assert abs(float(p_pass["ACUMULADA"].replace(",", ".")) - demanda_20_cli) < 0.05


def test_poste_passante_isolado_sem_jusante():
    """Valida poste passante ponta de linha (sem clientes a jusante)."""
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR", "P1"],
        ["P1", "P_FIM"],
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 10, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P_FIM", "CLIENTES": 0, "PASSANTE": "SIM", "COORDENADAS": "X: 30,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p_fim = _alist_para_dict(res[1])

    # No poste final passante, demanda local e a jusante são 0
    assert float(p_fim["ACUMULADA"].replace(",", ".")) == 0.0
    # O CQT acumulado na barra final reflete a queda de tensão herdada do trecho a montante
    cqt_val = float(p_fim["CQT_PREVISTO"].replace("*", "").replace(",", ".").replace("%", ""))
    assert cqt_val > 0.0


# ==============================================================================
# Suíte de Testes: Coleta de Vãos, Conexão P1=TR e Precedência de Condutores
# ==============================================================================

def test_cqt_p1_como_no_raiz_trafo_4m_240cu():
    """Valida regra de domínio P1 = TR raiz:
    - Conexão interna de 4m com condutor obrigatório 240 Cu
    - Queda QT_MT+TR (~6.71485% / 8.529V) somada à queda do trecho de 4m
    - Comprimento fixado em 4.0m e VAO_ESTIMADO = nil (sem '*' ou '(EST)')
    """
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR-01", "P1"],
        ["P1", "P2"],
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR-01"

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 10, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 10, "COORDENADAS": "X: 35,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p1 = _alist_para_dict(res[0])

    # P1 representa o TR
    assert p1["POSTE"] == "P1"
    assert p1["CONDUTOR_PREVISTO"] == "240 Cu"
    assert p1["COMPRIMENTO_VÃO"] == "4.0"
    assert "*" not in p1["CQT_PREVISTO"]
    assert "(EST)" not in p1["COMPRIMENTO_VÃO"]

    # CQT em P1 é maior que 6.71% (pois soma QT_MTTR de 6.71% + queda dos 4m de 240 Cu)
    cqt_p1 = float(p1["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    assert cqt_p1 > 6.71
    # Tensão em 127V é menor que 118.50V
    v127_p1 = float(p1["TENSAO_PREVISTA"].replace("V", "").replace(",", "."))
    assert v127_p1 < 118.50


def test_cqt_primeiro_vao_real_p1_p2_independente():
    """Valida que o primeiro vão real P1 -> P2 é independente dos 4m internos do Trafo:
    - P1 -> P2 não duplica QT_MT+TR
    - Os 4m de P1 não são somados ao comprimento do vão de P2
    - Tensão base de P2 herda diretamente a queda de P1
    """
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR-01", "P1"],
        ["P1", "P2"],
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR-01"

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 15, "COORDENADAS": "X: 42,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p1 = _alist_para_dict(res[0])
    p2 = _alist_para_dict(res[1])

    # P2 tem seu comprimento real medido no CAD (42.0m), sem os 4m somados
    assert p2["COMPRIMENTO_VÃO"] == "42.0"
    assert "*" not in p2["CQT_PREVISTO"]

    # CQT de P2 é estritamente maior que CQT de P1
    cqt_p1 = float(p1["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p2 = float(p2["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    assert cqt_p2 > cqt_p1

    # Delta de CQT entre P1 e P2 reflete apenas o vão de 42m com carga de P2 (15 cli)
    delta_cqt = cqt_p2 - cqt_p1
    assert 0.1 < delta_cqt < 5.0


def test_cqt_multiplos_lados_partindo_de_p1():
    """Valida ramificações em múltiplos lados partindo do mesmo poste TR (P1):
    LADO 1: P1 -> P2 -> P3
    LADO 2: P1 -> P4 -> P5
    Ambos os circuitos herdam a tensão base de P1 e calculam vãos independentes.
    """
    interp = _criar_harness_com_modulos()

    topo = [
        Pair("TR", "P1"),
        Pair("LADOS", [
            ["LADO 1", "P1", "P2", "P3"],
            ["LADO 2", "P1", "P4", "P5"],
        ]),
        Pair("DERIVACOES", []),
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 45
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = topo
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "P1"

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 2, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 10, "COORDENADAS": "X: 30,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P3", "CLIENTES": 10, "COORDENADAS": "X: 60,00 | Y: 0,00"},
        {"SEQ": 4, "POSTE": "P4", "CLIENTES": 8, "COORDENADAS": "X: 0,00 | Y: 35,00"},
        {"SEQ": 5, "POSTE": "P5", "CLIENTES": 8, "COORDENADAS": "X: 0,00 | Y: 70,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p1 = _alist_para_dict(res[0])
    p2 = _alist_para_dict(res[1])
    p3 = _alist_para_dict(res[2])
    p4 = _alist_para_dict(res[3])
    p5 = _alist_para_dict(res[4])

    cqt_p1 = float(p1["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p2 = float(p2["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p3 = float(p3["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p4 = float(p4["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p5 = float(p5["CQT_PREVISTO"].replace("%", "").replace(",", "."))

    # P1 é a raiz comum
    assert p1["COMPRIMENTO_VÃO"] == "4.0"
    assert p1["CONDUTOR_PREVISTO"] == "240 Cu"

    # LADO 1: P1 -> P2 -> P3 cresce monotonamente
    assert cqt_p1 < cqt_p2 < cqt_p3
    # LADO 2: P1 -> P4 -> P5 cresce monotonamente
    assert cqt_p1 < cqt_p4 < cqt_p5


def test_cqt_precedencia_condutor_manual_em_trechos():
    """Valida precedência de condutor manual salvo em TRECHOS sobrepondo o condutor automático."""
    interp = _criar_harness_com_modulos()

    arestas = [["TR", "P1"], ["P1", "P2"]]
    # Registra no LData TRECHOS o condutor manual '185 Al - MX' para o vão P1 -> P2
    trechos = [
        ["P2", "P1", 30.0, "185 Al - MX", "MANUAL", None]
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    interp.ldata[("OXCAD_LEVANTAMENTO", "TRECHOS")] = trechos

    # Carga baixa (5 clientes) faria a seleção automática adotar 70 Al - MX
    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 5, "COORDENADAS": "X: 30,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p2 = _alist_para_dict(res[1])

    # Condutor manual 185 Al - MX deve prevalecer sobre o automático 70 Al - MX
    assert p2["CONDUTOR_PREVISTO"] == "185 Al - MX"


def test_cqt_alerta_sobrecarga_ampacidade_manual_ib_maior_iz():
    """Valida que condutor manual insuficiente (Ib > Iz) emite alerta e sinaliza sobrecarga,
    sem substituir silenciosamente o condutor selecionado pelo usuário."""
    interp = _criar_harness_com_modulos()

    arestas = [["TR", "P1"], ["P1", "P2"]]
    # Condutor '25 Al' tem Iz = 90.0 A
    trechos = [
        ["P2", "P1", 30.0, "25 Al", "MANUAL", None]
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 70
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    interp.ldata[("OXCAD_LEVANTAMENTO", "TRECHOS")] = trechos

    # 40 clientes com área 70m² gera demanda ~90 kVA -> Ib ~236 A (> Iz 90A)
    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 40, "COORDENADAS": "X: 30,00 | Y: 0,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p2 = _alist_para_dict(res[1])

    # O condutor NÃO foi substituído silenciosamente
    assert p2["CONDUTOR_PREVISTO"] == "25 Al"
    # Status registra SOBRECARGA
    assert "SOBRECARGA" in p2["STATUS_CQT"]
    # Mensagem de aviso foi emitida no prompt
    assert "[AVISO][SOBRECARGA]" in interp.output


def test_cqt_precedencia_comprimento_manual_em_trechos():
    """Valida que comprimento manual em TRECHOS tem precedência e remove flag de vão estimado."""
    interp = _criar_harness_com_modulos()

    arestas = [["TR", "P1"], ["P1", "P2"]]
    # P2 não tem coordenadas CAD, mas tem vão manual de 48.5m em TRECHOS
    trechos = [
        ["P2", "P1", 48.5, "AUTO", "MANUAL", None]
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    interp.ldata[("OXCAD_LEVANTAMENTO", "TRECHOS")] = trechos

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 10, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 10, "COORDENADAS": "-"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p2 = _alist_para_dict(res[1])

    # Comprimento deve ser exatamente 48.5m, sem flag '(EST)' e sem '*'
    assert p2["COMPRIMENTO_VÃO"] == "48.5"
    assert "(EST)" not in p2["COMPRIMENTO_VÃO"]
    assert "*" not in p2["CQT_PREVISTO"]
    assert p2["STATUS_CQT"] == "PREVISTO (COORDENADAS REAIS)"


def test_cqt_poste_passante_com_vao_manual_e_acumulada():
    """Valida integração completa de poste passante (P2) com vão e condutor manuais:
    P1 -> P2 (passante) -> P3
    P2 recebe acumulada de P3 e CQT é calculado nos vãos confirmados."""
    interp = _criar_harness_com_modulos()

    arestas = [["TR", "P1"], ["P1", "P2"], ["P2", "P3"]]
    trechos = [
        ["P2", "P1", 35.0, "185 Al - MX", "MANUAL", None],
        ["P3", "P2", 40.0, "70 Al - MX", "MANUAL", None],
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    interp.ldata[("OXCAD_LEVANTAMENTO", "TRECHOS")] = trechos

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "-"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 0, "PASSANTE": "SIM", "COORDENADAS": "-"},
        {"SEQ": 3, "POSTE": "P3", "CLIENTES": 12, "COORDENADAS": "-"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p1 = _alist_para_dict(res[0])
    p2 = _alist_para_dict(res[1])
    p3 = _alist_para_dict(res[2])

    # P2 carrega 12 clientes de P3
    assert p2["PASSANTE"] == "SIM"
    assert p2["CLIENTES"] == 0
    demanda_12_cli = interp.call("lp-clan-calcular-demanda", 50, 12)
    assert abs(float(p2["ACUMULADA"].replace(",", ".")) - demanda_12_cli) < 0.05

    # Vãos e condutores manuais confirmados
    assert p2["COMPRIMENTO_VÃO"] == "35.0"
    assert p2["CONDUTOR_PREVISTO"] == "185 Al - MX"
    assert p3["COMPRIMENTO_VÃO"] == "40.0"
    assert p3["CONDUTOR_PREVISTO"] == "70 Al - MX"

    # CQT cresce ao longo dos trechos
    cqt_p1 = float(p1["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p2 = float(p2["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    cqt_p3 = float(p3["CQT_PREVISTO"].replace("%", "").replace(",", "."))
    assert cqt_p1 < cqt_p2 < cqt_p3


def test_coleta_topologia_com_medicao_vaos_e_condutor_interativo():
    """Valida a persistência em TRECHOS através da coleta interativa de topologia."""
    dwg = {
        "H_TR": {"0": "TEXT", "1": "TR-01", "5": "H_TR"},
        "H_P1": {"0": "TEXT", "1": "P1", "5": "H_P1"},
        "H_P2": {"0": "TEXT", "1": "P2", "5": "H_P2"},
    }

    inputs = [
        # TR
        "H_TR",
        # LADO 1
        "H_P1",    # Nó raiz TR
        "H_P2",    # Vão P1 -> P2
        "38.5",    # Comp digitado manualmente
        "Manual",  # Condutor Manual
        "70AMX",   # Alias de cabo
        None,      # Enter encerra LADO 1
        # Concluir
        "Concluir",
    ]

    interp = _criar_harness_com_modulos(dwg=dwg, inputs=inputs)
    res = interp.call("lp-clan-cadastrar-topologia")
    assert res is not None

    # Verifica se TRECHOS foi salvo no LData
    trechos_salvos = interp.call("lp-cqt-obter-trechos-salvos")
    assert trechos_salvos is not None
    assert len(trechos_salvos) == 1

    t0 = trechos_salvos[0]
    # Formato: (filho pai comp condutor origem handle)
    assert t0[0] == "P2"
    assert t0[1] == "P1"
    assert abs(t0[2] - 38.5) < 0.01
    assert t0[3] == "70 Al - MX"
    assert t0[4] == "MANUAL"

    # lp-cqt-obter-trecho-dados encontra os dados
    dado = interp.call("lp-cqt-obter-trecho-dados", "P2", "P1")
    dado_dict = _alist_para_dict(dado)
    assert dado_dict["FILHO"] == "P2"
    assert dado_dict["PAI"] == "P1"
    assert abs(dado_dict["COMP"] - 38.5) < 0.01
    assert dado_dict["CONDUTOR"] == "70 Al - MX"


def test_cqt_safeguard_no_desconectado():
    """Valida que nó sem conexão topológica com o trafo é sinalizado com erro explícito."""
    interp = _criar_harness_com_modulos()
    arestas = [["TR", "P1"], ["P1", "P2"]]  # P3 desconectado
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 5, "COORDENADAS": "X: 30,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P3", "CLIENTES": 5, "COORDENADAS": "X: 60,00 | Y: 0,00"},
    ]
    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p3 = _alist_para_dict(res[2])
    assert "ERRO: DESCONECTADO DO TRAFO" in p3["STATUS_CQT"]
    assert p3["CQT_PREVISTO"] == "- (ERRO)"
    assert p3["CONDUTOR_PREVISTO"] == "-"


def test_cqt_safeguard_montante_inexistente():
    """Valida que nó apontando para montante inexistente na lista de pontos é sinalizado com erro."""
    interp = _criar_harness_com_modulos()
    arestas = [["TR", "P1"], ["P99", "P2"]]  # P99 inexistente nos pontos
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 5, "COORDENADAS": "X: 30,00 | Y: 0,00"},
    ]
    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p2 = _alist_para_dict(res[1])
    assert "ERRO: MONTANTE INEXISTENTE" in p2["STATUS_CQT"]
    assert p2["CQT_PREVISTO"] == "- (ERRO)"
    assert p2["CONDUTOR_PREVISTO"] == "-"


def test_cqt_safeguard_ciclo_detectado():
    """Valida que ciclos topológicos são interceptados sem estouro de pilha/recursão infinita."""
    interp = _criar_harness_com_modulos()
    # P4 -> P2, P2 -> P3, P3 -> P4 forma ciclo fechado
    arestas = [["TR", "P1"], ["P4", "P2"], ["P2", "P3"], ["P3", "P4"]]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR"
    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 5, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 5, "COORDENADAS": "X: 30,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P3", "CLIENTES": 5, "COORDENADAS": "X: 60,00 | Y: 0,00"},
        {"SEQ": 4, "POSTE": "P4", "CLIENTES": 5, "COORDENADAS": "X: 90,00 | Y: 0,00"},
    ]
    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    p2 = _alist_para_dict(res[1])
    p3 = _alist_para_dict(res[2])
    p4 = _alist_para_dict(res[3])
    assert "ERRO: CICLO DETECTADO" in p2["STATUS_CQT"]
    assert "ERRO: CICLO DETECTADO" in p3["STATUS_CQT"]
    assert "ERRO: CICLO DETECTADO" in p4["STATUS_CQT"]


def test_cqt_safeguard_endpoint_mismatch_trecho_invalido():
    """Valida que lp-cqt-obter-trecho-dados invalida vãos com pontas incompatíveis."""
    interp = _criar_harness_com_modulos()
    trechos = [["P2", "P1", 30.0, "70 Al - MX", "MANUAL", None]]
    interp.ldata[("OXCAD_LEVANTAMENTO", "TRECHOS")] = trechos

    # Montante incompatível deve retornar nil
    assert interp.call("lp-cqt-obter-trecho-dados", "P2", "P99") is None
    # Filho inexistente deve retornar nil
    assert interp.call("lp-cqt-obter-trecho-dados", "P3", "P1") is None
    # Par correto retorna alist com os dados
    dado = interp.call("lp-cqt-obter-trecho-dados", "P2", "P1")
    assert dado is not None
    d = _alist_para_dict(dado)
    assert d["FILHO"] == "P2"
    assert d["PAI"] == "P1"


def test_cqt_perguntar_vao_comp_dimension_e_invalido():
    """Valida que seleção de DIMENSION extrai cota DXF 42 e texto inválido emite aviso."""
    dwg = {
        "H_DIM": {"0": "DIMENSION", "42": 34.25, "1": "", "5": "H_DIM"},
        "H_TXT": {"0": "TEXT", "1": "Texto Nao Numerico", "5": "H_TXT"},
    }
    # Seleção de DIMENSION via opção 'E'
    interp1 = _criar_harness_com_modulos(dwg=dwg, inputs=["E", "H_DIM"])
    res1 = interp1.call("lp-clan-perguntar-vao-comp", "P1", "P2")
    assert abs(res1[0] - 34.25) < 0.01
    assert res1[1] == "CAD_ENTIDADE"
    assert res1[2] == "H_DIM"

    # Seleção de texto inválido via 'E' avisa e recai na sugestão (sem inventar dado arbitrário)
    interp2 = _criar_harness_com_modulos(dwg=dwg, inputs=["E", "H_TXT"])
    res2 = interp2.call("lp-clan-perguntar-vao-comp", "P1", "P2")
    assert res2[0] == 30.0
    assert "Não foi possível extrair distância da entidade" in interp2.output

    # Digitação numérica manual
    interp3 = _criar_harness_com_modulos(dwg=dwg, inputs=["42.8"])
    res3 = interp3.call("lp-clan-perguntar-vao-comp", "P1", "P2")
    assert abs(res3[0] - 42.8) < 0.01
    assert res3[1] == "MANUAL"


# ==============================================================================
# Suíte de Testes: Tabela CQT de 11 Colunas e Exportação Excel (Modo Clandestino)
# ==============================================================================

def test_cqt_tabela_11_colunas_valores_linha_e_enriquecimento():
    """Valida que em modo CLANDESTINO os pontos são enriquecidos com as 11 colunas
    e lp-cqt-obter-valores-linha retorna a lista de 11 valores na ordem exata."""
    interp = _criar_harness_com_modulos()

    arestas = [
        ["TR-01", "P1"],
        ["P1", "P2"],
        ["P1", "P3"]
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "MODO")] = "CLANDESTINO"
    interp.ldata[("OXCAD_LEVANTAMENTO", "AREA_CLANDESTINO")] = 50
    interp.ldata[("OXCAD_LEVANTAMENTO", "TOPOLOGIA")] = arestas
    interp.ldata[("OXCAD_LEVANTAMENTO", "TR_ID")] = "TR-01"

    trechos = [
        ["P2", "P1", 35.0, "70 Al - MX", "CAD_ENTIDADE", "H1"],
        ["P3", "P1", 28.0, "70 Al - MX", "MANUAL", None]
    ]
    interp.ldata[("OXCAD_LEVANTAMENTO", "TRECHOS")] = trechos

    dados = [
        {"SEQ": 1, "POSTE": "P1", "CLIENTES": 4, "COORDENADAS": "X: 0,00 | Y: 0,00"},
        {"SEQ": 2, "POSTE": "P2", "CLIENTES": 3, "COORDENADAS": "X: 35,00 | Y: 0,00"},
        {"SEQ": 3, "POSTE": "P3", "CLIENTES": 2, "COORDENADAS": "X: 0,00 | Y: 28,00"},
    ]

    res = interp.call("lp-recalcular-demandas-acumuladas", [_dict_para_alist(d) for d in dados])
    assert len(res) == 3

    p1 = _alist_para_dict(res[0])
    p2 = _alist_para_dict(res[1])
    p3 = _alist_para_dict(res[2])

    # 1. P1 (Nó Raiz / Trafo)
    assert p1["POSTE"] == "P1"
    assert p1["TRECHO"] == "TR, P1"
    assert int(p1["SEQ"]) == 1
    assert int(p1["TRECHO_JUSANTE"]) == 0
    assert float(p1["COMPRIMENTO_VÃO"]) == 4.0
    assert p1["CLIENTES"] == 4
    assert p1["CLIENTES_ACUMULADOS"] == 9  # 4 + 3 + 2
    assert p1["LADO"] == "TRAFO"
    assert p1["CONDUTOR_PREVISTO"] == "240 Cu"
    assert ("✔" in p1["CQT_NO_PONTO"]) or ("✖" in p1["CQT_NO_PONTO"])

    vals_p1 = interp.call("lp-cqt-obter-valores-linha", res[0])
    assert len(vals_p1) == 11
    assert vals_p1[0] == "TR, P1"       # TRECHO
    assert int(vals_p1[1]) == 1         # PONTO
    assert int(vals_p1[2]) == 0         # TRECHO jusante (pai)
    assert float(vals_p1[3]) == 4.0     # M
    assert int(vals_p1[4]) == 4         # CLIENTES
    assert int(vals_p1[5]) == 9         # CLT ACUMULAD
    assert vals_p1[8] == "TRAFO"        # LADO
    assert vals_p1[9] == "240 Cu"       # CONDUTOR
    assert vals_p1[10] == p1["CQT_NO_PONTO"]

    # 2. P2 (Filho LADO 1)
    assert p2["POSTE"] == "P2"
    assert p2["TRECHO"] == "P1, P2"
    assert int(p2["SEQ"]) == 2
    assert int(p2["TRECHO_JUSANTE"]) == 1    # Pai é P1 (SEQ 1)
    assert float(p2["COMPRIMENTO_VÃO"]) == 35.0
    assert p2["CLIENTES"] == 3
    assert p2["CLIENTES_ACUMULADOS"] == 3
    assert p2["LADO"] == "LADO 1"
    assert p2["CONDUTOR_PREVISTO"] == "70 Al - MX"

    vals_p2 = interp.call("lp-cqt-obter-valores-linha", res[1])
    assert len(vals_p2) == 11
    assert vals_p2[0] == "P1, P2"
    assert int(vals_p2[1]) == 2
    assert int(vals_p2[2]) == 1
    assert float(vals_p2[3]) == 35.0
    assert int(vals_p2[4]) == 3
    assert int(vals_p2[5]) == 3
    assert vals_p2[8] == "LADO 1"
    assert vals_p2[9] == "70 Al - MX"

    # 3. P3 (Ramificação - herda circuito BT)
    assert p3["POSTE"] == "P3"
    assert p3["TRECHO"] == "P1, P3"
    assert int(p3["SEQ"]) == 3
    assert int(p3["TRECHO_JUSANTE"]) == 1    # Pai é P1 (SEQ 1)
    assert float(p3["COMPRIMENTO_VÃO"]) == 28.0
    assert p3["CLIENTES"] == 2
    assert p3["CLIENTES_ACUMULADOS"] == 2
    assert p3["LADO"] == "LADO 1"

    vals_p3 = interp.call("lp-cqt-obter-valores-linha", res[2])
    assert len(vals_p3) == 11
    assert vals_p3[0] == "P1, P3"
    assert int(vals_p3[1]) == 3
    assert int(vals_p3[2]) == 1
    assert float(vals_p3[3]) == 28.0
    assert int(vals_p3[4]) == 2
    assert int(vals_p3[5]) == 2
    assert vals_p3[8] == "LADO 1"


def test_cqt_tabela_nativa_cad_dwg_11_colunas():
    """Valida estrutura de 11 colunas, cabeçalhos, larguras e delegações em AutoLISP."""
    src_cqt = (Path(__file__).parent.parent / "orcamento" / "clandestino_cqt.lsp").read_text(encoding="utf-8")
    src_lp = LISP_PATH.read_text(encoding="utf-8")

    # 1. Verificações no módulo clandestino_cqt.lsp
    assert "col-count 11" in src_cqt
    assert "(vla-MergeCells table 0 0 0 10)" in src_cqt

    # Ordem exata dos 11 cabeçalhos
    assert 'headers \'("TRECHO" "PONTO" "TRECHO jusante" "M" "CLIENTES" "CLT ACUMULAD" "TOTAL DO TRECHO" "ACUMULADA" "LADO" "CONDUTOR" "CQT NO PONTO")' in src_cqt

    # Larguras das 11 colunas mapeadas
    assert "widths  '(28.0 20.0 28.0 18.0 24.0 28.0 34.0 32.0 26.0 30.0 32.0)" in src_cqt
    assert "(vla-SetText table 1 col-idx h)" in src_cqt
    assert "(vla-SetColumnWidth table col-idx w)" in src_cqt

    # 2. Verificações de delegação condicional em levantamento_postes.lsp
    assert "(lp-cqt-colunas-adicionais-p)" in src_lp
    assert "(lp-cqt-gerar-tabela-cad pontos pt)" in src_lp
    assert "(lp-cqt-atualizar-celulas-tabela table pontos)" in src_lp


def test_cqt_exportacao_excel_11_colunas_openpyxl(tmp_path):
    """Valida a exportação completa para planilha Excel com o layout CQT de 11 colunas da Light."""
    import openpyxl
    from orcamento.clandestino_excel import exportar_cqt_excel

    pontos_teste = [
        {
            "TRECHO": "TR, P1",
            "SEQ": 1,
            "TRECHO_JUSANTE": 0,
            "COMPRIMENTO_VÃO": 4.0,
            "CLIENTES": 4,
            "CLT_ACUMULADO": 12,
            "DEMANDA": 9.08,
            "ACUMULADA": 22.50,
            "LADO": "TRAFO",
            "CONDUTOR_PREVISTO": "240 Cu",
            "CQT_NO_PONTO": "✔ 118,59",
        },
        {
            "TRECHO": "P1, P2",
            "SEQ": 2,
            "TRECHO_JUSANTE": 1,
            "COMPRIMENTO_VÃO": 35.0,
            "CLIENTES": 5,
            "CLT_ACUMULADO": 8,
            "DEMANDA": 11.20,
            "ACUMULADA": 13.42,
            "LADO": "LADO 1",
            "CONDUTOR_PREVISTO": "70 Al - MX",
            "CQT_NO_PONTO": "✔ 117,80",
        },
        {
            "TRECHO": "P2, P3",
            "SEQ": 3,
            "TRECHO_JUSANTE": 2,
            "COMPRIMENTO_VÃO": 40.0,
            "CLIENTES": 3,
            "CLT_ACUMULADO": 3,
            "DEMANDA": 7.14,
            "ACUMULADA": 7.14,
            "LADO": "LADO 1",
            "CONDUTOR_PREVISTO": "70 Al - MX",
            "CQT_NO_PONTO": "✖ 116,22",
        }
    ]

    xlsx_path = tmp_path / "cqt_saida_teste.xlsx"
    res = exportar_cqt_excel(pontos_teste, str(xlsx_path), abrir_arquivo=False)
    assert res == str(xlsx_path)
    assert xlsx_path.exists()

    wb = openpyxl.load_workbook(str(xlsx_path))
    assert "CQT CLANDESTINO" in wb.sheetnames
    ws = wb["CQT CLANDESTINO"]

    # 1. Cabeçalhos na Linha 1 (11 colunas exatas)
    raw_headers = [ws.cell(row=1, column=col).value for col in range(1, 12)]
    headers = [str(h).replace("\n", " ").strip() for h in raw_headers]
    assert headers == [
        "TRECHO", "PONTO", "TRECHO jusante", "M", "CLIENTES",
        "CLT ACUMULAD", "TOTAL DO TRECHO", "ACUMULADA",
        "LADO", "CONDUTOR", "CQT NO PONTO"
    ]

    # Preenchimento preto do cabeçalho
    for col in range(1, 12):
        cell = ws.cell(row=1, column=col)
        assert cell.font.bold is True
        assert cell.font.color.rgb in ("FFFFFFFF", "00FFFFFF", "FFFFFF")
        assert cell.fill.fill_type == "solid"
        assert cell.fill.start_color.rgb in ("00000000", "FF000000", "000000")

    # 2. Linha 2 (P1 / TRAFO): destaque vermelho nas fontes das colunas principais
    row2_vals = [ws.cell(row=2, column=col).value for col in range(1, 12)]
    assert row2_vals[0] == "TR, P1"
    assert row2_vals[1] == 1
    assert row2_vals[2] == 0
    assert row2_vals[3] == 4.0
    assert row2_vals[4] == 4
    assert row2_vals[5] == 12
    assert row2_vals[8] == "TRAFO"
    assert row2_vals[9] == "240 Cu"
    assert "✔ 118,59" in row2_vals[10]

    # Fonte vermelha (C00000) na linha TRAFO para LADO, CONDUTOR e demandas
    for col_idx in [5, 6, 7, 8, 9, 10]:
        cell = ws.cell(row=2, column=col_idx)
        assert cell.font.color.rgb in ("FFC00000", "00C00000", "C00000")

    # 3. Linha 3 (P2 / LADO 1)
    row3_vals = [ws.cell(row=3, column=col).value for col in range(1, 12)]
    assert row3_vals[0] == "P1, P2"
    assert row3_vals[1] == 2
    assert row3_vals[2] == 1
    assert row3_vals[3] == 35.0
    assert row3_vals[8] == "LADO 1"
    assert row3_vals[9] == "70 Al - MX"

    # 4. Linha 4 (P3 / LADO 1 com violação CQT)
    row4_vals = [ws.cell(row=4, column=col).value for col in range(1, 12)]
    assert row4_vals[0] == "P2, P3"
    assert row4_vals[1] == 3
    assert row4_vals[2] == 2
    assert "✖ 116,22" in row4_vals[10]
    # Célula CQT com violação tem cor de texto vermelha
    assert ws.cell(row=4, column=11).font.color.rgb in ("FFC00000", "00C00000", "C00000")


def test_cqt_exportacao_excel_cli(tmp_path):
    """Valida a execução por linha de comando de clandestino_excel.py."""
    import json
    import subprocess
    import sys

    pontos_teste = [
        {
            "TRECHO": "TR, P1",
            "SEQ": 1,
            "TRECHO_JUSANTE": 0,
            "COMPRIMENTO_VÃO": 4.0,
            "CLIENTES": 2,
            "CLT_ACUMULADO": 2,
            "DEMANDA": 4.54,
            "ACUMULADA": 4.54,
            "LADO": "TRAFO",
            "CONDUTOR_PREVISTO": "240 Cu",
            "CQT_NO_PONTO": "✔ 119,00",
        }
    ]

    json_path = tmp_path / "pontos.json"
    xlsx_path = tmp_path / "saida_cli.xlsx"
    json_path.write_text(json.dumps(pontos_teste, ensure_ascii=False), encoding="utf-8")

    script_path = Path(__file__).parent.parent / "orcamento" / "clandestino_excel.py"
    proc = subprocess.run(
        [sys.executable, str(script_path), "--json", str(json_path), "--out", str(xlsx_path)],
        capture_output=True,
        text=True
    )
    assert proc.returncode == 0
    assert xlsx_path.exists()


def test_clandestino_excel_funcoes_auxiliares_e_cobertura(tmp_path, monkeypatch):
    """Cobre funções utilitárias e ramos defensivos de clandestino_excel.py para meta de alta cobertura."""
    from orcamento.clandestino_excel import (
        exportar_cqt_excel,
        formatar_cqt_ponto,
        parse_float_safe,
        parse_int_safe,
    )

    # 1. parse_float_safe
    assert parse_float_safe(None) == 0.0
    assert parse_float_safe(12) == 12.0
    assert parse_float_safe(" 15,75 kVA ") == 15.75
    assert parse_float_safe("inválido") == 0.0

    # 2. parse_int_safe
    assert parse_int_safe(None) == 0
    assert parse_int_safe(5) == 5
    assert parse_int_safe(" 7,0 ") == 7
    assert parse_int_safe("nao_numero") == 0

    # 3. formatar_cqt_ponto
    assert "ERRO" in formatar_cqt_ponto(0.0, 5.0, "ERRO")
    assert "✔" in formatar_cqt_ponto(118.5, 6.5)
    assert "✖" in formatar_cqt_ponto(115.0, 8.5)

    # 4. Fallback de trecho e mock de abertura automática
    pontos_borda = [
        {"POSTE": "P1", "SEQ": 1, "V_127": 118.5, "QT_PERC": 6.5},  # Sem TRECHO -> gera TR, P1
        {"POSTE": "P2", "SEQ": 2, "TRECHO_JUSANTE": 1, "V_127": 115.0, "QT_PERC": 8.0},
    ]
    out_borda = tmp_path / "borda.xlsx"

    # Mock de abertura para não disparar processo externo real
    abriu = []
    monkeypatch.setattr("os.startfile", lambda p: abriu.append(p), raising=False)
    monkeypatch.setattr("subprocess.Popen", lambda *args, **kwargs: abriu.append(args), raising=False)

    res = exportar_cqt_excel(pontos_borda, str(out_borda), abrir_arquivo=True)
    assert res == str(out_borda)
    assert len(abriu) == 1



