"""Auditoria P0/P1: isolamento A/B sem alterar comportamento de produto."""

import json
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from workspace.workspace_context_cli import (
    _executar_acao_operacional_impl,
    executar_acao_operacional,
    extrair_contexto_operacional,
)
from workspace.workspace_actions import WorkspaceActionDispatcher
from workspace.workspace_delivery import approve, generate, verify
from workspace.workspace_pending import create, get, list_pending, update
from workspace.workspace_validation import ValidationContext, ValidationEngine
from workspace.workspace_versions import capture, inspect, restore


def _project(parent: Path, sob_id: str) -> ValidationContext:
    root = parent / f"SOB_{sob_id}"
    root.mkdir()
    dwg = root / f"PE_{sob_id}_DESENHO.dwg"
    dwg.write_bytes(b"synthetic DWG for isolation test")
    (root / ".sob_state.json").write_text(
        json.dumps({
            "sob_id": sob_id, "status": "CONCLUÍDO", "base_dwg": str(dwg)
        }),
        encoding="utf-8",
    )
    workbook_path = root / f"CQT_{sob_id}.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "CQT ATUAL"
    sheet["A11"] = "test-only segment"
    sheet["C11"] = 1.0
    sheet["L11"] = 6.2
    workbook.save(workbook_path)
    workbook.close()
    return ValidationContext(
        execution_id=f"exec-{sob_id}",
        correlation_id=f"corr-{sob_id}",
        sob_id=sob_id,
        dwg_path=str(dwg),
        dwg_identity=f"fp-{sob_id}",
        database_fingerprint=f"fp-{sob_id}",
        project_root=root,
        cqt_path=workbook_path,
    )


def test_ab_isolation_across_p0_and_all_p1_contracts(tmp_path: Path):
    context_a = _project(tmp_path, "A057349839")
    context_b = _project(tmp_path, "B123456789")

    resolved_a = extrair_contexto_operacional(
        dwg_path=context_a.dwg_path,
        sob_id_esperada=context_a.sob_id,
        require_project_scope=True,
    )
    resolved_b = extrair_contexto_operacional(
        dwg_path=context_b.dwg_path,
        sob_id_esperada=context_b.sob_id,
        require_project_scope=True,
    )
    assert resolved_a["project_scope_status"] == "Unverified"
    assert resolved_b["project_scope_status"] == "Unverified"
    assert resolved_a["artefatos"] == resolved_b["artefatos"] == []

    results_a = [
        item.to_dict() for item in ValidationEngine().validate(context_a)
    ]
    results_b = [
        item.to_dict() for item in ValidationEngine().validate(context_b)
    ]
    assert {item["sob_id"] for item in results_a} == {context_a.sob_id}
    assert {item["sob_id"] for item in results_b} == {context_b.sob_id}
    assert {item["execution_id"] for item in results_a} == {
        context_a.execution_id
    }
    assert {item["execution_id"] for item in results_b} == {
        context_b.execution_id
    }

    snapshot_a = capture(context_a, context_a.cqt_path.name)["snapshot_id"]
    assert inspect(context_b, snapshot_a)["status"] == "MISSING_SNAPSHOT"
    assert restore(context_b, snapshot_a)["status"] == "MISSING_SNAPSHOT"
    assert not list(context_b.project_root.glob("*.recuperado-*"))

    pending_a = create(
        context_a, "Synthetic check", "ALERTA", validation=results_a[0]
    )
    assert list_pending(context_b) == []
    with pytest.raises(ValueError, match="PENDING_NOT_FOUND"):
        get(context_b, pending_a["id"])
    with pytest.raises(ValueError, match="PENDING_NOT_FOUND"):
        update(context_b, pending_a["id"], status="RESOLVIDA")
    assert get(context_a, pending_a["id"])["status"] == "ABERTA"

    package_a = generate(
        context_a, [context_a.cqt_path.name], results_a, [snapshot_a]
    )
    package_id = package_a["package_id"]
    assert verify(context_b, package_id)["status"] == "PACKAGE_NOT_FOUND"
    with pytest.raises(FileNotFoundError):
        approve(context_b, package_id, confirmed=True)
    assert verify(context_a, package_id)["status"] == "UNAPPROVED"

    blocked = executar_acao_operacional(
        "versao_capturar",
        dwg_path=context_b.dwg_path,
        sob_id_esperada=context_a.sob_id,
        dwg_identity=context_a.dwg_identity,
        execution_id=context_a.execution_id,
        correlation_id=context_a.correlation_id,
        alvo=context_b.cqt_path.name,
    )
    assert blocked["sucesso"] is False
    assert blocked["estado"] == "Requer especificação"
    assert "execution_id" not in blocked
    assert not (context_b.project_root / ".auxcad_versions").exists()


def test_same_project_stale_dwg_cannot_mutate_or_approve(tmp_path: Path):
    context = _project(tmp_path, "A057349839")
    results = [item.to_dict() for item in ValidationEngine().validate(context)]
    snapshot_id = capture(context, context.cqt_path.name)["snapshot_id"]
    pending = create(context, "Synthetic check", "ALERTA")
    package_id = generate(
        context, [context.cqt_path.name], results, [snapshot_id]
    )["package_id"]
    second_dwg = context.project_root / "SECOND.dwg"
    second_dwg.write_bytes(b"different synthetic DWG")
    stale = replace(
        context, dwg_path=str(second_dwg),
        dwg_identity="fp-other", database_fingerprint="fp-other",
    )
    assert inspect(stale, snapshot_id)["status"] == "CONTEXT_MISMATCH"
    with pytest.raises(ValueError, match="CONTEXT_MISMATCH"):
        update(stale, pending["id"], status="RESOLVIDA")
    assert verify(stale, package_id)["status"] == "CONTEXT_MISMATCH"
    with pytest.raises(ValueError, match="CONTEXT_MISMATCH"):
        approve(stale, package_id, confirmed=True)
    assert get(context, pending["id"])["status"] == "ABERTA"
    assert verify(context, package_id)["status"] == "UNAPPROVED"


def test_audit_characterizes_unbound_scope_fallback(tmp_path: Path):
    context = _project(tmp_path, "A057349839")
    (context.project_root / ".sob_state.json").unlink()

    observed = executar_acao_operacional(
        "versao_capturar",
        dwg_path=context.dwg_path,
        sob_id_esperada="B123456789",
        dwg_identity="caller-supplied-fingerprint",
        alvo=context.cqt_path.name,
    )

    assert observed["sucesso"] is False
    assert observed["project_scope_status"] == "Unverified"
    assert observed["erro"] == "DWG_CONTEXT_NOT_VERIFIED"
    assert "version_result" not in observed
    assert not (context.project_root / ".auxcad_versions").exists()


def test_audit_characterizes_optional_validation_scope(tmp_path: Path):
    context = _project(tmp_path, "A057349839")

    observed = executar_acao_operacional(
        "validar",
        dwg_path=context.dwg_path,
        sob_id_esperada="B123456789",
        dwg_identity="caller-supplied-fingerprint",
    )

    assert observed["sucesso"] is False
    assert observed["project_scope_status"] == "Unverified"
    assert observed["erro"] == "DWG_CONTEXT_NOT_VERIFIED"
    assert "validation_results" not in observed


def test_audit_characterizes_stale_validation_in_approved_package(
    tmp_path: Path
):
    context = _project(tmp_path, "A057349839")
    workbook = openpyxl.load_workbook(context.cqt_path)
    workbook["CQT ATUAL"]["L11"] = 5.0
    workbook.save(context.cqt_path)
    workbook.close()
    previous = [
        item.to_dict() for item in ValidationEngine().validate(context)
    ]
    assert previous[0]["status"] == "OK"

    workbook = openpyxl.load_workbook(context.cqt_path)
    workbook["CQT ATUAL"]["L11"] = 7.0
    workbook.save(context.cqt_path)
    workbook.close()
    current = [item.to_dict() for item in ValidationEngine().validate(context)]
    assert current[0]["status"] == "ALERTA"

    with pytest.raises(ValueError, match="VALIDATION_SOURCE_CHANGED"):
        generate(context, [context.cqt_path.name], previous)
    assert not list((context.project_root / ".oxcad" / "deliveries").glob("*/approval.json"))


def test_audit_characterizes_exit_zero_without_generated_artifact(
    tmp_path: Path, monkeypatch
):
    context = _project(tmp_path, "A057349839")

    def empty_process(*args, **kwargs):
        return subprocess.Popen([sys.executable, "-c", "pass"])

    monkeypatch.setattr(
        WorkspaceActionDispatcher, "executar_cqt", empty_process
    )
    result = _executar_acao_operacional_impl(
        "cqt", dwg_path=context.dwg_path,
        sob_id_esperada=context.sob_id, dwg_identity=context.dwg_identity,
        require_project_scope=True, wait_for_completion=True,
    )
    assert result["estado"] == "Falhou"
    assert result["erro"] == "ARTIFACT_NOT_VERIFIED"


def test_audit_characterizes_missing_checksum_as_missing_package(
    tmp_path: Path
):
    context = _project(tmp_path, "A057349839")
    results = [item.to_dict() for item in ValidationEngine().validate(context)]
    package = generate(context, [context.cqt_path.name], results)
    package_id = package["package_id"]
    digest = (
        context.project_root / ".oxcad" / "deliveries" / package_id
        / "manifest.sha256"
    )
    digest.unlink()
    checked = verify(context, package_id)
    assert checked["status"] == "PACKAGE_NOT_FOUND"
    assert checked["approval_valid"] is False
