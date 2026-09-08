import os
import pytest
from cqt_excel_io import (
    load_cables_from_database,
    load_scenario_from_excel,
    update_cqt_excel,
    DEFAULT_TEMPLATE_PATH,
)


def test_load_cables_fallback():
    cables = load_cables_from_database("caminho_inexistente_123.xlsm")
    assert isinstance(cables, list)
    assert len(cables) > 0
    assert "3x150+70 - Al - F.P 1" in cables
    assert "3x35+1x54,6 - Al - F.P 1" in cables


def test_load_scenario_nonexistent():
    data = load_scenario_from_excel("inexistente.xlsm", "ATUAL")
    assert data == {}


def test_update_cqt_excel_not_found():
    with pytest.raises(FileNotFoundError):
        update_cqt_excel("modelo_inexistente.xlsm", "dest.xlsm", {})


def test_cqt_cad_input_json_interchange(tmp_path):
    """Testa a geração e parsing do JSON de intercâmbio do comando CQTCAD."""
    import json

    json_file = tmp_path / "cqt_cad_input.json"
    dados_cad = {
        "projeto": "A057100646",
        "ponto_trafo": "A55323",
        "potencia_trafo": "45",
        "cenario": "PROJETADO",
        "salvar_em": "C:/Projetos/CQT_A057100646.xlsm",
        "trechos": [
            {"letra": "A", "montante": "TRAFO", "vaos": 2, "comprimento_m": 75.50},
            {"letra": "B", "montante": "A", "vaos": 1, "comprimento_m": 38.20}
        ]
    }
    json_file.write_text(json.dumps(dados_cad, ensure_ascii=False, indent=2), encoding="utf-8")

    assert os.path.exists(json_file)
    with open(json_file, "r", encoding="utf-8") as f:
        carregado = json.load(f)

    assert carregado["projeto"] == "A057100646"
    assert carregado["ponto_trafo"] == "A55323"
    assert carregado["cenario"] == "PROJETADO"
    assert len(carregado["trechos"]) == 2
    assert carregado["trechos"][0]["letra"] == "A"
    assert carregado["trechos"][0]["comprimento_m"] == 75.50


def test_load_cables_from_valid_excel(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BASE DE DADOS"
    ws.cell(4, 1, "3x70+1x70 - Al - F.P 1")
    ws.cell(5, 1, "3x95+1x54,6 - Al - F.P 1")
    fpath = str(tmp_path / "cables_test.xlsx")
    wb.save(fpath)

    cables = load_cables_from_database(fpath)
    assert "3x70+1x70 - Al - F.P 1" in cables
    assert "3x95+1x54,6 - Al - F.P 1" in cables


def test_load_scenario_and_update_excel(tmp_path):
    import openpyxl

    # 1. Cria template mínimo
    template_path = str(tmp_path / "template_cqt.xlsx")
    dest_path = str(tmp_path / "output_cqt.xlsx")

    wb = openpyxl.Workbook()
    ws_proj = wb.active
    ws_proj.title = "PROJETADO"
    ws_cqt = wb.create_sheet("CQT PROJETADO")

    # Estrutura esperada
    ws_cqt.cell(1, 2, "PROJ_ORIGINAL")
    ws_cqt.cell(4, 6, "TRAFO_ORIG")
    ws_proj.cell(3, 3, "CLASSE B")
    ws_proj.cell(10, 3, 30.0)

    # Trecho de exemplo
    ws_cqt.cell(11, 1, "A")
    ws_cqt.cell(11, 2, "TRAFO")
    ws_cqt.cell(11, 3, 0.50)  # 50m
    ws_cqt.cell(11, 9, "3x70+1x70 - Al - F.P 1")
    ws_proj.cell(3, 12, 5)  # mono

    wb.save(template_path)

    # 2. Testa leitura do cenário
    scen_data = load_scenario_from_excel(template_path, "PROJETADO")
    assert scen_data["projeto"] == "PROJ_ORIGINAL"
    assert scen_data["ponto_trafo"] == "TRAFO_ORIG"
    assert len(scen_data["trechos"]) >= 1
    assert scen_data["trechos"][0]["letra"] == "A"
    assert scen_data["trechos"][0]["comprimento_m"] == 50.0

    # 3. Testa gravação de novos dados via update_cqt_excel
    novos_dados = {
        "projeto": "PROJ_NOVO_99",
        "ponto_trafo": "TRAFO_99",
        "classe": "CLASSE C",
        "potencia_trafo": "45",
        "scenarios": {
            "PROJETADO": [
                {
                    "letra": "A",
                    "montante": "TRAFO",
                    "comprimento_m": 80.0,
                    "condutor": "3x70+1x70 - Al - F.P 1",
                    "mono": 3,
                    "bi": 2,
                    "tri": 1,
                    "ip70": 2,
                }
            ]
        }
    }

    res = update_cqt_excel(template_path, dest_path, novos_dados)
    assert "PROJETADO" in res
    assert os.path.exists(dest_path)

    # 4. Verifica se gravou corretamente no destino
    wb_dest = openpyxl.load_workbook(dest_path)
    assert wb_dest["CQT PROJETADO"].cell(1, 2).value == "PROJ_NOVO_99"
    assert wb_dest["CQT PROJETADO"].cell(4, 6).value == "TRAFO_99"
    assert wb_dest["CQT PROJETADO"].cell(11, 3).value == 0.80  # 80m / 100
    assert wb_dest["PROJETADO"].cell(3, 12).value == 3


