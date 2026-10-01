import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from common.logging_config import configurar_logging
from workspace.workspace_context_cli import executar_acao_operacional
from workspace.workspace_delivery import approve, generate, revoke, verify
from workspace.workspace_pending import create, update
from workspace.workspace_validation import ValidationContext, ValidationEngine
from workspace.workspace_versions import capture


def project(parent: Path, sob: str) -> tuple[ValidationContext, Path, Path]:
    root = parent / f"SOB_{sob}"
    root.mkdir()
    dwg = root / f"PE_{sob}_DESENHO.dwg"
    dwg.write_bytes(b"synthetic DWG context for test")
    (root / ".sob_state.json").write_text(
        json.dumps({
            "sob_id": sob, "status": "CONCLUÍDO", "base_dwg": str(dwg)
        }),
        encoding="utf-8",
    )
    cqt = root / f"CQT_{sob}.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "CQT ATUAL"
    sheet["A11"] = "test segment"
    sheet["C11"] = 1.0
    sheet["L11"] = 6.25
    workbook.save(cqt)
    workbook.close()
    return ValidationContext(
        execution_id=f"exec-{sob}", correlation_id=f"corr-{sob}",
        sob_id=sob, dwg_path=str(dwg), dwg_identity=f"fp-{sob}",
        database_fingerprint=f"fp-{sob}", project_root=root,
        cqt_path=cqt, origin="oxcad.palette",
    ), dwg, cqt


def validation(context: ValidationContext) -> list[dict]:
    return [item.to_dict() for item in ValidationEngine().validate(context)]


def test_complete_cycle_from_sob_through_divergent_delivery(tmp_path: Path):
    context, dwg, cqt = project(tmp_path, "A057349839")
    result = validation(context)
    assert result[0]["status"] == "ALERTA"
    pending = create(
        context, "Verificar queda calculada", "ALERTA", validation=result[0]
    )
    snapshot_id = capture(context, cqt.name)["snapshot_id"]
    package = generate(context, [cqt.name], result, [snapshot_id])
    manifest = package["manifest"]
    assert package["status"] == "UNAPPROVED"
    assert verify(context, package["package_id"])["approval_valid"] is False
    assert [item["path"] for item in manifest["files"]] == sorted(
        [cqt.name, dwg.name]
    )
    assert manifest["schema"] == 1
    assert manifest["revision"] is None
    assert manifest["dwg_identity"] == context.dwg_identity
    assert manifest["execution_id"] == context.execution_id
    assert manifest["correlation_id"] == context.correlation_id
    assert manifest["validations"][0]["status"] == "ALERTA"
    assert manifest["pending"][0]["id"] == pending["id"]
    assert manifest["snapshots"][0]["snapshot_id"] == snapshot_id
    assert manifest["files"][0]["sha256"] == hashlib.sha256(
        cqt.read_bytes()
    ).hexdigest()
    with pytest.raises(ValueError, match="ISSUES_REQUIRE_DECISION"):
        approve(context, package["package_id"], confirmed=True)
    approval = approve(
        context, package["package_id"], confirmed=True,
        acknowledge_issues=True,
        decision_note="Revisão consciente de alerta de teste",
    )
    assert approval["approver"] is None
    assert approval["issues"]["open_pending"] == [pending["id"]]
    assert verify(context, package["package_id"])["status"] == "CONFIRMED"
    cqt.write_bytes(b"x" * cqt.stat().st_size)
    divergence = verify(context, package["package_id"])
    assert divergence["status"] == "DIVERGENT"
    assert divergence["approval_valid"] is False
    assert any(item["path"] == cqt.name and item["status"] == "DIVERGENT"
               for item in divergence["files"])


def test_empty_incomplete_or_wrong_validation_rejected(tmp_path: Path):
    context, dwg, cqt = project(tmp_path, "A057349839")
    results = validation(context)
    for paths, values in (([], results), ([dwg.name], results),
                          ([cqt.name], [])):
        with pytest.raises(ValueError, match="PACKAGE_INCOMPLETE"):
            generate(context, paths, values)
    other, _, _ = project(tmp_path, "B123456789")
    with pytest.raises(ValueError, match="VALIDATION_CONTEXT_MISMATCH"):
        generate(context, [cqt.name], validation(other))
    with pytest.raises(ValueError, match="SNAPSHOT_INCOMPATIBLE"):
        generate(context, [cqt.name], results, ["0" * 32])
    assert not (context.project_root / ".oxcad" / "deliveries").exists()


def test_missing_file_and_inconsistent_manifest_or_approval(tmp_path: Path):
    context, _, cqt = project(tmp_path, "A057349839")
    package = generate(context, [cqt.name], validation(context))
    package_id = package["package_id"]
    approve(context, package_id, confirmed=True)
    cqt.unlink()
    missing = verify(context, package_id)
    assert missing["status"] == "MISSING"
    assert missing["approval_valid"] is False
    path = (
        context.project_root / ".oxcad" / "deliveries" / package_id
        / "manifest.json"
    )
    path.write_text("{broken", encoding="utf-8")
    assert verify(context, package_id)["status"] == "INVALID"


def test_approval_identity_tampering_is_invalid(tmp_path: Path):
    context, _, cqt = project(tmp_path, "A057349839")
    package = generate(context, [cqt.name], validation(context))
    package_id = package["package_id"]
    approve(context, package_id, confirmed=True)
    path = (
        context.project_root / ".oxcad" / "deliveries" / package_id
        / "approval.json"
    )
    data = json.loads(path.read_text(encoding="utf-8"))
    data["database_fingerprint"] = "another-dwg"
    path.write_text(json.dumps(data), encoding="utf-8")
    checked = verify(context, package_id)
    assert checked["status"] == "INVALID"
    assert checked["approval_valid"] is False


def test_approval_is_explicit_and_failed_write_never_approves(
    tmp_path: Path, monkeypatch
):
    context, _, cqt = project(tmp_path, "A057349839")
    package = generate(context, [cqt.name], validation(context))
    package_id = package["package_id"]
    with pytest.raises(ValueError, match="APPROVAL_NOT_CONFIRMED"):
        approve(context, package_id)
    from workspace import workspace_delivery as delivery

    def fail_write(path, data):
        raise OSError("simulated write failure")

    monkeypatch.setattr(delivery, "_write_approval", fail_write)
    with pytest.raises(OSError, match="simulated write failure"):
        approve(context, package_id, confirmed=True)
    assert verify(context, package_id)["status"] == "UNAPPROVED"
    assert not (context.project_root / ".oxcad" / "deliveries"
                / package_id / "approval.json").exists()


def test_two_projects_and_stale_dwg_never_share_approval(tmp_path: Path):
    first, _, cqt = project(tmp_path, "A057349839")
    second, _, _ = project(tmp_path, "B123456789")
    package_id = generate(first, [cqt.name], validation(first))["package_id"]
    with ThreadPoolExecutor(max_workers=2) as pool:
        valid = pool.submit(verify, first, package_id)
        foreign = pool.submit(verify, second, package_id)
        assert valid.result()["status"] == "UNAPPROVED"
        assert foreign.result()["status"] == "PACKAGE_NOT_FOUND"
    with pytest.raises(FileNotFoundError):
        approve(second, package_id, confirmed=True)
    with pytest.raises(ValueError, match="CONTEXT_MISMATCH"):
        approve(
            replace(first, dwg_identity="stale"), package_id,
            confirmed=True,
        )
    with pytest.raises(ValueError, match="CONTEXT_MISMATCH"):
        approve(replace(first, database_fingerprint="stale"), package_id,
                confirmed=True)
    with pytest.raises(ValueError, match="DWG_MISMATCH"):
        approve(replace(first, dwg_path=None), package_id, confirmed=True)
    assert verify(first, package_id)["status"] == "UNAPPROVED"


def test_pending_change_invalidates_approval_and_revocation(tmp_path: Path):
    context, _, cqt = project(tmp_path, "A057349839")
    pending = create(context, "Revisar", "ALERTA")
    package = generate(context, [cqt.name], validation(context))
    package_id = package["package_id"]
    approve(context, package_id, confirmed=True, acknowledge_issues=True,
            decision_note="Decisão explícita sobre pendência de teste")
    update(context, pending["id"], "RESOLVIDA")
    assert verify(context, package_id)["status"] == "DIVERGENT"
    revoked = revoke(context, package_id, "Mudança do estado aprovado")
    assert revoked["state"] == "REVOGADO"
    assert revoked["revocation"]["execution_id"] == context.execution_id
    assert verify(context, package_id)["approval_valid"] is False
    with pytest.raises(ValueError, match="PACKAGE_NOT_APPROVABLE"):
        approve(context, package_id, confirmed=True)


def test_error_requires_acknowledgement_but_not_automatic_rejection(
    tmp_path: Path
):
    context, _, cqt = project(tmp_path, "A057349839")
    workbook = openpyxl.load_workbook(cqt)
    workbook["CQT ATUAL"]["L11"] = "invalid test value"
    workbook.save(cqt)
    workbook.close()
    results = validation(context)
    assert results[0]["status"] == "ERRO"
    package_id = generate(context, [cqt.name], results)["package_id"]
    with pytest.raises(ValueError, match="ISSUES_REQUIRE_DECISION"):
        approve(context, package_id, confirmed=True)
    approved = approve(
        context, package_id, confirmed=True,
        acknowledge_issues=True, decision_note="Revisão explícita",
    )
    assert approved["issues"]["errors"] == [results[0]["rule_id"]]
    assert verify(context, package_id)["status"] == "CONFIRMED"


def test_ok_and_unavailable_remain_distinct_without_inventing_results(
    tmp_path: Path
):
    context, _, cqt = project(tmp_path, "A057349839")
    workbook = openpyxl.load_workbook(cqt)
    workbook["CQT ATUAL"]["L11"] = 4.0
    workbook.save(cqt)
    workbook.close()
    ok = validation(context)
    assert ok[0]["status"] == "OK"
    ok_package = generate(context, [cqt.name], ok)
    assert ok_package["manifest"]["validations"][0]["status"] == "OK"

    unavailable = validation(
        replace(context, cqt_path=cqt.with_name("missing.xlsx"))
    )
    assert unavailable[0]["status"] == "NÃO DISPONÍVEL"
    package = generate(context, [cqt.name], unavailable)
    assert package["manifest"]["validations"][0]["status"] == "NÃO DISPONÍVEL"
    assert verify(context, package["package_id"])["status"] == "UNAPPROVED"


def test_missing_dwg_and_preapproval_checksum_corruption(tmp_path: Path):
    context, dwg, cqt = project(tmp_path, "A057349839")
    package = generate(context, [cqt.name], validation(context))
    package_id = package["package_id"]
    manifest_path = (
        context.project_root / ".oxcad" / "deliveries" / package_id
        / "manifest.json"
    )
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["files"][0]["size"] += 1
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    assert verify(context, package_id)["status"] == "INVALID"
    with pytest.raises(ValueError, match="PACKAGE_INVALID"):
        approve(context, package_id, confirmed=True)
    manifest_path.unlink()
    assert verify(context, package_id)["status"] == "PACKAGE_NOT_FOUND"

    other_id = generate(context, [cqt.name], validation(context))["package_id"]
    approve(context, other_id, confirmed=True)
    dwg.unlink()
    assert verify(context, other_id)["status"] == "MISSING"


def test_manifest_write_failure_never_publishes_package(
    tmp_path: Path, monkeypatch
):
    context, _, cqt = project(tmp_path, "A057349839")
    original = Path.write_text

    def fail_checksum(path, content, *args, **kwargs):
        if path.name == "manifest.sha256":
            raise OSError("write interrupted")
        return original(path, content, *args, **kwargs)

    monkeypatch.setattr(Path, "write_text", fail_checksum)
    with pytest.raises(OSError, match="write interrupted"):
        generate(context, [cqt.name], validation(context))
    store = context.project_root / ".oxcad" / "deliveries"
    assert list(store.iterdir()) == []


def test_cli_jsonl_events_and_stale_document(
    tmp_path: Path, simulated_host_protocol
):
    configurar_logging(diretorio_logs=tmp_path)
    context, _, cqt = project(tmp_path, "A057349839")

    def action(name, active=context, **kwargs):
        return executar_acao_operacional(
            name, dwg_path=active.dwg_path, sob_id_esperada=active.sob_id,
            dwg_identity=active.dwg_identity, execution_id=active.execution_id,
            correlation_id=active.correlation_id,
            execution_origin=active.origin,
            host_context=simulated_host_protocol(
                active.dwg_path, active.sob_id, active.dwg_identity,
                active.execution_id, active.correlation_id,
            ),
            **kwargs,
        )

    created = action(
        "entrega_gerar", delivery_artifacts=[cqt.name],
        delivery_validations=validation(context),
    )
    assert created["sucesso"] is True
    package_id = created["delivery_result"]["package_id"]
    stale_path = context.project_root / "SECOND.dwg"
    stale_path.write_bytes(b"synthetic different DWG")
    stale = replace(context, dwg_path=str(stale_path), dwg_identity="other-fp")
    blocked = action("entrega_aprovar", stale, package_id=package_id,
                     approval_confirmed=True)
    assert blocked["sucesso"] is False
    current = action("entrega_verificar", package_id=package_id)
    assert current["delivery_result"]["status"] == "UNAPPROVED"
    approved = action("entrega_aprovar", package_id=package_id,
                      approval_confirmed=True)
    assert approved["delivery_result"]["status"] == "CONFIRMED"
    cqt.write_bytes(b"modified")
    failed = action("entrega_verificar", package_id=package_id)
    assert failed["sucesso"] is False
    assert failed["delivery_result"]["status"] == "DIVERGENT"
    logs = (tmp_path / "logs" / "application.jsonl").read_text(
        encoding="utf-8"
    ).splitlines()
    events = [json.loads(line) for line in logs]
    assert any(event.get("delivery_package_id") == package_id
               and event.get("delivery_status") == "DIVERGENT"
               and event.get("dwg_identity") == context.dwg_identity
               for event in events)
    assert any(event.get("delivery_package_id") == package_id
               and event.get("delivery_approval_state") == "APROVADO"
               and event.get("execution_id") == context.execution_id
               and event.get("correlation_id") == context.correlation_id
               for event in events)
    legacy = json.loads('{"level":"INFO","message":"old event"}')
    assert legacy["message"] == "old event"
